#!/usr/bin/env python3
"""Generate the government review workbook from the run SQLite database only.

Sheets follow skills/silverleaf-create-lead-list/references/workbook-layout.md: Organisations holds the government
offices and Contacts the official posts, and Parent Enquiries stays empty by design. Government sheets follow
plans/b2b-government-leads-plan.md §5.8. The workbook is written to outputs/runs/<run-id>/, never inside
outputs/master/.
"""
from __future__ import annotations

import argparse
import json
import sqlite3
from collections import Counter
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import gov_lib as G

BLUE, WHITE = "002368", "FFFFFF"
RISK_FILL = {"low": "E2F0D9", "medium": "FFF2CC", "risky": "F8CBAD"}
WRAP = {"track_reason", "location_note", "convening_forum", "size_evidence", "pdpa_risk_reason", "missing_information", "evidence_excerpt",
        "all_source_urls", "convening_role", "notes", "channel_attribution", "census_source_location", "detail", "values_json", "action", "title",
        "proposed_treatment", "note", "score_factors", "meaning", "silverleaf_contribution", "do_not_imply", "gate", "activation_gate", "text"}
ORG_HEADERS = ["organisation_id", "record_key", "name", "office_level", "proposed_government_track", "track_reason", "ward_tier", "ward_score",
               "rank_in_cluster", "campus_cluster", "factors_known", "council_name", "region", "admin_unit_name", "parent_admin_unit",
               "census_2022_population", "census_source_location", "catchment_wards", "catchment_population_2022", "campus", "distance_km",
               "transport_band", "nearest_primary_campus", "distance_to_primary_km", "primary_band", "geocode_precision", "location_source",
               "latitude", "longitude", "location_note", "website", "public_email", "public_phone", "po_box", "convening_forum", "size_evidence",
               "verification_status", "pdpa_risk", "pdpa_risk_reason", "osm_id", "master_organisation_id", "missing_information", "source_url",
               "evidence_basis", "evidence_excerpt", "all_source_urls"]
CONTACT_HEADERS = ["contact_id", "organisation", "office_level", "role", "appointment_type", "name", "holder_verified_on", "tenure_source_url",
                   "role_certainty", "contact_route", "role_phone", "published_role_email", "organisation_phone", "channel_attribution",
                   "b2g_check_required", "convening_role", "pdpa_risk", "pdpa_risk_reason", "proposed_government_track", "verification_status",
                   "source_url", "evidence_excerpt", "notes"]
UNIT_HEADERS = ["unit_key", "level", "name", "council", "region", "in_catchment", "ward_tier", "ward_score", "rank_in_cluster", "campus_cluster",
                "population_2022", "households_2022", "avg_household_size_2022", "male_2022", "female_2022", "census_source_location", "campus",
                "distance_km", "transport_band", "nearest_primary_campus", "distance_to_primary_km", "primary_band", "geocode_precision",
                "location_source", "latitude", "longitude", "location_note", "location_url", "factors_known", "score_factors",
                "office_organisation_id", "census_url"]
EVENT_HEADERS = ["event_code", "organisation_id", "event_type", "event_date", "event_status", "expected_attendance", "attendance_estimate",
                 "attendance_method", "opt_ins", "enquiries", "tours", "applications", "enrolments", "cost_tzs"]
TRIAGE_HEADERS = ["source", "source_id", "name", "category", "detail", "tags", "osm_district", "campus", "distance_km", "linked_organisation_id", "note"]
MASTER_HEADERS = ["master_organisation_id", "name", "triage_group", "category", "detail", "proposed_treatment", "master_campus", "master_distance_km",
                  "matched_osm_id", "linked_organisation_id"]
SIGNAL_HEADERS = ["item_date", "organisation", "kind", "title", "tags", "item_id", "listing_url"]
REVIEW_HEADERS = ["review_id", "kind", "entity_type", "entity_id", "field", "values_json", "action", "status"]
FORUM_ROWS = [
    ["Village assembly (mkutano mkuu wa kijiji)", "Village chairperson, with the village executive officer (VEO) as secretary",
     "All residents aged 18 and over; at least quarterly. Rural councils (Arusha District, Meru, Hai).", "Main v1 forum in rural wards"],
    ["Mtaa meeting", "Mtaa chairperson, with the mtaa executive officer (MEO)", "Urban wards in Arusha City and Moshi Municipal. Confirm how often they meet.",
     "Main v1 forum in urban wards"],
    ["Ward public meeting", "Ward executive officer (WEO), often with the councillor",
     "The Ward Development Committee (at least four meetings a year) is a committee, not a public forum.", "v1 forum; reached after the council introduction"],
    ["District or regional public meeting (mkutano wa hadhara)", "District or Regional Commissioner's office", "Large, formal protocol.", "Not a v1 target"],
    ["Government-organised public events", "Council or region", "Institutions can take stalls, e.g. Nane Nane at the Njiro grounds (August).",
     "Later; T-21 production clock"],
    ["Community development groups", "Council and ward community development officers", "Women's, youth and savings groups.",
     "Savings groups go to corporate campaign C02"],
]


NUMERIC = {"census_2022_population", "catchment_wards", "catchment_population_2022", "ward_score", "rank_in_cluster", "distance_to_primary_km",
           "distance_km", "latitude", "longitude", "desk_score"}


def as_number(header, value):
    """Intake payload values are strings; show the numeric ones as numbers."""
    if header not in NUMERIC or value in (None, ""):
        return value
    try:
        number = float(value)
    except (TypeError, ValueError):
        return value
    return int(number) if number.is_integer() and header not in ("latitude", "longitude", "distance_km", "distance_to_primary_km", "ward_score") else number


def add_sheet(wb, title, subtitle, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws["A1"], ws["A2"] = title, subtitle
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color=BLUE)
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="818283")
    ws.append([])
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    last = 4 + len(rows)
    for col, header in enumerate(headers, 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = (widths or {}).get(header, 48 if header in WRAP else min(40, max(12, len(header) + 2)))
        if header in WRAP:
            for cell in ws[letter][4:last]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if header.endswith("url") or header in ("website", "all_source_urls"):
            for cell in ws[letter][4:last]:
                if isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value.split("\n")[0]
                    cell.font = Font(name="Arial", size=10, color="1F4E9A", underline="single")
        if header == "pdpa_risk":
            for cell in ws[letter][4:last]:
                if cell.value in RISK_FILL:
                    cell.fill = PatternFill("solid", fgColor=RISK_FILL[cell.value])
    ws.freeze_panes = "C5"
    if rows:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{last}"
    return ws


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-id", required=True)
    args = parser.parse_args()
    cfg = G.load_config(args.run_id)
    db, workbook_path = cfg["paths"]["db"], cfg["paths"]["workbook"]
    master_dir = (G.ROOT / "outputs" / "master").resolve()
    if master_dir in workbook_path.resolve().parents:
        raise SystemExit("Refusing to write the government workbook inside outputs/master: it must stay separate from the company-leads master.")
    con = sqlite3.connect(f"file:{db.as_posix()}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    payload = {}
    for r in con.execute("SELECT es.entity_type, es.entity_id, sr.payload_json FROM entity_sources es JOIN source_records sr "
                         "ON sr.source_record_id = es.source_record_id"):
        payload.setdefault(r["entity_id"], json.loads(r["payload_json"]))

    org_rows, names, evidence, sources = [], {}, [], {}
    orgs = [dict(r) for r in con.execute("SELECT o.*, g.office_level, g.proposed_government_track FROM organisations o "
                                         "JOIN government_office_profiles g USING(organisation_id)")]
    level_order = {"council": 0, "ward": 1, "village": 2, "district": 3, "region": 4}
    orgs.sort(key=lambda o: (not str(o["proposed_government_track"]).startswith("GA01"), level_order.get(o["office_level"], 9),
                             -(o["desk_score"] or -1), o["name"]))
    for o in orgs:
        p = payload.get(o["organisation_id"], {})
        srcs = json.loads(p.get("all_sources_json") or "[]")
        merged = {**p, **{k: v for k, v in o.items() if v not in (None, "")}, "name": o["name"], "latitude": o["latitude"], "longitude": o["longitude"],
                  "public_email": o["email"], "public_phone": o["phone"], "all_source_urls": "\n".join(s["url"] for s in srcs if s.get("url"))}
        names[o["organisation_id"]] = o["name"]
        org_rows.append([as_number(h, merged.get(h, "")) for h in ORG_HEADERS])
        for s in srcs:
            evidence.append(["organisation", o["organisation_id"], o["name"], s.get("url"), s.get("title"), s.get("accessed_on"), s.get("evidence_basis"),
                             "yes" if s.get("fetched") else "no", ", ".join(s.get("facts_supported") or []), s.get("evidence_excerpt", "")])
            entry = sources.setdefault(s.get("url"), {"title": s.get("title"), "basis": set(), "cited": 0})
            entry["basis"].add(s.get("evidence_basis"))
            entry["cited"] += 1

    contact_rows = []
    for c in con.execute("SELECT c.*, o.name AS organisation, op.office_title, op.appointment_type, op.b2g_check_required, op.convening_role, "
                         "op.holder_verified_on, op.tenure_source_url, op.pdpa_risk, op.pdpa_risk_reason, g.office_level, g.proposed_government_track "
                         "FROM contacts c JOIN organisations o USING(organisation_id) JOIN official_posts op ON op.contact_id = c.contact_id "
                         "JOIN government_office_profiles g ON g.organisation_id = c.organisation_id ORDER BY o.name, op.post_key"):
        c = dict(c)
        p = payload.get(c["contact_id"], {})
        merged = {**p, **{k: v for k, v in c.items() if v not in (None, "")}, "name": c["name"] or "", "organisation": c["organisation"]}
        contact_rows.append([merged.get(h, "") for h in CONTACT_HEADERS])
        for s in json.loads(p.get("all_sources_json") or "[]"):
            evidence.append(["contact", c["contact_id"], f"{c['role']} — {c['organisation']}", s.get("url"), s.get("title"), s.get("accessed_on"),
                             s.get("evidence_basis"), "yes" if s.get("fetched") else "no", ", ".join(s.get("facts_supported") or []), s.get("evidence_excerpt", "")])
            entry = sources.setdefault(s.get("url"), {"title": s.get("title"), "basis": set(), "cited": 0})
            entry["basis"].add(s.get("evidence_basis"))
            entry["cited"] += 1

    units = [dict(r) for r in con.execute("SELECT * FROM admin_units")]
    units.sort(key=lambda u: (u["level"] != "ward", {"yes": 0, "unknown": 1, "no": 2}.get(u["in_catchment"], 3), u["campus_cluster"] or "",
                              u["rank_in_cluster"] or 999, u["council"], u["name"]))
    unit_rows = [[u.get(h) for h in UNIT_HEADERS] for u in units]
    event_rows = [[r[h] for h in EVENT_HEADERS] for r in con.execute("SELECT * FROM community_events ORDER BY event_date")]
    triage_rows = [[r[h] for h in TRIAGE_HEADERS] for r in con.execute(
        "SELECT source, source_id, name, category, detail, tags, osm_district, campus, distance_km, linked_organisation_id, note FROM office_triage "
        "ORDER BY CASE category WHEN 'convener' THEN 0 WHEN 'excluded' THEN 1 WHEN 'not_government' THEN 2 WHEN 'agency' THEN 3 ELSE 4 END, distance_km")]
    master_rows = [[r[h] for h in MASTER_HEADERS] for r in con.execute("SELECT * FROM master_triage ORDER BY triage_group, master_distance_km")]
    signal_rows = [[r[h] for h in SIGNAL_HEADERS] for r in con.execute("SELECT * FROM convening_signals ORDER BY item_date DESC, organisation")]
    review_rows = [[r[h] for h in REVIEW_HEADERS] for r in con.execute(f"SELECT {', '.join(REVIEW_HEADERS)} FROM review ORDER BY kind, entity_id")]
    source_rows = [[u, (u or "").split("/")[2] if "//" in (u or "") else "", v["title"], v["cited"], ", ".join(sorted(b for b in v["basis"] if b))]
                   for u, v in sorted(sources.items(), key=lambda kv: -kv[1]["cited"]) if u]
    counts = {t: con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in
              ("organisations", "contacts", "enquiries", "admin_units", "community_events", "office_triage", "master_triage", "convening_signals",
               "review", "source_records", "source_files")}
    wards = [u for u in units if u["level"] == "ward"]
    catch = [u for u in wards if u["in_catchment"] == "yes"]
    track = Counter(o["proposed_government_track"] for o in orgs)
    level = Counter(o["office_level"] for o in orgs)
    tier = Counter(u["ward_tier"] for u in catch)
    risk = Counter(r[CONTACT_HEADERS.index("pdpa_risk")] for r in contact_rows)
    named = sum(1 for r in contact_rows if r[CONTACT_HEADERS.index("name")])
    with_phone = sum(1 for r in contact_rows if r[CONTACT_HEADERS.index("role_phone")])
    locsrc = Counter(u["location_source"] for u in wards)
    excluded = Counter(r[4] for r in triage_rows if r[3] == "excluded")
    core = {c["name"] for c in cfg["councils"] if c.get("core")}
    clusters = {}
    for u in sorted((u for u in catch if u["council"] in core), key=lambda u: u["rank_in_cluster"]):
        if len(clusters.setdefault(u["campus_cluster"], [])) < 5:
            clusters[u["campus_cluster"]].append(u)
    wb = Workbook()
    ws = wb.active
    ws.title = "Read Me"
    lines = [(f"Silverleaf government leads — research run {args.run_id}", "title"),
             (f"Generated {generated} from {db.name} (SQLite is the source of truth; this workbook is a generated view). "
              "Separate from the company-leads master in outputs/master/ and from the welfare runs.", "sub"), ("", ""),
             ("What this is", "h"),
             ("Local government offices that can convene community meetings (barazas) near Silverleaf's five campuses, so that Silverleaf can meet "
              "parents at official forums. A government lead is the office, not the person holding it. Parents are never collected here: "
              "they opt in with Silverleaf at an event, into a separate consented store.", ""),
             ("Research only: no letter, message or outreach plan was drafted, sent or scheduled, and every automation stays disabled.", ""), ("", ""),
             ("Start here: counts", "h"),
             (f"Offices {counts['organisations']} (" + ", ".join(f"{k} {v}" for k, v in sorted(level.items(), key=lambda kv: level_order.get(kv[0], 9))) +
              f") · Official posts {counts['contacts']} ({named} with a holder named by an official source, {with_phone} with a phone the council published) · "
              f"Wards {len(wards)} ({len(catch)} within {cfg['catchment_km']} km of a campus) · Review items {counts['review']}", ""),
             ("Offices by proposed track: " + "; ".join(f"{k} {v}" for k, v in track.most_common()), ""),
             ("Catchment wards by priority tier: " + "; ".join(f"{k} {v}" for k, v in sorted(tier.items())), ""),
             ("Official posts by data-protection risk: " + "; ".join(f"{k} {v}" for k, v in risk.most_common()), ""),
             ("Ward locations by source: " + "; ".join(f"{k} {v}" for k, v in locsrc.most_common()), ""),
             ("Office triage exclusions (never contacted): " + ("; ".join(f"{k} {v}" for k, v in excluded.most_common()) or "none"), ""), ("", ""),
             ("Top five wards per campus cluster in the core councils (pilot candidates, plan §5.10; rank_in_cluster counts every catchment ward)", "h")]
    for cluster, items in sorted(clusters.items()):
        lines.append((f"{cluster}: " + "; ".join(f"{u['name']} ({u['council'].replace(' Council', '')}, {u['population_2022']:,} people, "
                                                 f"{u['distance_km']} km, score {u['ward_score']})" for u in sorted(items, key=lambda u: u['rank_in_cluster'])), ""))
    lines += [("", ""), ("How to use it", "h"),
              ("1. GA01 councils first: deliver the Kiswahili introduction letter to each council director (plans/b2b-government-leads-plan.md §5.6), "
               "call after 5 working days, then visit. Ask for guidance, the ward meeting calendars and an introduction to the ward executive officers.", ""),
              ("2. Then choose wards by tier and cluster rank. GA02 (a convening request) needs the council introduction and the ward post verified "
               "within 90 days, so no row is GA02 yet.", ""),
              ("3. Check with the B2G owner before contacting any education officer (b2g_check_required = yes).", ""),
              ("4. Work through the Review Queue: unlocated wards, conflicting points, roster spellings and master corrections.", ""), ("", ""),
              ("Scores and tiers", "h"),
              ("ward_score (0–100, plan §5.5): 2022 population relative to the other catchment wards (30), transport band to a campus (25; half for the "
               "nearest campus, half for the nearest primary campus), access readiness (20), convening opportunity (15) and observed yield (10). "
               "The last three are zero until introductions, confirmed meetings and events are recorded, so today's maximum is 55. "
               "Tiers: P1 40+, P2 30–39.9, P3 20–29.9, P4 below 20. The score never uses political affiliation, ethnicity, religion or income.", ""),
              ("Ward locations are estimates (geocode_precision = estimated): a mapped ward office where one exists; otherwise a Wikipedia ward point "
               "checked against the council's district and corroborated by OpenStreetMap; otherwise OpenStreetMap places named like the ward. "
               "Distances are straight lines to approximate campus centroids.", ""), ("", ""),
              ("Data protection (pdpa_risk)", "h"),
              ("low — offices, role desks and organisation-level facts.", ""),
              ("medium — an official's name as published by the office (council profiles and councillor lists), and phone numbers the council "
               "published for contacting its councillors. Use them only in the official role and after the council introduction.", ""),
              ("risky — none in this run. Parent, resident, voter, beneficiary and pupil data are never collected; party affiliation is not recorded "
               "even where a list shows it.", ""),
              ("Tanzania's Personal Data Protection Act 2022 requires explicit consent for direct marketing (per DLA Piper's summary). Whether a formal "
               "letter to a public office counts as direct marketing is decision 4 in the plan. Not legal advice.", ""), ("", ""),
              ("Coverage and gaps", "h")]
    lines += [(note, "") for note in cfg.get("readme_notes", [])]
    lines += [("", ""), ("Sheets", "h"),
              ("Organisations (government offices) · Contacts (official posts) · Administrative Units (wards and councils, census 2022) · "
               "Community Events (counts only; empty until events run) · Convening Forums · Convening Signals (council news and notices) · "
               "Office Triage (OpenStreetMap offices and exclusions) · Master Triage (the master's 29 office:government records) · Review Queue · "
               "Evidence · Sources · Parent Enquiries (empty by design) · then the layout sheets, which hold design-only material.", "")]
    for i, (text, kind) in enumerate(lines, 1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(name="Arial", size=16 if kind == "title" else 11 if kind == "h" else 10, bold=kind in ("title", "h"),
                         italic=kind == "sub", color=BLUE if kind in ("title", "h") else "000000")
    ws.column_dimensions["A"].width = 150

    add_sheet(wb, "Organisations", "Government offices: one row per convening office. GA01 councils first, then wards by score. Filter office_level, "
              "proposed_government_track, ward_tier and campus_cluster.", ORG_HEADERS, org_rows, {"name": 44, "track_reason": 50})
    add_sheet(wb, "Contacts", "Official posts: role desks, with a holder only where an official source names one. Check pdpa_risk and "
              "b2g_check_required before any use.", CONTACT_HEADERS, contact_rows, {"role": 44, "organisation": 40})
    add_sheet(wb, "Administrative Units", "Every ward of the councils studied, with 2022 census population (table, page and row), an estimated "
              "location and campus distances. Catchment wards first, by campus cluster and rank.", UNIT_HEADERS, unit_rows, {"name": 22})
    add_sheet(wb, "Community Events", "Counts only (plan §5.2): no names, phone numbers or other person-level fields. Empty until events are held.",
              EVENT_HEADERS, event_rows)
    add_sheet(wb, "Convening Forums", "The forums behind 'baraza' (plan §1). Confirm the legal sections before quoting them externally.",
              ["forum", "convened_by", "meaning", "use_in_v1"], FORUM_ROWS, {"forum": 40, "convened_by": 45, "use_in_v1": 36})
    add_sheet(wb, "Convening Signals", "Council and regional news or notices that mention public meetings, events, ward or village meetings, "
              "training for executive officers, or elections. Past events show how each office convenes; election notices mark hold periods.",
              SIGNAL_HEADERS, signal_rows, {"title": 80, "tags": 40})
    add_sheet(wb, "Office Triage", "Every named OpenStreetMap government office in the catchment, classified. Excluded rows are never contacted; "
              "only a documented review decision can change one.", TRIAGE_HEADERS, triage_rows, {"name": 44})
    add_sheet(wb, "Master Triage", "The company master's office:government records, classified as plan §2.1 proposes. Apply changes to the "
              "master only through silverleaf-update-lead-list.", MASTER_HEADERS, master_rows, {"name": 44})
    add_sheet(wb, "Review Queue", "Decisions for a person: unlocated wards, conflicting points, roster spellings, estimated village wards and "
              "master corrections.", REVIEW_HEADERS, review_rows, {"entity_id": 40, "values_json": 70})
    add_sheet(wb, "Evidence", "One row per entity and source, with the facts it supports.",
              ["entity_type", "entity_id", "entity", "url", "title", "accessed_on", "evidence_basis", "fetched", "facts_supported", "evidence_excerpt"],
              evidence, {"entity": 44, "url": 50})
    add_sheet(wb, "Sources", "Unique source URLs.", ["url", "domain", "title", "times_cited", "evidence_basis"], source_rows, {"url": 70, "title": 50})
    add_sheet(wb, "Parent Enquiries", "Empty by design: parents opt in with Silverleaf at events, into a separate consented store (plan §3). "
              "The master and this workbook hold counts only.", ["enquiry_id", "type", "enquiry_date", "request", "platform", "source_url"], [])
    note = "Design only; nothing is drafted, scheduled or enabled in this research run. See plans/b2b-government-leads-plan.md §5.6–5.7."
    add_sheet(wb, "Outreach Plans", note, ["status"], [["No outreach plans drafted in this research-only run."]])
    add_sheet(wb, "Campaigns", note, ["campaign_id", "name", "audience", "activation_gate", "status"],
              [["C11", "Community convening via local government", "GA01 councils, then GA02 ward, village and mtaa offices",
                "Council introduction recorded; approved Kiswahili session content, guide and privacy notice; live consented parent store or an explicit "
                "inbound-only decision; presenter briefed on neutrality and safeguarding; materials budget; B2G owner informed", "Design only; not activated"]])
    add_sheet(wb, "Touchpoints", note, ["flow_id", "step", "trigger", "action", "pipeline_stage"],
              [["F15", 1, "Council or district office selected for GA01", "Deliver the formal Kiswahili letter", "letter_sent"],
               ["F15", 2, "No response after 5 working days", "Call the office", "letter_sent"],
               ["F15", 3, "After the call", "Courtesy visit; record guidance and the introduction", "introduced"],
               ["F15", 4, "Introduction recorded and the ward post verified within 90 days", "GA02: learn the meeting calendar; request a slot in writing", "convening_requested"],
               ["F15", 5, "Slot agreed in writing", "Create a community_events row (counts only) and hand over to F16", "convening_agreed"],
               ["F16", 1, "T-7 (T-21 for bespoke events)", "Print the guide and forms with the event code; brief the presenter", "scheduled"],
               ["F16", 2, "T-2", "Confirm logistics with the convener", "scheduled"],
               ["F16", 3, "Event day", "VM14 talk, VM15 guide, VM16 voluntary opt-in; record the attendance estimate and method", "event_held"],
               ["F16", 4, "Within 24 hours", "Opt-ins into the consented parent store; secure or destroy paper forms", "event_held"],
               ["F16", 5, "Within 5 working days", "Thank-you letter with aggregate feedback (VM17)", "event_held / repeat_convener"]])
    add_sheet(wb, "Assignments", note, ["status"], [["No campaign assignments in this research-only run."]])
    add_sheet(wb, "Strategies", note, ["module_id", "name", "silverleaf_contribution", "do_not_imply", "gate"],
              [["VM14", "Community education session", "15–20 minute Kiswahili talk: school readiness, early learning at home, choosing a school, admissions calendar",
                "Government endorsement; criticism of public schools", "Content approved once, with native-speaker review"],
               ["VM15", "Take-home parent guide", "One-page Kiswahili guide with published levels, fees and intake windows", "Scholarships or discounts unless approved",
                "Marketing and admissions approval"],
               ["VM16", "Voluntary opt-in", "Silverleaf's own form with a Kiswahili privacy notice; the office never collects or shares details",
                "That the office endorses signing up", "Notice reviewed; consented parent store live"],
               ["VM17", "Aggregate feedback", "Attendance and follow-up counts, no names, within 5 working days", "Any obligation or performance claim", "None"],
               ["VM18", "Approved scholarship or bursary information", "Details of an approved scheme", "Any unapproved scheme", "Blocked until approved"]])
    add_sheet(wb, "Automation", note, ["flow_id", "status"], [["F15", "DISABLED — design only"], ["F16", "DISABLED — design only"]])
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(workbook_path)
    sheet_counts = {"Organisations": len(org_rows), "Contacts": len(contact_rows), "Administrative Units": len(unit_rows),
                    "Community Events": len(event_rows), "Office Triage": len(triage_rows), "Master Triage": len(master_rows),
                    "Convening Signals": len(signal_rows), "Review Queue": len(review_rows)}
    db_counts = {"Organisations": counts["organisations"], "Contacts": counts["contacts"], "Administrative Units": counts["admin_units"],
                 "Community Events": counts["community_events"], "Office Triage": counts["office_triage"], "Master Triage": counts["master_triage"],
                 "Convening Signals": counts["convening_signals"], "Review Queue": counts["review"]}
    mismatch = {k: (sheet_counts[k], v) for k, v in db_counts.items() if sheet_counts[k] != v}
    con.close()
    print(json.dumps({"workbook": str(workbook_path.relative_to(G.ROOT)), "sheet_rows": sheet_counts, "mismatches": mismatch}, indent=1))
    return 0 if not mismatch else 1


if __name__ == "__main__":
    raise SystemExit(main())
