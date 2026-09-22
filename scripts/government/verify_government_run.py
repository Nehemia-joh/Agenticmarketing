#!/usr/bin/env python3
"""Verify a government run: integrity, reconciliation, exclusions, data-protection guards, offer-aligned letters and master separation.

Writes outputs/runs/<run-id>/run-verification.json and exits 1 if any check fails. run_pipeline.py records the
company-master baseline at the start of the run (runtime/government/<run-id>/master-baseline.json).
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sqlite3
import sys
from datetime import datetime, timezone

from openpyxl import load_workbook

import gov_lib as G

sys.path.insert(0, str(G.ROOT / "scripts" / "messaging"))
import offer_lib as OL  # noqa: E402  (offer-aligned draft checks)

W = G.W
LAYOUT_SHEETS = ["Read Me", "Organisations", "Contacts", "Parent Enquiries", "Outreach Plans", "Campaigns", "Touchpoints", "Assignments",
                 "Strategies", "Evidence", "Sources", "Review Queue", "Automation", "Administrative Units", "Community Events"]
SENSITIVE_COLUMNS = re.compile(r"party|chama|politic|religio|dini|ethnic|kabila|tribe|voter|resident_|pupil|student|child", re.I)
PERSON_COLUMNS = re.compile(r"name|phone|email|mobile|profile|whatsapp|address", re.I)


def has_key(value, key) -> bool:
    if isinstance(value, dict):
        return key in value or any(has_key(v, key) for v in value.values())
    if isinstance(value, list):
        return any(has_key(v, key) for v in value)
    return False


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-id", required=True)
    args = parser.parse_args()
    cfg = G.load_config(args.run_id)
    paths = cfg["paths"]
    master_dir = (G.ROOT / "outputs" / "master").resolve()
    baseline_path = paths["work"] / "master-baseline.json"
    baseline = json.loads(baseline_path.read_text(encoding="utf-8")) if baseline_path.exists() else {}
    con = sqlite3.connect(f"file:{paths['db'].as_posix()}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    integrity = con.execute("PRAGMA integrity_check").fetchone()[0]
    fks = [tuple(r) for r in con.execute("PRAGMA foreign_key_check").fetchall()]
    tables = ("organisations", "contacts", "enquiries", "outreach_plans", "automation_recipes", "admin_units", "government_office_profiles",
              "official_posts", "community_events", "office_triage", "master_triage", "convening_signals", "review", "source_records", "source_files")
    counts = {t: con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in tables}
    columns = {t: [r[1] for r in con.execute(f'PRAGMA table_info("{t}")')] for (t,) in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    sensitive = [f"{t}.{c}" for t, cols in columns.items() for c in cols if SENSITIVE_COLUMNS.search(c)]
    event_person = [c for c in columns.get("community_events", []) if PERSON_COLUMNS.search(c)]
    payloads = [json.loads(p) for (p,) in con.execute("SELECT payload_json FROM source_records")]
    missing_risk = [p.get("record_key") for p in payloads if p.get("pdpa_risk") not in ("low", "medium", "risky")]
    excluded_names = [r["name"] for r in con.execute("SELECT name FROM organisations") if G.exclusion_reason(r["name"])]
    named_without_tenure = [r["record_key"] for r in con.execute("SELECT record_key FROM official_posts WHERE holder_name != '' AND "
                                                                 "(holder_verified_on = '' OR tenure_source_url = '')")]
    phone_rows = [dict(r) for r in con.execute("SELECT c.contact_id, c.role_phone, c.channel_attribution, op.pdpa_risk FROM contacts c "
                                               "JOIN official_posts op ON op.contact_id = c.contact_id WHERE c.role_phone != ''")]
    phones_unattributed = [r["contact_id"] for r in phone_rows if not r["channel_attribution"] or r["pdpa_risk"] == "low"]
    ga02 = con.execute("SELECT COUNT(*) FROM government_office_profiles WHERE proposed_government_track LIKE 'GA02%'").fetchone()[0]
    excluded_linked = con.execute("SELECT COUNT(*) FROM office_triage WHERE category = 'excluded' AND linked_organisation_id IS NOT NULL").fetchone()[0]
    b2g_unflagged = con.execute("SELECT COUNT(*) FROM official_posts WHERE post_key IN ('primary_education_head', 'regional_education') "
                                "AND b2g_check_required != 'yes'").fetchone()[0]
    message_checks = OL.run_message_checks(con, convening=True)
    undrafted_offices = con.execute("SELECT COUNT(*) FROM government_office_profiles g WHERE NOT EXISTS (SELECT 1 FROM outreach_plans p WHERE p.target_id = g.organisation_id)").fetchone()[0]
    catch_ward_offices = con.execute("SELECT COUNT(*) FROM admin_units WHERE level = 'ward' AND in_catchment = 'yes' AND office_organisation_id IS NULL").fetchone()[0]
    con.close()
    totals = json.loads((paths["raw"] / f"nbs_2022_councils_{cfg['research_date']}.json").read_text(encoding="utf-8"))
    census_ok = all(v["reconciles"] for v in totals["reconciliation"].values()) and not totals["councils_missing"]
    cms_leaks = [p.name for p in (paths["raw"] / "council_sites").glob("*.json")
                 if has_key(json.loads(p.read_text(encoding="utf-8")), "createdBy")]
    with open(paths["intake"], encoding="utf-8") as handle:
        intake_rows = list(csv.DictReader(handle))
    intake_by_type = {k: sum(1 for r in intake_rows if r["record_type"] == k) for k in ("organisation", "contact", "enquiry")}
    wb = load_workbook(paths["workbook"], read_only=True)
    sheet_rows = {name: max(0, wb[name].max_row - 4) for name in ("Organisations", "Contacts", "Administrative Units", "Community Events",
                                                                    "Office Triage", "Master Triage", "Convening Signals", "Review Queue",
                                                                    "Outreach Plans")}
    missing_sheets = [s for s in LAYOUT_SHEETS if s not in wb.sheetnames]
    drift = []
    if G.SKILL_SCRIPTS.is_dir():
        canonical = {p.name: W.sha256_file(p) for p in G.CANONICAL_SCRIPTS.glob("*.py")}
        bundled = {p.name: W.sha256_file(p) for p in G.SKILL_SCRIPTS.glob("*.py")}
        drift = sorted(n for n in canonical.keys() | bundled.keys() if canonical.get(n) != bundled.get(n))
    checks = {
        "integrity": integrity == "ok",
        "foreign_keys": not fks,
        "organisations_reconcile": counts["organisations"] == intake_by_type["organisation"] == sheet_rows["Organisations"] == counts["government_office_profiles"],
        "contacts_reconcile": counts["contacts"] == sheet_rows["Contacts"] == counts["official_posts"] and counts["contacts"] <= intake_by_type["contact"],
        "admin_units_reconcile": counts["admin_units"] == sheet_rows["Administrative Units"],
        "triage_and_signals_reconcile": (counts["office_triage"] == sheet_rows["Office Triage"] and counts["master_triage"] == sheet_rows["Master Triage"]
                                         and counts["convening_signals"] == sheet_rows["Convening Signals"]),
        "review_reconcile": counts["review"] == sheet_rows["Review Queue"],
        "census_wards_reconcile_to_council_totals": census_ok,
        "every_catchment_ward_has_an_office": catch_ward_offices == 0,
        "every_row_risk_labelled": not missing_risk,
        "no_excluded_office_in_organisations": not excluded_names,
        "no_excluded_office_linked": excluded_linked == 0,
        "no_sensitive_columns": not sensitive,
        "community_events_counts_only": not event_person and sheet_rows["Community Events"] == counts["community_events"],
        "no_parent_or_resident_data": counts["enquiries"] == 0 and intake_by_type["enquiry"] == 0,
        "named_holders_have_tenure_source": not named_without_tenure,
        "published_phones_attributed_and_not_low_risk": not phones_unattributed,
        "education_posts_flagged_for_b2g": b2g_unflagged == 0,
        "no_ga02_without_introduction": ga02 == 0,
        "no_cms_editor_data_in_raw_captures": not cms_leaks,
        "outreach_drafts_only_no_automation": set(message_checks["statuses"]) <= OL.DRAFT_STATUSES and counts["automation_recipes"] == 0,
        "letters_conform_to_offer_register_and_offer_officials_nothing": not message_checks["issues"] and message_checks["orphan_messages"] == 0,
        "every_office_has_a_draft": undrafted_offices == 0,
        "outreach_plans_reconcile": counts["outreach_plans"] == sheet_rows["Outreach Plans"],
        "layout_sheets_present": not missing_sheets,
        "raw_evidence_registered": counts["source_files"] > 1,
        "outputs_separate_from_master": master_dir not in paths["workbook"].resolve().parents and master_dir not in paths["db"].resolve().parents,
        "company_master_db_unchanged": bool(baseline) and W.sha256_file(W.MASTER) == baseline.get("master_db_sha256"),
        "company_master_workbook_unchanged": bool(baseline) and W.sha256_file(master_dir / "Silverleaf Master Database - Consolidated.xlsx") == baseline.get("master_workbook_sha256"),
        "skill_script_copies_in_sync": not drift,
    }
    report = {"checked_at": datetime.now(timezone.utc).isoformat(), "run_id": args.run_id, "result": "passed" if all(checks.values()) else "failed",
              "checks": checks, "db_counts": counts, "intake_by_type": intake_by_type, "sheet_rows": sheet_rows,
              "details": {"missing_risk": missing_risk[:20], "excluded_names": excluded_names, "sensitive_columns": sensitive,
                          "community_event_person_columns": event_person, "named_without_tenure": named_without_tenure[:20],
                          "phones_unattributed": phones_unattributed, "cms_leaks": cms_leaks, "missing_sheets": missing_sheets,
                          "foreign_key_violations": fks[:10], "master_baseline": baseline, "script_copies_out_of_sync": drift,
                          "catchment_wards_without_office": catch_ward_offices, "ga02_rows": ga02,
                          "message_issues": message_checks["issues"][:20], "messages_checked": message_checks["messages_checked"]}}
    (paths["run_out"] / "run-verification.json").write_text(json.dumps(report, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"result": report["result"], "failed": [k for k, v in checks.items() if not v], "db_counts": counts}, indent=1))
    return 0 if report["result"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
