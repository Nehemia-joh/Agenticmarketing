#!/usr/bin/env python3
"""Link VICOBA savings groups to their introducers through the reviewed links file, and record the band properties.

Decided 13 October 2026 (docs/plans/vicoba-integration-plan.md, decision 3; the user's approval of the same day): group-level
records only; introducers (NGO field offices, bank branches, faith offices, employers) live in the master; CDO routes come through
the government run and are referenced there, never copied; a group leader's phone counts only when published for the role; a
number given with consent never enters the repository; women's groups go first.

Links go only through data/reference/vicoba-introducers.json (the reviewed file). An intake's introducer_ref, sponsor_ref or
employer_ref is a proposal, never a link. The file holds:
- "links": one entry per group and introducer. Fields: id, group (the vicoba_group_key, or the master organisation_id), group_name,
  introducer_kind (cdo_ward, cdo_council, ngo_field_office, bank_branch, faith_office, employer, group_published), introducer_ref
  (gov:<run-id>:<office organisation_id> for a CDO; a master organisation_id otherwise; blank for group_published), introducer_name,
  basis, source_url, source_date, evidence_excerpt, introducer_outcome ('' | 'introduced <date>' | 'declined <date>'),
  officer_known_through_introducer ('' | '<date>': the introducer named an officer with the officer's agreement; no name is kept),
  status (proposed, accepted or rejected), reviewed_by (a role), reviewed_on, notes. Only accepted entries are used.
- "officer_routes": a group officer's role phone that a public body or the group publishes for the role. Fields: id, group, role,
  phone, phone_basis (must be published_for_role), source_url, source_date, evidence_excerpt, confirmed_on, status, reviewed_by,
  reviewed_on. It becomes a role-only contact (no name) with contacts.role_phone and the contact fact phone_basis. A consented
  number is never accepted here: consent arrives only as call results with no number (import_call_results.py --vicoba).

Checks on every accepted link (any failure stops the apply; nothing is written):
- cdo_ward: the reference resolves to a ward office in that government run (government_office_profiles.office_level 'ward') that
  holds the 'Ward Community Development Officer' post, in the group's council and ward; cdo_council: a council office with the
  'Head of Community Development Division' post, in the group's council. The government run is read only.
- ngo_field_office, bank_branch, faith_office, employer: a master organisation that is not a savings group.
- group_published: an accepted officer route for the group.

Call results: data/raw/contact-research/vicoba_calls_<date>.jsonl (import_call_results.py --vicoba) carry no number and no name. A
reached officer who agreed becomes the group fact phone_basis 'consent_recorded <date>'; a refusal suppresses the officer route; an
introducer's answer becomes the introducer outcome; the meeting day the committee states becomes meeting_cadence; an offered member
list is refused and recorded as list_refused.

Band properties (one current value each, on the group, rewritten on every apply; the band itself is computed by
scripts/messaging/first_touch.vicoba_case(), section 4.1). See BAND_FIELDS and docs/methodology/contact-research.md ("VICOBA").

--propose writes runtime/master/vicoba-introducer-proposals.json from the groups' intake references, a group intake's published
role phones (--intake) and an introducer intake (--introducer-intake, T1's data/raw/vicoba/introducers_intake_<date>.csv) for a
person to review and copy into the reviewed file. Preview by default (runtime/master/vicoba-introducer-links-preview.json);
--apply writes in one transaction and rolls back on any failed check; a re-run adds nothing. --check validates the reviewed file
only. --self-test runs the T3 checks on temporary copies of the given databases. Afterwards:
  python scripts/master/recompute_campus_distances.py --apply   (ward points and bands)
  python scripts/messaging/draft_master_messages.py; python scripts/master/refresh_acquisition_metadata.py
  python scripts/master/export_master_workbook_data.py; npm run build:workbook; python scripts/master/verify_master.py
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts" / "master"))
sys.path.insert(0, str(ROOT / "scripts" / "government"))
import add_vicoba_groups as V  # noqa: E402  the VICOBA contract and readers (T2)

MASTER = V.MASTER
LINKS_FILE = ROOT / "data" / "reference" / "vicoba-introducers.json"
CALLS_DIR = ROOT / "data" / "raw" / "contact-research"
CALLS_GLOB = "vicoba_calls_*.jsonl"
WORK = ROOT / "runtime" / "master"
DEFAULT_RUN = "arusha-government-2026-09"

KINDS = V.INTRODUCER_KINDS + ("employer",)
CDO_KINDS = {"cdo_ward": ("ward", "Ward Community Development Officer"),
             "cdo_council": ("council", "Head of Community Development Division")}
MASTER_KINDS = ("ngo_field_office", "bank_branch", "faith_office", "employer")
KIND_ORDER = ("employer", "cdo_ward", "ngo_field_office", "bank_branch", "cdo_council", "faith_office", "group_published")
STATUSES = ("proposed", "accepted", "rejected")
OUTCOME = re.compile(r"^(introduced|declined) \d{4}-\d{2}-\d{2}$")
DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
PHONE = re.compile(r"\+?\d[\d ()./-]{7,}\d")
EMAIL = re.compile(r"[^@\s]+@[^@\s]+\.[A-Za-z]{2,}")
ROLES = {"mwenyekiti": "Mwenyekiti (chair)", "chair": "Mwenyekiti (chair)", "chairperson": "Mwenyekiti (chair)",
         "katibu": "Katibu (secretary)", "secretary": "Katibu (secretary)",
         "mweka hazina": "Mweka Hazina (treasurer)", "treasurer": "Mweka Hazina (treasurer)"}

# The band properties: one current value each, on the group (facts, entity_type 'organisation'), rewritten on every apply.
# first_touch.vicoba_case() (T4) and export_resend_contacts.vicoba_rows() (T5) read them; 'none' means not known or not linked.
BAND_FIELDS = {
    "introducer_status": "linked (an accepted link resolves) | to_confirm (the intake names one, no accepted link) | none",
    "linked_introducer_kind": "cdo_ward | cdo_council | ngo_field_office | bank_branch | faith_office | employer | group_published | none",
    "linked_introducer_ref": "gov:<run-id>:<office id> | <master organisation_id> | none",
    "linked_introducer_name": "the introducer's name as the master or the government run holds it | none",
    "introducer_ready": "yes | no: <reason> (a CDO waits for its council's recorded introduction, GA01)",
    "introducer_outcome": "introduced <YYYY-MM-DD> | declined <YYYY-MM-DD> | none",
    "officer_route": "published_for_role | consent_recorded <YYYY-MM-DD> | known_through_introducer <YYYY-MM-DD> | refused <YYYY-MM-DD> | none",
    "officer_confirmed_on": "<YYYY-MM-DD> | none (when the officer route was last confirmed)",
    "channel": "call | letter | email",
}
NONE = "none"


def hid(prefix: str, *parts) -> str:
    return V.hid(prefix, *parts)


def norm(text) -> str:
    return " ".join(str(text or "").strip().casefold().split())


def place_key(text) -> str:
    try:
        import gov_lib  # noqa: E402  ward and council names as the government run compares them
        return gov_lib.ward_key(text)
    except Exception:  # pragma: no cover - gov_lib is always present in the repository
        return re.sub(r"[^a-z0-9 ]", "", norm(text).replace("'", ""))


def digits(value) -> str:
    return re.sub(r"\D", "", str(value or ""))


def has_number(text) -> bool:
    """A phone number (9 or more digits) in free text; ISO dates are not numbers."""
    text = re.sub(r"\d{4}-\d{2}-\d{2}", " ", str(text or ""))
    return any(len(digits(m.group())) >= 9 for m in PHONE.finditer(text))


def role_label(role) -> str:
    return ROLES.get(norm(role), str(role or "").strip())


def link_id(group: str, kind: str, ref: str) -> str:
    return "VL-" + hashlib.sha256(f"{norm(group)}|{kind}|{ref}".encode()).hexdigest()[:10]


def route_id(group: str, role: str) -> str:
    return "VO-" + hashlib.sha256(f"{norm(group)}|{norm(role_label(role))}".encode()).hexdigest()[:10]


def last(group: dict, field: str, default: str = "") -> str:
    values = group["facts"].get(field) or []
    return values[-1] if values else default


def band_properties(group: dict) -> dict:
    """The group's band properties as the last apply recorded them ('none' when the linker has not run for it)."""
    return {f: last(group, f, NONE) for f in BAND_FIELDS}


# --- The reviewed file --------------------------------------------------------------------------------------------------------------
def empty_links() -> dict:
    return {"_about": "Reviewed links between VICOBA savings groups and their introducers (scripts/master/link_vicoba_introducers.py; "
                      "docs/methodology/contact-research.md, 'VICOBA savings groups'). Only entries with status 'accepted', a reviewer "
                      "role and a review date are used. A consented number is never written here.",
            "version": "", "links": [], "officer_routes": []}


def load_links(path: Path) -> dict:
    if not path.is_file():
        return empty_links()
    data = json.loads(path.read_text(encoding="utf-8"))
    data.setdefault("links", [])
    data.setdefault("officer_routes", [])
    return data


def file_issues(data: dict) -> list[str]:
    """Structure rules of the reviewed file, before anything is resolved against a database."""
    issues, seen = [], set()
    for section in ("links", "officer_routes"):
        for i, e in enumerate(data.get(section, [])):
            where = f"{section}[{i}] {e.get('id', '')}".strip()
            if e.get("id") in seen:
                issues.append(f"{where}: duplicate id")
            seen.add(e.get("id"))
            if e.get("status") not in STATUSES:
                issues.append(f"{where}: status must be one of {list(STATUSES)}")
            if not e.get("group"):
                issues.append(f"{where}: no group (vicoba_group_key or organisation_id)")
            if e.get("status") == "accepted" and not (e.get("reviewed_by") and DATE.match(str(e.get("reviewed_on") or ""))):
                issues.append(f"{where}: accepted without reviewed_by and reviewed_on (YYYY-MM-DD)")
            if e.get("status") == "accepted" and not e.get("source_url"):
                issues.append(f"{where}: accepted without a source_url")
            for field, value in e.items():
                if section == "links" and field in ("group_name", "introducer_name", "basis", "evidence_excerpt", "notes") \
                        and isinstance(value, str) and has_number(value):
                    issues.append(f"{where}: field '{field}' holds a phone number; a link never carries one")
            if section == "links":
                if e.get("introducer_kind") not in KINDS:
                    issues.append(f"{where}: introducer_kind must be one of {list(KINDS)}")
                if e.get("introducer_outcome") and not OUTCOME.match(e["introducer_outcome"]):
                    issues.append(f"{where}: introducer_outcome must be 'introduced <date>' or 'declined <date>'")
                if e.get("officer_known_through_introducer") and not DATE.match(e["officer_known_through_introducer"]):
                    issues.append(f"{where}: officer_known_through_introducer must be a date")
            else:
                basis = str(e.get("phone_basis") or "")
                if basis != "published_for_role":
                    issues.append(f"{where}: phone_basis must be 'published_for_role'"
                                  + (" (a consented number never enters the repository)" if basis.startswith("consent") else ""))
                if len(digits(e.get("phone"))) < 9:
                    issues.append(f"{where}: no phone number")
                if not e.get("role"):
                    issues.append(f"{where}: no role (Mwenyekiti, Katibu or Mweka Hazina)")
                if e.get("name"):
                    issues.append(f"{where}: an officer route holds no name (group-level records only)")
                if e.get("confirmed_on") and not DATE.match(e["confirmed_on"]):
                    issues.append(f"{where}: confirmed_on must be a date")
    return issues


# --- The government run (read only) ---------------------------------------------------------------------------------------------
class GovRuns:
    """Read-only access to government runs by run id: offices, posts, council protocol status and ward reference points."""

    def __init__(self, overrides: dict[str, Path] | None = None):
        self.overrides = overrides or {}
        self.cons: dict[str, sqlite3.Connection | None] = {}

    def path(self, run: str) -> Path:
        return self.overrides.get(run) or ROOT / "outputs" / "runs" / run / "lead-database.sqlite"

    def con(self, run: str):
        if run not in self.cons:
            path = self.path(run)
            if path.is_file():
                c = sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True)
                c.row_factory = sqlite3.Row
                self.cons[run] = c
            else:
                self.cons[run] = None
        return self.cons[run]

    def office(self, ref: str) -> tuple[dict | None, str]:
        m = re.fullmatch(r"gov:([a-z0-9\-]+):([A-Za-z0-9_\-]+)", ref or "")
        if not m:
            return None, f"'{ref}' is not gov:<run-id>:<office organisation_id>"
        run, oid = m.groups()
        c = self.con(run)
        if c is None:
            return None, f"government run '{run}' not found at {self.path(run)} (pass --government-db {run}=<copy>)"
        r = c.execute("""SELECT o.organisation_id, o.name, o.phone, o.email, g.office_level, g.admin_unit_key, g.admin_unit_name,
                                g.council_name, g.protocol_status FROM organisations o JOIN government_office_profiles g USING(organisation_id)
                         WHERE o.organisation_id=?""", (oid,)).fetchone()
        if not r:
            return None, f"{ref} is not a government office in run '{run}'"
        office = dict(r) | {"run": run}
        office["posts"] = [p[0] for p in c.execute("SELECT office_title FROM official_posts WHERE organisation_id=?", (oid,))]
        council = c.execute("SELECT protocol_status, organisation_id FROM government_office_profiles WHERE office_level='council' AND council_name=?",
                            (office["council_name"],)).fetchone()
        office["council_protocol_status"] = council["protocol_status"] if council else "none"
        office["council_office"] = council["organisation_id"] if council else ""
        unit = c.execute("SELECT * FROM admin_units WHERE office_organisation_id=? OR unit_key=? ORDER BY office_organisation_id=? DESC",
                         (oid, office["admin_unit_key"] or "", oid)).fetchone()
        office["unit"] = dict(unit) if unit else None
        return office, ""

    def ward_unit(self, run: str, council: str, ward: str) -> dict | None:
        """The ward's admin unit (reference point) by council and ward name, with the run's reviewed ward aliases."""
        c = self.con(run)
        if c is None or not council or not ward:
            return None
        aliases = {}
        links = ROOT / "data" / "runs" / run / "links.json"
        if links.is_file():
            for cname, mapping in (json.loads(links.read_text(encoding="utf-8")).get("ward_aliases") or {}).items():
                if isinstance(mapping, dict):
                    aliases[place_key(cname)] = {place_key(k): place_key(v) for k, v in mapping.items()}
        ck, wk = place_key(council), place_key(ward)
        wk = aliases.get(ck, {}).get(wk, wk)
        for u in c.execute("SELECT * FROM admin_units WHERE level='ward'"):
            if place_key(u["council"]) == ck and place_key(u["name"]) == wk:
                return dict(u)
        return None


# --- Resolution ----------------------------------------------------------------------------------------------------------------------
def master_index(con: sqlite3.Connection) -> tuple[dict, dict, set]:
    con.row_factory = sqlite3.Row
    orgs = {r["organisation_id"]: dict(r) for r in con.execute(
        "SELECT organisation_id, name, segment, subtype, email, phone, campus, distance_km FROM organisations")}
    routes: dict[str, dict] = {}
    for r in con.execute("SELECT organisation_id, shared_email, published_role_email, named_email, organisation_phone FROM contacts"):
        d = routes.setdefault(r["organisation_id"], {"email": False, "phone": False})
        d["email"] |= any("@" in str(r[k] or "") for k in ("shared_email", "published_role_email", "named_email"))
        d["phone"] |= len(digits(r["organisation_phone"])) >= 9
    closed = {r[0] for r in con.execute("SELECT entity_id FROM review WHERE kind='Possible closure'")}
    return orgs, routes, closed


def find_group(ref: str, groups: dict, by_key: dict) -> str | None:
    if ref in groups:
        return ref
    return by_key.get(norm(ref))


def resolve_link(e: dict, group: dict, orgs: dict, gov: GovRuns) -> tuple[dict | None, list[str]]:
    kind, ref, errors = e.get("introducer_kind"), str(e.get("introducer_ref") or ""), []
    council, ward = last(group, "council"), last(group, "ward")
    if kind in CDO_KINDS:
        office, why = gov.office(ref)
        if not office:
            return None, [why]
        level, post = CDO_KINDS[kind]
        if office["office_level"] != level:
            errors.append(f"{ref} is a {office['office_level']} office, not a {level} office")
        if post not in office["posts"]:
            errors.append(f"{ref} holds no '{post}' post")
        if place_key(office["council_name"]) != place_key(council):
            errors.append(f"{ref} is in {office['council_name']}, the group in {council or 'no council'}")
        if kind == "cdo_ward":
            if not ward:
                errors.append("the group has no ward: link the council CDO (cdo_council) until the ward is known")
            elif place_key(office["admin_unit_name"]) != place_key(ward):
                unit = gov.ward_unit(office["run"], council, ward)
                if not unit or unit.get("office_organisation_id") != office["organisation_id"]:
                    errors.append(f"{ref} is the office of {office['admin_unit_name']} ward, the group is in {ward}")
        return (None if errors else {"kind": kind, "ref": ref, "name": office["name"], "office": office}), errors
    if kind in MASTER_KINDS:
        org = orgs.get(ref)
        if not V.MASTER_ID.match(ref) or not org:
            return None, [f"introducer_ref '{ref}' is not a master organisation"]
        if org["segment"] == V.SEGMENT:
            return None, [f"{ref} ({org['name']}) is a savings group, not an introducer"]
        return {"kind": kind, "ref": ref, "name": org["name"], "org": org}, []
    if kind == "group_published":
        return {"kind": kind, "ref": "", "name": group["name"] + " (route the group publishes)"}, []
    return None, [f"introducer_kind '{kind}' is not one of {list(KINDS)}"]


def readiness(link: dict | None, orgs: dict, routes: dict, closed: set, officer_route: str) -> str:
    if not link:
        return "no: no introducer linked"
    kind = link["kind"]
    if kind in CDO_KINDS:
        status = link["office"]["council_protocol_status"]
        if status == "introduced":
            return "yes"
        if status == "declined":
            return "no: the council declined the introduction"
        return "no: council introduction (GA01) not recorded"
    if kind == "group_published":
        return "yes" if officer_route == "published_for_role" else "no: no route the group publishes is on record"
    org = link["org"]
    if link["ref"] in closed:
        return "no: introducer possibly closed (review item)"
    if kind == "employer":
        return "yes"
    r = routes.get(link["ref"], {})
    if "@" in str(org.get("email") or "") or len(digits(org.get("phone"))) >= 9 or r.get("email") or r.get("phone"):
        return "yes"
    return "no: introducer has no published route"


def channel_for(link: dict | None, officer_route: str, routes: dict) -> str:
    if officer_route == "published_for_role" or officer_route.startswith("consent_recorded"):
        return "call"
    if not link or link["kind"] in CDO_KINDS or link["kind"] == "group_published":
        return "letter"
    org, r = link["org"], routes.get(link["ref"], {})
    if "@" in str(org.get("email") or "") or r.get("email"):
        return "email"
    if len(digits(org.get("phone"))) >= 9 or r.get("phone"):
        return "call"
    return "letter"


def load_calls(paths: list[Path]) -> list[dict]:
    out = []
    for path in paths:
        for n, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            if line.strip():
                rec = json.loads(line)
                rec["_file"], rec["_line"] = path, n
                out.append(rec)
    return out


def call_issues(calls: list[dict]) -> list[str]:
    """A call record never carries a number, an email or a name."""
    issues = []
    for rec in calls:
        for k, v in rec.items():
            if k.startswith("_") or k in ("organisation_id", "covers_groups", "record_id", "organisation_name", "source"):
                continue
            text = json.dumps(v, ensure_ascii=False) if not isinstance(v, str) else v
            if has_number(text) or EMAIL.search(text):
                issues.append(f"{rec['_file'].name} line {rec['_line']}: field '{k}' holds a number or an email; call records never do")
        if rec.get("contact_name") or rec.get("name") or rec.get("phone"):
            issues.append(f"{rec['_file'].name} line {rec['_line']}: a call record holds no name or number")
    return issues


# --- The plan -------------------------------------------------------------------------------------------------------------------------
def build_plan(con: sqlite3.Connection, links: dict, calls: list[dict], gov: GovRuns, as_of: date) -> dict:
    groups = V.load_vicoba_groups(con)
    by_key = {norm(g["group_key"]): oid for oid, g in groups.items() if g["group_key"]}
    orgs, routes, closed = master_index(con)
    plan = {"errors": file_issues(links) + call_issues(calls), "groups": {}, "link_records": [], "officer_contacts": [], "reviews": [],
            "call_facts": [], "unknown_groups": [], "rejected": 0, "proposed_unreviewed": 0}
    accepted_links: dict[str, list] = {}
    for e in links.get("links", []):
        if e.get("status") != "accepted":
            plan["rejected" if e.get("status") == "rejected" else "proposed_unreviewed"] += 1
            continue
        oid = find_group(e["group"], groups, by_key)
        if not oid:
            plan["errors"].append(f"link {e.get('id')}: group '{e['group']}' is not a VICOBA group in this database")
            continue
        resolved, errors = resolve_link(e, groups[oid], orgs, gov)
        if errors:
            plan["errors"] += [f"link {e.get('id')} ({groups[oid]['name']}): {x}" for x in errors]
            continue
        resolved["entry"] = e
        accepted_links.setdefault(oid, []).append(resolved)
    routes_by_group: dict[str, list] = {}
    for e in links.get("officer_routes", []):
        if e.get("status") != "accepted":
            continue
        oid = find_group(e["group"], groups, by_key)
        if not oid:
            plan["errors"].append(f"officer route {e.get('id')}: group '{e['group']}' is not a VICOBA group in this database")
            continue
        routes_by_group.setdefault(oid, []).append(e)
    # Call results: consent (no number), refusals, introducer outcomes, the committee's meeting day, a refused member list.
    events: dict[str, list] = {}
    for rec in calls:
        oid, when = rec.get("organisation_id"), rec.get("call_date") or ""
        if not DATE.match(when):
            plan["errors"].append(f"{rec['_file'].name} line {rec['_line']}: call_date must be YYYY-MM-DD")
            continue
        if rec.get("row_kind") == "group_officer":
            if oid not in groups:
                plan["errors"].append(f"{rec['_file'].name} line {rec['_line']}: {oid} is not a VICOBA group")
                continue
            events.setdefault(oid, []).append(rec)
        elif rec.get("row_kind") == "introducer":
            for g in rec.get("covers_groups") or []:
                if g in groups:
                    events.setdefault(g, []).append(rec)
        else:
            plan["errors"].append(f"{rec['_file'].name} line {rec['_line']}: row_kind must be group_officer or introducer")
    for oid, g in groups.items():
        f = g["facts"]
        # Officer route: the latest dated event wins; on the same day a refusal beats consent, consent beats a published route.
        officer = [("none", "", 0)]
        for e in routes_by_group.get(oid, []):
            officer.append(("published_for_role", e.get("confirmed_on") or e.get("reviewed_on"), 1))
        for value in f.get("phone_basis", []):
            if V.CONSENT.match(value):
                officer.append((value, value.split()[-1], 3))
        group_links = sorted(accepted_links.get(oid, []), key=lambda l: (KIND_ORDER.index(l["kind"]), l["ref"]))
        for l in group_links:
            if l["entry"].get("officer_known_through_introducer"):
                d = l["entry"]["officer_known_through_introducer"]
                officer.append((f"known_through_introducer {d}", d, 0.5))
        outcomes = [(l["entry"]["introducer_outcome"], l["entry"]["introducer_outcome"].split()[-1], l["ref"])
                    for l in group_links if l["entry"].get("introducer_outcome")]
        for rec in events.get(oid, []):
            when = rec["call_date"]
            outcome_code = rec.get("outcome")
            if rec["row_kind"] == "group_officer":
                if outcome_code == "not_interested":
                    officer.append((f"refused {when}", when, 4))
                elif outcome_code == "reached" and (rec.get("letter_or_sms_agreed") or rec.get("committee_meeting_agreed")):
                    officer.append((f"consent_recorded {when}", when, 3))
                    plan["call_facts"].append((oid, "phone_basis", f"consent_recorded {when}", rec))
                if outcome_code == "reached":
                    plan["call_facts"].append((oid, "last_confirmed", when, rec))
            if rec["row_kind"] == "introducer":
                if outcome_code == "not_interested":
                    outcomes.append((f"declined {when}", when, rec["organisation_id"]))
                elif outcome_code == "reached" and rec.get("introduces_groups"):
                    outcomes.append((f"introduced {when}", when, rec["organisation_id"]))
            if rec.get("meeting_cadence") and rec["row_kind"] == "group_officer":
                plan["call_facts"].append((oid, "meeting_cadence", rec["meeting_cadence"], rec))
            if rec.get("list_refused"):
                plan["call_facts"].append((oid, "list_refused", when, rec))
        officer.sort(key=lambda x: (x[1] or "", x[2]))
        route = officer[-1][0]
        route_kind = route.split()[0]
        confirmed = officer[-1][1] if route != "none" else NONE
        # Primary link: an employer first (the workplace group joins its deal), then a link that is ready, then by kind.
        ranked = sorted(group_links, key=lambda l: (l["kind"] != "employer",
                                                    readiness(l, orgs, routes, closed, route_kind) != "yes",
                                                    not any(o[0].startswith("introduced") and o[2] == l["ref"] for o in outcomes),
                                                    KIND_ORDER.index(l["kind"])))
        primary = ranked[0] if ranked else None
        named_in_intake = bool(f.get("introducer_kind") or f.get("introducer_ref") or f.get("sponsor_ref") or f.get("employer_ref"))
        status = "linked" if primary else "to_confirm" if named_in_intake else "none"
        mine = [o for o in outcomes if primary and o[2] == primary["ref"]]
        outcome = sorted(mine, key=lambda o: o[1])[-1][0] if mine else NONE
        props = {
            "introducer_status": status,
            "linked_introducer_kind": primary["kind"] if primary else NONE,
            "linked_introducer_ref": (primary["ref"] or NONE) if primary else NONE,
            "linked_introducer_name": primary["name"] if primary else NONE,
            "introducer_ready": readiness(primary, orgs, routes, closed, route_kind),
            "introducer_outcome": outcome,
            "officer_route": route,
            "officer_confirmed_on": confirmed or NONE,
            "channel": channel_for(primary, route_kind, routes),
        }
        # The record status follows the links: a group that now has an introducer route and is otherwise complete becomes complete.
        row = {"council": last(g, "council"), "ward": last(g, "ward"),
               "introducer_kind": (primary["kind"] if primary else last(g, "introducer_kind")) or ("group_published" if route_kind in ("published_for_role", "consent_recorded") else ""),
               "last_confirmed": max([*(f.get("last_confirmed") or []), *[c[2] for c in plan["call_facts"] if c[0] == oid and c[1] == "last_confirmed"], ""])}
        computed = V.compute_status(row, as_of)
        plan["groups"][oid] = {"name": g["name"], "women_group": g["women_group"], "record_status": g["record_status"],
                               "record_status_new": computed if computed != g["record_status"] else "", "properties": props,
                               "links": [{"id": l["entry"].get("id"), "kind": l["kind"], "ref": l["ref"]} for l in group_links],
                               "primary_record": hid("R", "vicoba-link", primary["entry"].get("id"), json.dumps(primary["entry"], sort_keys=True)) if primary else ""}
        for l in group_links:
            plan["link_records"].append({"oid": oid, "link": l, "record_id": hid("R", "vicoba-link", l["entry"].get("id"), json.dumps(l["entry"], sort_keys=True))})
        for e in routes_by_group.get(oid, []):
            plan["officer_contacts"].append({"oid": oid, "entry": e, "contact_id": hid("C", "vicoba", oid, role_label(e["role"])),
                                             "record_id": hid("R", "vicoba-officer-route", e.get("id"), json.dumps(e, sort_keys=True)),
                                             "refused": route_kind == "refused"})
        if status == "to_confirm":
            refs = "; ".join(f"{k} {', '.join(f.get(k))}" for k in ("introducer_kind", "introducer_ref", "sponsor_ref", "employer_ref") if f.get(k))
            plan["reviews"].append((hid("V", "vicoba-link-to-confirm", oid), "VICOBA introducer to confirm", oid, (f.get("introducer_ref") or [""])[0],
                                    "introducer_ref", f"named in the intake, not in the reviewed file: {refs}"[:500],
                                    "Review the introducer and add an accepted (or rejected) entry to data/reference/vicoba-introducers.json "
                                    "(link_vicoba_introducers.py --propose drafts it). Intake references are never links."))
    return plan


# --- Apply ----------------------------------------------------------------------------------------------------------------------------
def upsert_fact(con, oid, field, value, record, entity_type="organisation") -> int:
    """One current value per group and band field: a stable fact id, the value and its record replaced."""
    fid = hid("X", oid, field, "vicoba-band")
    before = con.execute("SELECT value, record_id FROM facts WHERE fact_id=?", (fid,)).fetchone()
    if before and tuple(before) == (str(value), record):
        return 0
    con.execute("INSERT INTO facts (fact_id, entity_type, entity_id, field, value, record_id) VALUES (?,?,?,?,?,?) "
                "ON CONFLICT(fact_id) DO UPDATE SET value=excluded.value, record_id=excluded.record_id",
                (fid, entity_type, oid, field, str(value), record))
    return 1


def add_fact(con, oid, field, value, record, entity_type="organisation", salt="") -> int:
    if value in (None, ""):
        return 0
    cur = con.execute("INSERT OR IGNORE INTO facts (fact_id, entity_type, entity_id, field, value, record_id) VALUES (?,?,?,?,?,?)",
                      (hid("X", oid, field, value, *([salt] if salt else [])), entity_type, oid, field, str(value), record))
    return cur.rowcount


def add_record(con, record_id, source_id, location, payload) -> int:
    cur = con.execute("INSERT OR IGNORE INTO source_records (record_id, source_id, location, payload_json) VALUES (?,?,?,?)",
                      (record_id, source_id, location, json.dumps(payload, ensure_ascii=False, sort_keys=True)))
    return cur.rowcount


def add_entity_source(con, entity_type, entity_id, record_id) -> int:
    cur = con.execute("INSERT INTO entity_sources (entity_type, entity_id, record_id) SELECT ?, ?, ? WHERE NOT EXISTS "
                      "(SELECT 1 FROM entity_sources WHERE entity_type=? AND entity_id=? AND record_id=?)",
                      (entity_type, entity_id, record_id, entity_type, entity_id, record_id))
    return cur.rowcount


def rel(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return path.as_posix()


def apply(con: sqlite3.Connection, plan: dict, links_path: Path, call_paths: list[Path]) -> dict:
    con.row_factory = None
    tables = ("organisations", "contacts", "outreach_plans", "source_records", "entity_sources", "facts", "review")
    before = {t: con.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0] for t in tables}
    org_ids = {r[0] for r in con.execute("SELECT organisation_id FROM organisations")}
    contact_ids = {r[0] for r in con.execute("SELECT contact_id FROM contacts")}
    raw = links_path.read_bytes() if links_path.is_file() else json.dumps(empty_links()).encode()
    sha = hashlib.sha256(raw).hexdigest()
    sid = hid("F", "vicoba-introducers", sha)
    written = {"band_facts_changed": 0, "facts": 0, "records": 0, "contacts": 0, "reviews": 0, "record_status_changed": 0}
    try:
        con.execute("BEGIN")
        con.execute("INSERT OR IGNORE INTO source_files (source_id, path, sha256, bytes, kind, encoding, content) VALUES (?,?,?,?,?,?,?)",
                    (sid, rel(links_path), sha, len(raw), "json; reviewed VICOBA introducer links", "utf-8", raw))
        derived = hid("R", "vicoba-band", sha)
        written["records"] += add_record(con, derived, sid, "band properties derived by link_vicoba_introducers.py",
                                         {"what": "VICOBA band properties for groups without an accepted link", "links_file": rel(links_path), "sha256": sha})
        call_sources = {}
        for path in call_paths:
            body = path.read_bytes()
            csha = hashlib.sha256(body).hexdigest()
            call_sources[path] = hid("F", "vicoba-calls", path.name, csha)
            con.execute("INSERT OR IGNORE INTO source_files (source_id, path, sha256, bytes, kind, encoding, content) VALUES (?,?,?,?,?,?,?)",
                        (call_sources[path], rel(path), csha, len(body), "jsonl; VICOBA call results (no numbers, no names)", "utf-8", body))
        for lr in plan["link_records"]:
            e = lr["link"]["entry"]
            written["records"] += add_record(con, lr["record_id"], sid, f"links: {e.get('id')}", e)
            # The intake's introducer_kind and introducer_ref facts stay as the intake gave them (proposals); the reviewed link is
            # the linked_introducer_* band properties, traceable to this entry's source record.
            written["records"] += add_entity_source(con, "organisation", lr["oid"], lr["record_id"])
        for oc in plan["officer_contacts"]:
            e, oid, cid = oc["entry"], oc["oid"], oc["contact_id"]
            written["records"] += add_record(con, oc["record_id"], sid, f"officer_routes: {e.get('id')}", e)
            existing = con.execute("SELECT role_phone, verification FROM contacts WHERE contact_id=?", (cid,)).fetchone()
            campus = con.execute("SELECT campus FROM organisations WHERE organisation_id=?", (oid,)).fetchone()[0]
            verification = (f"Role phone published for the role by {e['source_url']} ({e.get('source_date') or e.get('reviewed_on')}); "
                            f"phone_basis published_for_role; confirmed {e.get('confirmed_on') or e.get('reviewed_on')}; reviewed "
                            f"{e['reviewed_on']} ({e['reviewed_by']}). Role only: no name is recorded (group-level record).")
            if not existing:
                con.execute("""INSERT INTO contacts (contact_id, organisation_id, name, role, campus, source_url, verification, shared_email,
                               organisation_phone, contact_route, published_role_email, role_phone, named_email) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                            (cid, oid, "", role_label(e["role"]), campus, e["source_url"], verification, "", "",
                             "Role phone published for the role", "", e["phone"], ""))
                written["contacts"] += 1
            elif digits(existing[0]) != digits(e["phone"]):
                cur = con.execute('INSERT OR IGNORE INTO review (review_id, kind, entity_id, related_id, field, "values", action) VALUES (?,?,?,?,?,?,?)',
                                  (hid("V", "vicoba-officer-phone", cid, e.get("id")), V.CONFLICT_KIND, cid, oc["record_id"], "role_phone",
                                   f"{e.get('id')}: the reviewed file gives a different role phone", "Two sources disagree: confirm the current "
                                   "role phone with the source; the stored one is kept."))
                written["reviews"] += cur.rowcount
            written["records"] += add_entity_source(con, "contact", cid, oc["record_id"])
            for field, value in (("phone_basis", "published_for_role"), ("pdpa_risk", "medium"), ("role", role_label(e["role"])),
                                 ("source_url", e["source_url"])):
                written["facts"] += add_fact(con, cid, field, value, oc["record_id"], entity_type="contact")
            if oc["refused"]:
                written["facts"] += add_fact(con, cid, "pdpa_risk", "risky", oc["record_id"], entity_type="contact")
                con.execute("UPDATE contacts SET contact_route='Suppressed: refused by phone' WHERE contact_id=?", (cid,))
        for oid, field, value, rec in plan["call_facts"]:
            record = hid("R", "vicoba-call", rec["_file"].name, rec["_line"], rec.get("record_id", ""))
            payload = {k: v for k, v in rec.items() if not k.startswith("_")}
            written["records"] += add_record(con, record, call_sources[rec["_file"]], f"{rec['_file'].name} line {rec['_line']}", payload)
            written["records"] += add_entity_source(con, "organisation", oid, record)
            if field == "meeting_cadence":
                earlier = [r[0] for r in con.execute("SELECT value FROM facts WHERE entity_id=? AND field='meeting_cadence'", (oid,))]
                if earlier and norm(value) not in {norm(v) for v in earlier}:
                    cur = con.execute('INSERT OR IGNORE INTO review (review_id, kind, entity_id, related_id, field, "values", action) VALUES (?,?,?,?,?,?,?)',
                                      (hid("V", "vicoba", V.CONFLICT_KIND, oid, record, field), V.CONFLICT_KIND, oid, record, field,
                                       f"{' | '.join(earlier)} / {value} (stated by the committee)"[:500],
                                       "Two sources disagree: the committee's own statement is the current value; confirm it."))
                    written["reviews"] += cur.rowcount
            written["facts"] += add_fact(con, oid, field, value, record)
        for oid, g in plan["groups"].items():
            record = g["primary_record"] or derived
            for field, value in g["properties"].items():
                written["band_facts_changed"] += upsert_fact(con, oid, field, value, record)
            if g["record_status_new"]:
                current = [r[0] for r in con.execute("SELECT value FROM facts WHERE entity_id=? AND field='record_status' ORDER BY rowid", (oid,))]
                if not current or current[-1] != g["record_status_new"]:
                    written["record_status_changed"] += add_fact(con, oid, "record_status", g["record_status_new"], record, salt=sha)
        for review in plan["reviews"]:
            cur = con.execute('INSERT OR IGNORE INTO review (review_id, kind, entity_id, related_id, field, "values", action) VALUES (?,?,?,?,?,?,?)', review)
            written["reviews"] += cur.rowcount
        checks = transaction_checks(con, plan, org_ids, contact_ids, before)
        if not all(checks.values()):
            raise RuntimeError(f"checks failed: {[k for k, ok in checks.items() if not ok]}")
        con.execute("COMMIT")
    except Exception:
        con.execute("ROLLBACK")
        raise
    after = {t: con.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0] for t in tables}
    return {"before": before, "after": after, "written": written, "checks": checks}


def transaction_checks(con, plan, org_ids, contact_ids, before) -> dict:
    vic = V.vicoba_ids(con)
    marks = ",".join("?" * len(vic)) or "''"
    phones = con.execute(f"SELECT contact_id FROM contacts WHERE organisation_id IN ({marks}) AND (COALESCE(role_phone,'')<>'' OR "
                         "COALESCE(organisation_phone,'')<>'')", tuple(vic)).fetchall()
    basis = {r[0] for r in con.execute("SELECT entity_id FROM facts WHERE entity_type='contact' AND field='phone_basis' AND value='published_for_role'")}
    consent_with_number = con.execute("SELECT COUNT(*) FROM facts WHERE field='phone_basis' AND value LIKE 'consent%' AND value GLOB '*[0-9][0-9][0-9][0-9][0-9][0-9][0-9][0-9][0-9]*'").fetchone()[0]
    band = {r[0]: r[1] for r in con.execute("SELECT entity_id, COUNT(DISTINCT field) FROM facts WHERE field IN (%s) GROUP BY entity_id" % ",".join("?" * len(BAND_FIELDS)), tuple(BAND_FIELDS))}
    return {
        "integrity": con.execute("PRAGMA integrity_check").fetchone()[0] == "ok",
        "foreign_keys": not con.execute("PRAGMA foreign_key_check").fetchall(),
        "organisations_unchanged": {r[0] for r in con.execute("SELECT organisation_id FROM organisations")} == org_ids,
        "earlier_contacts_kept": contact_ids <= {r[0] for r in con.execute("SELECT contact_id FROM contacts")},
        "contacts_only_officer_routes": con.execute("SELECT COUNT(*) FROM contacts").fetchone()[0] - before["contacts"] <= len(plan["officer_contacts"]),
        "every_vicoba_contact_phone_has_basis": all(c in basis for (c,) in phones),
        "no_consented_number": not consent_with_number,
        "no_vicoba_phone_or_email": not con.execute(f"SELECT COUNT(*) FROM organisations WHERE organisation_id IN ({marks}) AND "
                                                    "(COALESCE(phone,'')<>'' OR COALESCE(email,'')<>'')", tuple(vic)).fetchone()[0],
        "every_group_has_band_properties": all(band.get(oid) == len(BAND_FIELDS) for oid in plan["groups"]),
        "automations_disabled": not con.execute("SELECT COUNT(*) FROM automation_configuration WHERE enabled=1").fetchone()[0],
    }


# --- Proposals ------------------------------------------------------------------------------------------------------------------------
def propose(con: sqlite3.Connection, links: dict, gov: GovRuns, intake: Path | None, introducer_intake: Path | None, out: Path) -> dict:
    """Draft entries for the reviewed file from the intake references; a person reviews them and copies the accepted ones."""
    groups = V.load_vicoba_groups(con)
    by_key = {norm(g["group_key"]): oid for oid, g in groups.items() if g["group_key"]}
    orgs, routes, closed = master_index(con)
    known = {e.get("id") for s in ("links", "officer_routes") for e in links.get(s, [])}
    by_name = {}
    for oid, o in orgs.items():
        if o["segment"] != V.SEGMENT:
            by_name.setdefault(norm(o["name"]), []).append(oid)
    out_links, out_routes, unresolved, without_groups = [], [], [], []
    today = date.today().isoformat()

    def add(oid, kind, ref, name, basis, source_url, source_date, excerpt):
        g = groups[oid]
        entry = {"id": link_id(g["group_key"], kind, ref), "group": g["group_key"], "group_name": g["name"], "introducer_kind": kind,
                 "introducer_ref": ref, "introducer_name": name, "basis": basis, "source_url": source_url, "source_date": source_date,
                 "evidence_excerpt": excerpt[:300], "introducer_outcome": "", "officer_known_through_introducer": "", "status": "proposed",
                 "reviewed_by": "", "reviewed_on": "", "notes": ""}
        if entry["id"] in known or any(x["id"] == entry["id"] for x in out_links):
            return
        resolved, errors = resolve_link(entry, g, orgs, gov)
        entry["_check"] = {"resolves": not errors, "errors": errors,
                           "ready": readiness(resolved, orgs, routes, closed, "none") if resolved else ""}
        if resolved and kind in CDO_KINDS:
            entry["introducer_name"] = resolved["name"]
        elif resolved and kind in MASTER_KINDS:
            entry["introducer_name"] = resolved["name"]
        out_links.append(entry)

    sponsor_kind = {"ngo": "ngo_field_office", "bank": "bank_branch", "church": "faith_office", "mosque": "faith_office", "employer": "employer"}
    for oid, g in groups.items():
        f = g["facts"]
        src = (f.get("sources_seen") or [g.get("source_url") or ""])[0].split(" (")[0]
        excerpt = f"Intake reference on the group record ({g['name']})."
        for kind, ref in zip(f.get("introducer_kind", []), f.get("introducer_ref", []) or [""] * len(f.get("introducer_kind", []))):
            if ref or kind == "group_published":
                add(oid, kind, ref, "", "intake introducer_ref", src, "", excerpt)
            else:
                unresolved.append({"group": g["group_key"], "group_name": g["name"], "what": f"introducer {kind} named without a reference"})
        for kind, ref in zip(f.get("sponsor_kind", []), f.get("sponsor_ref", [])):
            if ref and V.MASTER_ID.match(ref) and sponsor_kind.get(kind):
                add(oid, sponsor_kind[kind], ref, "", "intake sponsor_ref", src, "", excerpt)
        for ref in f.get("employer_ref", []):
            add(oid, "employer", ref, "", "intake employer_ref", src, "", excerpt)
        if not f.get("introducer_ref") and f.get("sponsor_name") and not f.get("sponsor_ref"):
            unresolved.append({"group": g["group_key"], "group_name": g["name"], "what": f"sponsor '{f['sponsor_name'][0]}' named without a reference"})
    if introducer_intake:
        # T1's introducer intake: rows that name a group become link proposals; rows that only name an introducer are listed for
        # add_introducers.py or a later link. Columns read (others ignored): vicoba_group_key | group_organisation_id | candidate_id,
        # introducer_kind, introducer_ref, introducer_name | organisation_name, source_url, source_date, evidence_excerpt.
        cand = {}
        for oid, g in groups.items():
            for note in [g.get("notes") or ""]:
                for m in re.finditer(r"\(([A-Za-z0-9\-]+)\)\.", note):
                    cand[m.group(1)] = oid
        for n, row in enumerate(csv.DictReader(introducer_intake.open(encoding="utf-8-sig", newline="")), 2):
            row = {k: (v or "").strip() for k, v in row.items() if k}
            oid = find_group(row.get("vicoba_group_key") or row.get("group_organisation_id") or "", groups, by_key) or cand.get(row.get("candidate_id", ""))
            name = row.get("introducer_name") or row.get("organisation_name") or ""
            ref = row.get("introducer_ref") or ""
            if not ref and name and len(by_name.get(norm(name), [])) == 1:
                ref = by_name[norm(name)][0]
            kind = row.get("introducer_kind") or ""
            if oid and kind:
                add(oid, kind, ref, name, f"introducer intake {introducer_intake.name} line {n}", row.get("source_url", ""),
                    row.get("source_date", ""), row.get("evidence_excerpt", ""))
            else:
                without_groups.append({"line": n, "introducer_name": name, "introducer_kind": kind, "master_match": ref,
                                       "what": "no group named: add the introducer with add_introducers.py if it is new, and link groups when it names them"})
    if intake:
        # Published role phones in a group intake (the adder withholds them): proposals for officer routes, for review.
        for n, row in enumerate(csv.DictReader(intake.open(encoding="utf-8-sig", newline="")), 2):
            row = {k: (v or "").strip() for k, v in row.items() if k}
            if row.get("phone_basis") != "published_for_role" or not row.get("role_phone"):
                continue
            key = row.get("vicoba_group_key") or V.group_key(row.get("council"), row.get("ward"), row.get("organisation_name"))
            oid = find_group(key, groups, by_key)
            if not oid:
                unresolved.append({"line": n, "group_name": row.get("organisation_name"), "what": "published role phone for a group not in the database"})
                continue
            entry = {"id": route_id(groups[oid]["group_key"], row.get("role") or "officer"), "group": groups[oid]["group_key"],
                     "group_name": groups[oid]["name"], "role": role_label(row.get("role") or ""), "phone": row["role_phone"],
                     "phone_basis": "published_for_role", "source_url": row.get("source_url", ""), "source_date": row.get("source_date", ""),
                     "evidence_excerpt": row.get("evidence_excerpt", "")[:300], "confirmed_on": row.get("verified_on", ""), "status": "proposed",
                     "reviewed_by": "", "reviewed_on": "", "notes": f"from {intake.name} line {n}"}
            if entry["id"] not in known:
                out_routes.append(entry)
    result = {"_about": f"Proposals for data/reference/vicoba-introducers.json, written {today}. Review each one against its source; copy "
                        "the ones you accept into the reviewed file with status 'accepted', reviewed_by (a role) and reviewed_on, and drop "
                        "the '_check' field. Nothing here is a link until then. This file is in runtime/ and never committed.",
              "links": out_links, "officer_routes": out_routes, "introducers_without_groups": without_groups, "unresolved": unresolved}
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
    return {"proposed_links": len(out_links), "resolving": sum(1 for e in out_links if e["_check"]["resolves"]),
            "proposed_officer_routes": len(out_routes), "introducers_without_groups": len(without_groups), "unresolved": len(unresolved),
            "file": str(out)}


# --- Self-test ------------------------------------------------------------------------------------------------------------------------
def self_test(args) -> int:
    """T3's checks (vicoba-integration-plan.md section 4.2) on temporary copies of the given master and government run."""
    if not args.database.is_file() or not args.government_db:
        raise SystemExit("--self-test needs --database <master copy> and --government-db [<run-id>=]<government run copy>")
    results, tmp, gov, con = {}, Path(tempfile.mkdtemp(prefix="vicoba-t3-")), None, None
    as_of = date(2026, 10, 13)
    try:
        db = tmp / "master.sqlite"
        shutil.copyfile(args.database, db)
        gov_overrides = {}
        for run, path in parse_gov(args.government_db).items():
            gov_overrides[run] = tmp / f"{run}.sqlite"
            shutil.copyfile(path, gov_overrides[run])
        gov_arg = ",".join(f"{k}={v}" for k, v in gov_overrides.items())
        run = next(iter(gov_overrides))
        gov = GovRuns(gov_overrides)
        con = sqlite3.connect(db, isolation_level=None)
        fixture = ROOT / "data" / "raw" / "vicoba" / "fixtures" / "t2_intake_fixture.csv"
        if not V.vicoba_ids(con):
            subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "master" / "add_vicoba_groups.py"), "--intake", str(fixture),
                            "--database", str(db), "--work", str(tmp / "work"), "--as-of", as_of.isoformat(), "--apply"],
                           check=True, capture_output=True, text=True, encoding="utf-8")
        groups = V.load_vicoba_groups(con)
        intake_ngo = [oid for oid, g in groups.items() if set(g["facts"].get("introducer_kind", [])) & {"ngo_field_office", "bank_branch"}
                      and any(V.MASTER_ID.match(r) for r in g["facts"].get("introducer_ref", []))]
        # 1. Proposals from the intake references; a reviewer accepts the ones that resolve and rejects the rest.
        stats = propose(con, empty_links(), gov, fixture, None, tmp / "proposals.json")
        proposals = json.loads((tmp / "proposals.json").read_text(encoding="utf-8"))
        reviewed = empty_links()
        for e in proposals["links"]:
            ok = e.pop("_check")["resolves"]
            e.update(status="accepted" if ok else "rejected", reviewed_by="self-test reviewer", reviewed_on=as_of.isoformat())
            e["source_url"] = e["source_url"] or "https://example.org/vicoba-fixture"
            reviewed["links"].append(e)
        for e in proposals["officer_routes"]:
            e.update(status="accepted", reviewed_by="self-test reviewer", reviewed_on=as_of.isoformat())
            reviewed["officer_routes"].append(e)
        links_path = tmp / "vicoba-introducers.json"
        links_path.write_text(json.dumps(reviewed, ensure_ascii=False, indent=1), encoding="utf-8")
        results["proposals"] = stats
        # 2. Entries that must be refused (each makes build_plan report an error, so nothing would be written).
        first_cdo = next(e for e in reviewed["links"] if e["introducer_kind"] == "cdo_ward" and e["status"] == "accepted")
        first = groups[find_group(first_cdo["group"], groups, {norm(x["group_key"]): o for o, x in groups.items()})]
        other = next(g for g in groups.values() if last(g, "council") == last(first, "council") and last(g, "ward")
                     and place_key(last(g, "ward")) != place_key(last(first, "ward")))
        council_office = gov.con(run).execute("SELECT organisation_id FROM government_office_profiles WHERE office_level='council'").fetchone()[0]
        village_office = gov.con(run).execute("SELECT organisation_id FROM government_office_profiles WHERE office_level='village'").fetchone()[0]
        saccos = con.execute("SELECT organisation_id FROM organisations WHERE subtype LIKE 'SACCOS%' LIMIT 1").fetchone()[0]
        bad = [
            ("cdo_ward link to a council office", "links", dict(first_cdo, id="VL-bad1", introducer_ref=f"gov:{run}:{council_office}")),
            ("cdo_ward link to a village office", "links", dict(first_cdo, id="VL-bad2", introducer_ref=f"gov:{run}:{village_office}")),
            ("cdo_ward link to another ward's office", "links", dict(first_cdo, id="VL-bad3", group=other["group_key"])),
            ("cdo_ward link to an office not in the run", "links", dict(first_cdo, id="VL-bad4", introducer_ref=f"gov:{run}:ORG_0000000000000000")),
            ("introducer that is not a master organisation", "links", dict(first_cdo, id="VL-bad5", introducer_kind="ngo_field_office", introducer_ref="O000000000000")),
            ("introducer that is a savings group", "links", dict(first_cdo, id="VL-bad6", introducer_kind="bank_branch", introducer_ref=saccos)),
            ("accepted without review", "links", dict(first_cdo, id="VL-bad7", reviewed_by="", reviewed_on="")),
            ("a link carrying a number", "links", dict(first_cdo, id="VL-bad8", notes="chair on 0754 000 009")),
            ("consented number in the reviewed file", "officer_routes",
             {"id": "VO-bad9", "group": first_cdo["group"], "role": "Mwenyekiti", "phone": "+255 000 000 009", "phone_basis": "consent_recorded 2026-10-13",
              "source_url": "call", "status": "accepted", "reviewed_by": "x", "reviewed_on": "2026-10-13"}),
            ("officer route without a basis", "officer_routes",
             {"id": "VO-bad10", "group": first_cdo["group"], "role": "Katibu", "phone": "+255 000 000 008", "phone_basis": "",
              "source_url": "https://example.org", "status": "accepted", "reviewed_by": "x", "reviewed_on": "2026-10-13"}),
        ]
        refused = {}
        for label, section, entry in bad:
            trial = json.loads(json.dumps(reviewed))
            trial[section].append(entry)
            refused[label] = bool(build_plan(con, trial, [], gov, as_of)["errors"])
        results["bad_entries_refused"] = refused
        # 3. Apply the reviewed links; a re-run adds nothing.
        plan = build_plan(con, reviewed, [], gov, as_of)
        if plan["errors"]:
            raise SystemExit(f"self-test: the reviewed fixture links do not resolve: {plan['errors']}")
        r1 = apply(con, plan, links_path, [])
        r2 = apply(con, build_plan(con, reviewed, [], gov, as_of), links_path, [])
        results["apply_checks"] = r1["checks"]
        results["rerun_adds_nothing"] = r2["before"] == r2["after"] and r2["written"]["band_facts_changed"] == 0
        groups = V.load_vicoba_groups(con)
        cdo = [(oid, band_properties(g)) for oid, g in groups.items() if band_properties(g)["linked_introducer_kind"] == "cdo_ward"]
        resolved = []
        for oid, p in cdo:
            office, _ = gov.office(p["linked_introducer_ref"])
            resolved.append(bool(office) and office["office_level"] == "ward" and "Ward Community Development Officer" in office["posts"]
                            and place_key(office["admin_unit_name"]) == place_key(last(groups[oid], "ward")))
        results["cdo_links_resolve_to_ward_offices"] = {"links": len(cdo), "all_resolve": bool(cdo) and all(resolved),
                                                        "held_until_council_introduction": all(p["introducer_ready"].startswith("no: council")
                                                                                               for _, p in cdo)}
        # NGO and bank links only through the reviewed file: unaccept the NGO entries and the intake reference alone links nothing.
        trial = json.loads(json.dumps(reviewed))
        for e in trial["links"]:
            if e["introducer_kind"] in ("ngo_field_office", "bank_branch"):
                e["status"] = "proposed"
        p_ngo = build_plan(con, trial, [], gov, as_of)
        results["ngo_links_only_through_reviewed_file"] = bool(intake_ngo) and all(
            p_ngo["groups"][oid]["properties"]["introducer_status"] == "to_confirm" for oid in intake_ngo) and all(
            g["properties"]["linked_introducer_kind"] not in ("ngo_field_office", "bank_branch") for g in p_ngo["groups"].values())
        import verify_master  # noqa: E402
        vchecks, _ = verify_master.vicoba_checks(con)
        results["phone_basis_for_every_vicoba_contact_phone"] = vchecks["vicoba_contact_phones_have_basis"] and vchecks["vicoba_no_organisation_phone"]
        # 4. The call sheet holds only office numbers and numbers published for the role; the withheld intake phone never appears.
        sheet = tmp / "sheet.xlsx"
        cs = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "contacts" / "build_call_sheet.py"), "--vicoba", "--database", str(db),
                             "--government-db", gov_arg, "--out", str(sheet)], capture_output=True, text=True, encoding="utf-8")
        from openpyxl import load_workbook
        rows = list(load_workbook(sheet, read_only=True)["Calls"].iter_rows(values_only=True)) if cs.returncode == 0 else [()]
        head = [str(h or "").split(" (")[0] for h in rows[0]]
        allowed = {digits(r[0]) for r in con.execute("SELECT c.role_phone FROM contacts c JOIN facts f ON f.entity_id=c.contact_id "
                                                     "AND f.entity_type='contact' AND f.field='phone_basis' AND f.value='published_for_role'")}
        allowed |= {digits(r[0]) for r in con.execute("SELECT phone FROM organisations WHERE segment<>?", (V.SEGMENT,))}
        allowed |= {digits(r[0]) for r in con.execute("SELECT organisation_phone FROM contacts")}
        numbers = [n for r in rows[1:] for n in str(r[head.index("number")] or "").split(", ") if n] if "number" in head else []
        bases = {str(r[head.index("number_basis")]).split(":")[0] for r in rows[1:]} if "number_basis" in head else set()
        kinds = {r[head.index("row_kind")] for r in rows[1:]} if "row_kind" in head else set()
        results["call_sheet"] = {"built": cs.returncode == 0, "rows": max(0, len(rows) - 1), "numbers": len(numbers),
                                 "has_officer_and_introducer_rows": kinds == {"group_officer", "introducer"},
                                 "only_office_or_published_for_role": bool(numbers) and all(digits(n) in allowed for n in numbers)
                                                                     and bases <= {"office", "published_for_role"},
                                 "withheld_intake_phone_absent": "000000001" not in "".join(digits(n) for n in numbers),
                                 "no_number_or_name_answer_column": not any(h in head for h in ("contact_name", "contact_work_phone", "contact_work_email"))}
        if cs.returncode:
            results["call_sheet"]["error"] = cs.stderr[-800:]
        # 5. Call results: a sheet with a number in an answer is refused; consent is recorded with no number.
        imported = {}
        wb = load_workbook(sheet) if cs.returncode == 0 else None
        ws = wb["Calls"] if wb else None
        col = {str(c.value or "").split(" (")[0]: i + 1 for i, c in enumerate(ws[1])} if ws else {}
        officer_row = next((i for i in range(2, ws.max_row + 1) if ws.cell(i, col["row_kind"]).value == "group_officer"), None) if ws else None
        calls_dir = tmp / "calls"
        imp = ROOT / "scripts" / "contacts" / "import_call_results.py"
        if officer_row:
            for k, v in (("reached", "Katibu"), ("outcome", "reached"), ("letter_or_sms_agreed", "yes"), ("committee_meeting_agreed", "yes"),
                         ("meeting_day", "Thursday"), ("meeting_time", "afternoon"), ("member_list_offered", "yes"), ("call_date", "2026-10-14"),
                         ("caller", "Silverleaf"), ("notes", "Asked to call back on +255 000 000 009 (her own mobile)")):
                ws.cell(officer_row, col[k]).value = v
            wb.save(tmp / "filled_bad.xlsx")
            ws.cell(officer_row, col["notes"]).value = "Committee meets after the weekly meeting."
            wb.save(tmp / "filled.xlsx")
            badrun = subprocess.run([sys.executable, "-B", str(imp), str(tmp / "filled_bad.xlsx"), "--vicoba", "--date", "2026-10-14",
                                     "--out-dir", str(calls_dir)], capture_output=True, text=True, encoding="utf-8")
            imported["sheet_with_a_number_refused"] = badrun.returncode != 0 and not list(calls_dir.glob("*.jsonl"))
            good = subprocess.run([sys.executable, "-B", str(imp), str(tmp / "filled.xlsx"), "--vicoba", "--date", "2026-10-14",
                                   "--out-dir", str(calls_dir)], capture_output=True, text=True, encoding="utf-8")
            files = sorted(calls_dir.glob(CALLS_GLOB))
            imported["clean_sheet_imported"] = good.returncode == 0 and len(files) == 1
            if files:
                text = files[0].read_text(encoding="utf-8")
                imported["calls_file_has_no_number"] = "000000009" not in digits(text) and not call_issues(load_calls(files))
                p3 = build_plan(con, reviewed, load_calls(files), gov, date(2026, 10, 14))
                r3 = apply(con, p3, links_path, files)
                groups = V.load_vicoba_groups(con)
                consent = [o for o, g in groups.items() if band_properties(g)["officer_route"] == "consent_recorded 2026-10-14"]
                imported["consent_recorded_without_number"] = len(consent) == 1 and \
                    "consent_recorded 2026-10-14" in groups[consent[0]]["facts"].get("phone_basis", [])
                imported["list_refused_recorded"] = bool(consent) and "2026-10-14" in groups[consent[0]]["facts"].get("list_refused", [])
                imported["apply_checks"] = all(r3["checks"].values())
                imported["consented_number_nowhere_in_database"] = not number_in_database(con, "000000009")
        else:
            imported["officer_row_on_sheet"] = False
        results["call_results"] = imported
        # 6. Ward precision: campus and band from the ward reference point.
        rc = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "master" / "recompute_campus_distances.py"), "--database", str(db),
                             "--government-db", gov_arg, "--report", str(tmp / "recompute.json"), "--apply"],
                            capture_output=True, text=True, encoding="utf-8")
        report = json.loads((tmp / "recompute.json").read_text(encoding="utf-8")) if rc.returncode == 0 else {}
        ward = report.get("ward_precision", {})
        results["ward_precision_recompute"] = {"ran": rc.returncode == 0, "groups_from_ward_point": ward.get("from ward reference point", 0),
                                               "every_ward_group_from_its_ward_point": rc.returncode == 0 and ward.get("from ward reference point", 0) > 0
                                               and not ward.get("ward not found in the government run")}
        if rc.returncode:
            results["ward_precision_recompute"]["error"] = rc.stderr[-800:]
        # 7. The contact planner still plans no web search for groups.
        results["planner_plans_no_group_search"] = planner_check(db, tmp)
        # 8. The research ledger is valid.
        lg = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "research" / "ledger.py"), "check"], capture_output=True, text=True, encoding="utf-8")
        results["ledger_check_passes"] = lg.returncode == 0 and '"problem_count": 0' in lg.stdout
    finally:
        if con is not None:
            con.close()
        for c in (gov.cons.values() if gov else []):
            if c is not None:
                c.close()
        shutil.rmtree(tmp, ignore_errors=True)
    print(json.dumps(results, indent=1, ensure_ascii=False))
    info = {"proposals", "rows", "numbers", "links", "groups", "groups_planned", "groups_from_ward_point", "error"}
    flags = []

    def walk(value):
        if isinstance(value, dict):
            for k, v in value.items():
                if k not in info:
                    walk(v)
        elif isinstance(value, bool):
            flags.append(value)
    walk(results)
    ok = bool(flags) and all(flags)
    print("self-test passed" if ok else "self-test FAILED")
    return 0 if ok else 1


def number_in_database(con: sqlite3.Connection, number_digits: str) -> bool:
    """Whether a number (its digits) appears in any text the master holds about people, facts or sources."""
    d = number_digits
    spellings = [d, f"{d[:3]} {d[3:6]} {d[6:]}", f"{d[:3]}-{d[3:6]}-{d[6:]}", f"{d[:4]} {d[4:7]} {d[7:]}"]
    queries = ["SELECT COUNT(*) FROM facts WHERE value LIKE ?", "SELECT COUNT(*) FROM source_records WHERE payload_json LIKE ?",
               "SELECT COUNT(*) FROM contacts WHERE COALESCE(role_phone,'')||COALESCE(organisation_phone,'')||COALESCE(verification,'') LIKE ?",
               "SELECT COUNT(*) FROM source_files WHERE CAST(content AS TEXT) LIKE ?", 'SELECT COUNT(*) FROM review WHERE "values" LIKE ?']
    return any(con.execute(q, (f"%{s}%",)).fetchone()[0] for q in queries for s in spellings)


def planner_check(db: Path, tmp: Path) -> dict:
    """Run plan_contact_research.py's own selection (dry run, ledger ignored) on the copy, with every VICOBA group and one ordinary
    organisation as profiles without a route: the ordinary one is planned (control), no group is."""
    sys.path.insert(0, str(ROOT / "scripts" / "contacts"))
    import contact_lib as C  # noqa: E402
    import plan_contact_research as P  # noqa: E402
    con = sqlite3.connect(f"file:{db.as_posix()}?mode=ro", uri=True)
    vic = V.vicoba_ids(con)
    control = con.execute("SELECT o.organisation_id, o.name, o.segment FROM organisations o JOIN outreach_plans p USING(organisation_id) "
                          "WHERE o.segment NOT IN ('Savings group','office:government') LIMIT 1").fetchone()
    names = {r[0]: (r[1], r[2]) for r in con.execute("SELECT organisation_id, name, segment FROM organisations")}
    con.close()
    profile = lambda oid: {"db": "master", "organisation_id": oid, "name": names[oid][0], "segment": names[oid][1],  # noqa: E731
                           "known": {"email": "", "phone": ""}, "emails": [], "phones": [], "websites": [], "people": []}
    work = ROOT / "runtime" / "contacts" / f"vicoba-self-test-{tmp.name}"  # the planner writes its prompts under the repository
    work.mkdir(parents=True, exist_ok=True)
    (work / "profiles.json").write_text(json.dumps({"profiles": [profile(o) for o in vic] + [profile(control[0])]}), encoding="utf-8")
    saved = (C.WORK, P.DBS, sys.argv)
    try:
        C.WORK, P.DBS = work, {"master": db}
        sys.argv = ["plan_contact_research.py", "--date", "2099-01-01", "--wave", "999", "--budget", "10", "--db", "master", "--dry-run",
                    "--ignore-ledger"]
        import contextlib
        import io
        with contextlib.redirect_stdout(io.StringIO()):
            P.main()
        planned = {o["organisation_id"] for path in (work / "dry-run" / "slices").glob("*.json")
                   for o in json.loads(path.read_text(encoding="utf-8"))}
    finally:
        C.WORK, P.DBS, sys.argv = saved
        shutil.rmtree(work, ignore_errors=True)
    return {"groups": len(vic), "groups_planned": len(planned & vic), "control_planned": control[0] in planned,
            "no_group_planned": not (planned & vic) and control[0] in planned}


# --- Main -----------------------------------------------------------------------------------------------------------------------------
def parse_gov(value) -> dict[str, Path]:
    """--government-db: '<path>' (the default run) or '<run-id>=<path>[,<run-id>=<path>]'."""
    out = {}
    for part in [p for p in str(value or "").split(",") if p]:
        run, _, path = part.rpartition("=") if "=" in part else (DEFAULT_RUN, "", part)
        out[run or DEFAULT_RUN] = Path(path)
    return out


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--database", type=Path, default=MASTER)
    parser.add_argument("--links", type=Path, default=LINKS_FILE, help="the reviewed links file")
    parser.add_argument("--government-db", help="a government run database copy: <path> for the default run, or <run-id>=<path>[,...] "
                                                "(default outputs/runs/<run-id>/lead-database.sqlite, read only)")
    parser.add_argument("--calls", type=Path, nargs="*", help=f"call-result files (default {CALLS_DIR.relative_to(ROOT).as_posix()}/{CALLS_GLOB})")
    parser.add_argument("--as-of", default=date.today().isoformat(), help="date for the 12-month staleness rule")
    parser.add_argument("--work", type=Path, default=WORK)
    parser.add_argument("--propose", action="store_true", help="write proposals for the reviewed file to runtime/master/")
    parser.add_argument("--intake", type=Path, help="with --propose: a group intake whose published role phones become officer-route proposals")
    parser.add_argument("--introducer-intake", type=Path, help="with --propose: T1's data/raw/vicoba/introducers_intake_<date>.csv")
    parser.add_argument("--check", action="store_true", help="validate the reviewed file against the databases; write nothing")
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    if args.self_test:
        return self_test(args)
    if not args.database.is_file():
        raise SystemExit(f"No database at {args.database}.")
    gov = GovRuns(parse_gov(args.government_db))
    links = load_links(args.links)
    con = sqlite3.connect(args.database, isolation_level=None)
    if args.propose:
        print(json.dumps(propose(con, links, gov, args.intake, args.introducer_intake, args.work / "vicoba-introducer-proposals.json"), indent=1))
        return 0
    call_paths = args.calls if args.calls is not None else sorted(CALLS_DIR.glob(CALLS_GLOB))
    plan = build_plan(con, links, load_calls(call_paths), gov, date.fromisoformat(args.as_of))
    counts = {}
    for g in plan["groups"].values():
        for field in ("introducer_status", "linked_introducer_kind", "introducer_ready", "officer_route", "channel"):
            value = g["properties"][field].split(" <")[0]
            value = re.sub(r" \d{4}-\d{2}-\d{2}$", " <date>", value)
            counts.setdefault(field, {}).setdefault(value, 0)
            counts[field][value] += 1
    summary = {"database": str(args.database), "links_file": rel(args.links), "groups": len(plan["groups"]),
               "women_groups": sum(1 for g in plan["groups"].values() if g["women_group"]),
               "accepted_links_used": len(plan["link_records"]), "proposed_not_reviewed": plan["proposed_unreviewed"], "rejected": plan["rejected"],
               "officer_routes": len(plan["officer_contacts"]), "call_records": len(load_calls(call_paths)), "call_facts": len(plan["call_facts"]),
               "record_status_changes": {oid: g["record_status_new"] for oid, g in plan["groups"].items() if g["record_status_new"]},
               "review_items": len(plan["reviews"]), "by_property": counts, "errors": plan["errors"]}
    order = V.wave_order(V.load_vicoba_groups(con))
    preview = {"summary": summary, "groups": [dict(organisation_id=oid, **plan["groups"][oid]) for oid in order if oid in plan["groups"]],
               "band_fields": BAND_FIELDS}
    args.work.mkdir(parents=True, exist_ok=True)
    out = args.work / "vicoba-introducer-links-preview.json"
    out.write_text(json.dumps(preview, indent=1, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary, indent=1, ensure_ascii=False))
    if plan["errors"]:
        print(f"{len(plan['errors'])} problem(s) in the reviewed file or the call results; nothing was written. See {out}.")
        return 1
    if args.check or not args.apply:
        print(f"Preview only: {out}")
        return 0
    result = apply(con, plan, args.links, call_paths)
    print(json.dumps(result, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
