#!/usr/bin/env python3
"""Verify a welfare run: database integrity, reconciliation, risk labels, child-data guards, offer-aligned drafts and master separation.

Writes outputs/runs/<run-id>/run-verification.json and exits 1 if any check fails. The company-master baseline is
recorded by run_pipeline.py at the start of the run (runtime/welfare/<run-id>/master-baseline.json).
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sqlite3
import sys
from datetime import datetime, timezone

from openpyxl import load_workbook

import welfare_lib as W

sys.path.insert(0, str(W.ROOT / "scripts" / "messaging"))
import offer_lib as OL  # noqa: E402  (offer-aligned draft checks)

CHILD_FIELDS = re.compile(r"child_name|pupil|student_name|learner|date_of_birth|dob\b|child_photo|resident_", re.I)
LAYOUT_SHEETS = ["Read Me", "Organisations", "Contacts", "Parent Enquiries", "Outreach Plans", "Campaigns", "Touchpoints", "Assignments",
                 "Strategies", "Evidence", "Sources", "Review Queue", "Automation"]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-id", required=True)
    args = parser.parse_args()
    cfg = W.load_config(args.run_id)
    paths = cfg["paths"]
    master_dir = (W.ROOT / "outputs" / "master").resolve()
    baseline_path = paths["work"] / "master-baseline.json"
    baseline = json.loads(baseline_path.read_text(encoding="utf-8")) if baseline_path.exists() else {}
    con = sqlite3.connect(f"file:{paths['db'].as_posix()}?mode=ro", uri=True)
    integrity = con.execute("PRAGMA integrity_check").fetchone()[0]
    fks = con.execute("PRAGMA foreign_key_check").fetchall()
    counts = {t: con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in
              ("organisations", "contacts", "enquiries", "outreach_plans", "automation_recipes", "organisation_relationships",
               "registry_ngos", "review", "source_records", "source_files")}
    columns = {t: [r[1] for r in con.execute(f'PRAGMA table_info("{t}")')] for (t,) in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    child_columns = [f"{t}.{c}" for t, cols in columns.items() for c in cols if CHILD_FIELDS.search(c)]
    payloads = [json.loads(p) for (p,) in con.execute("SELECT payload_json FROM source_records")]
    message_checks = OL.run_message_checks(con)
    org_payloads = {eid: json.loads(p) for eid, p in con.execute("SELECT es.entity_id, sr.payload_json FROM entity_sources es JOIN source_records sr "
                                                                  "ON sr.source_record_id = es.source_record_id WHERE es.entity_type = 'organisation'")}
    in_scope = {oid for oid, p in org_payloads.items() if p.get("proposed_welfare_track") != "Excluded" and p.get("segment") != "Out of scope"}
    drafted = {t for (t,) in con.execute("SELECT target_id FROM outreach_plans WHERE target_type = 'organisation'")}
    risky_routes = {str(v).lower() for p in payloads if p.get("record_type") == "contact" and p.get("pdpa_risk") == "risky"
                    for v in (p.get("named_email"), p.get("published_role_email"), p.get("shared_email"), p.get("role_phone")) if v}
    routes_used = {str(v).lower() for (v,) in con.execute("SELECT route_value FROM outreach_plans WHERE route_value != '' AND review_status = 'draft_ready'")}
    con.close()
    missing_risk = [p.get("record_key") or p.get("organisation_name") for p in payloads if p.get("pdpa_risk") not in ("low", "medium", "risky")]
    enquiry_not_risky = [p.get("record_key") for p in payloads if p.get("record_type") == "enquiry" and p.get("pdpa_risk") != "risky"]
    child_hits = [p.get("record_key") for p in payloads
                  if W.CHILD_ID_PATTERN.search(" ".join(str(p.get(k, "")) for k in ("notes", "evidence_excerpt", "request", "size_evidence")))]
    with open(paths["intake"], encoding="utf-8") as handle:
        intake_rows = list(csv.DictReader(handle))
    intake_by_type = {k: sum(1 for r in intake_rows if r["record_type"] == k) for k in ("organisation", "contact", "enquiry")}
    wb = load_workbook(paths["workbook"], read_only=True)
    sheet_rows = {name: max(0, wb[name].max_row - 4) for name in ("Organisations", "Contacts", "Parent Enquiries", "Relationships", "NGO Register", "Review Queue",
                                                                    "Outreach Plans")}
    missing_sheets = [s for s in LAYOUT_SHEETS if s not in wb.sheetnames]
    # scripts/welfare/ is canonical; the skill bundles an identical copy. Any difference means one side was edited alone.
    script_drift = []
    if W.SKILL_SCRIPTS.is_dir():
        canonical = {p.name: W.sha256_file(p) for p in W.CANONICAL_SCRIPTS.glob("*.py")}
        bundled = {p.name: W.sha256_file(p) for p in W.SKILL_SCRIPTS.glob("*.py")}
        script_drift = sorted(n for n in canonical.keys() | bundled.keys() if canonical.get(n) != bundled.get(n))
    checks = {
        "integrity": integrity == "ok",
        "foreign_keys": not fks,
        "organisations_reconcile": counts["organisations"] == intake_by_type["organisation"] == sheet_rows["Organisations"],
        "contacts_reconcile": counts["contacts"] == sheet_rows["Contacts"] and counts["contacts"] <= intake_by_type["contact"],
        "enquiries_reconcile": counts["enquiries"] == intake_by_type["enquiry"] == sheet_rows["Parent Enquiries"],
        "relationships_reconcile": counts["organisation_relationships"] == sheet_rows["Relationships"],
        "registry_reconcile": counts["registry_ngos"] == sheet_rows["NGO Register"],
        "review_reconcile": counts["review"] == sheet_rows["Review Queue"],
        "every_row_risk_labelled": not missing_risk,
        "enquiries_all_risky": not enquiry_not_risky,
        "no_child_level_columns": not child_columns,
        "no_child_identifier_text": not child_hits,
        "outreach_drafts_only_no_automation": set(message_checks["statuses"]) <= OL.DRAFT_STATUSES and counts["automation_recipes"] == 0,
        "messages_conform_to_offer_register": not message_checks["issues"] and message_checks["orphan_messages"] == 0,
        "every_in_scope_organisation_has_a_draft": in_scope <= drafted,
        "outreach_plans_reconcile": counts["outreach_plans"] == sheet_rows["Outreach Plans"],
        "no_risky_route_in_ready_drafts": not (risky_routes & routes_used),
        "layout_sheets_present": not missing_sheets,
        "raw_evidence_registered": counts["source_files"] > 1,
        "welfare_outputs_separate_from_master": master_dir not in paths["workbook"].resolve().parents and master_dir not in paths["db"].resolve().parents,
        "company_master_db_unchanged": bool(baseline) and W.sha256_file(W.MASTER) == baseline.get("master_db_sha256"),
        "company_master_workbook_unchanged": bool(baseline) and W.sha256_file(master_dir / "Silverleaf Master Database - Consolidated.xlsx") == baseline.get("master_workbook_sha256"),
        "skill_script_copies_in_sync": not script_drift,
    }
    report = {"checked_at": datetime.now(timezone.utc).isoformat(), "run_id": args.run_id, "result": "passed" if all(checks.values()) else "failed",
              "checks": checks, "db_counts": counts, "intake_by_type": intake_by_type, "sheet_rows": sheet_rows,
              "details": {"missing_risk": missing_risk[:20], "enquiry_not_risky": enquiry_not_risky, "child_columns": child_columns,
                          "child_identifier_hits": child_hits[:20], "missing_sheets": missing_sheets, "foreign_key_violations": fks[:10],
                          "master_baseline": baseline, "script_copies_out_of_sync": script_drift,
                          "message_issues": message_checks["issues"][:20], "messages_checked": message_checks["messages_checked"],
                          "undrafted_in_scope": sorted(in_scope - drafted)[:20]}}
    (paths["run_out"] / "run-verification.json").write_text(json.dumps(report, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"result": report["result"], "failed": [k for k, v in checks.items() if not v], "db_counts": counts}, indent=1))
    return 0 if report["result"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
