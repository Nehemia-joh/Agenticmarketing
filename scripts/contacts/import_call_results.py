#!/usr/bin/env python3
"""Turn a filled-in call sheet (build_call_sheet.py) into contact-research records the normal merge can take.

Reads the "Calls" sheet of the given workbook. A row is used only when the outcome is "reached" and a contact name is given. It becomes one
line of data/raw/contact-research/search_calls_<date>.jsonl in the research agents' format: the person's name and role, their work email and
phone as the person on the phone gave them, the source "tel:<number called>", and an excerpt that records who answered, the caller, the date
and whether they agreed to receive the email. A contact who did not agree is recorded with pdpa_risk "risky", so no draft uses that route.
Then run build_contact_profiles.py --date <date> and merge_master_contacts.py (preflight, then --apply), as for any research wave.

--vicoba reads a filled-in VICOBA call sheet (build_call_sheet.py --vicoba) and writes
data/raw/contact-research/vicoba_calls_<date>.jsonl: one record per answered row with the organisation ID, the groups it covers, the
role reached, the outcome, whether the committee agreed to a meeting or to a letter or SMS, whether the introducer will introduce
groups, the meeting day and time as the committee stated them, a refused member list, the date and the caller. It never holds a
number or a name: if any answer cell holds a phone number or an email, or the notes name someone (an honorific followed by a name),
the import stops and writes nothing, so a consented number can never enter the repository (it belongs in the partnerships call log).
scripts/master/link_vicoba_introducers.py --apply then records the results: consent as the group fact phone_basis
'consent_recorded <date>' (no number), a refusal suppresses the officer route, an introducer's answer becomes its outcome.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "raw" / "contact-research"
PHONE = re.compile(r"\+?\d[\d ()./-]{7,}\d")
EMAIL = re.compile(r"[^@\s]+@[^@\s]+\.[A-Za-z]{2,}")
# An honorific or title followed by a capitalised word reads as a person's name ("Mama Neema", "Bw. Juma", "Mr Peter").
NAMED = re.compile(r"\b(Mama|Baba|Bibi|Mzee|Bwana|Bw|Bi|Mhe|Mheshimiwa|Dada|Kaka|Mr|Mrs|Ms|Miss|Dr|Pastor|Mchungaji|Sheikh)\.?\s+[A-Z][a-z]+")
OUTCOMES = {"reached": "reached", "call back": "call_back", "wrong number": "wrong_number", "not interested": "not_interested", "no answer": "no_answer"}
VICOBA_ANSWERS = ("reached", "outcome", "committee_meeting_agreed", "letter_or_sms_agreed", "introduces_groups", "meeting_day", "meeting_time",
                  "member_list_offered", "call_date", "caller", "notes")


def has_number(text: str) -> bool:
    text = re.sub(r"\d{4}-\d{2}-\d{2}", " ", text)
    return any(len(re.sub(r"\D", "", m.group())) >= 9 for m in PHONE.finditer(text))


def vicoba(args) -> int:
    """VICOBA call results: no number and no name ever leaves the sheet; one problem stops the import with nothing written."""
    ws = load_workbook(args.workbook, read_only=True, data_only=True)["Calls"]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h or "") for h in next(rows)]
    col = {h.split(" (")[0]: i for i, h in enumerate(headers)}
    missing = [c for c in ("row_kind", "organisation_id", "group_ids", *VICOBA_ANSWERS) if c not in col]
    if missing:
        raise SystemExit(f"Not a VICOBA call sheet (build_call_sheet.py --vicoba): columns missing {missing}.")
    problems, out, skipped = [], [], 0
    for n, r in enumerate(rows, 2):
        get = lambda k: str(r[col[k]] if col[k] < len(r) and r[col[k]] is not None else "").strip()  # noqa: E731
        if not get("organisation_id"):
            continue
        for k in VICOBA_ANSWERS:
            v = get(k)
            if k != "call_date" and (has_number(v) or EMAIL.search(v)):
                problems.append(f"row {n}, column '{k}': a phone number or an email (write it only in the partnerships call log outside the repository)")
            if NAMED.search(v):
                problems.append(f"row {n}, column '{k}': reads as a person's name (record the role only)")
        outcome = OUTCOMES.get(get("outcome").lower())
        if not outcome:
            skipped += 1
            if get("outcome"):
                problems.append(f"row {n}: outcome '{get('outcome')}' is not one of {list(OUTCOMES)}")
            continue
        when = get("call_date")[:10] or args.date
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", when):
            problems.append(f"row {n}: call_date must be YYYY-MM-DD")
            continue
        yes = lambda k: get(k).lower().startswith("y")  # noqa: E731
        cadence = ", ".join(x for x in (get("meeting_day"), get("meeting_time")) if x)
        rec = {"db": "master", "kind": "vicoba_call", "row_kind": get("row_kind"), "organisation_id": get("organisation_id"),
               "organisation_name": get("organisation"), "covers_groups": [g for g in get("group_ids").split("; ") if g],
               "call_date": when, "caller": get("caller") or "Silverleaf", "reached_role": get("reached"), "outcome": outcome,
               "committee_meeting_agreed": yes("committee_meeting_agreed"), "letter_or_sms_agreed": yes("letter_or_sms_agreed"),
               "introduces_groups": yes("introduces_groups"),
               "meeting_cadence": f"{cadence} (stated by the committee on {when})" if cadence and get("row_kind") == "group_officer" else "",
               "list_refused": yes("member_list_offered"), "notes": get("notes"),
               "source": f"call sheet {args.workbook.name} row {n}"}
        rec["record_id"] = hashlib.sha256(json.dumps(rec, sort_keys=True).encode()).hexdigest()[:16]
        out.append(rec)
    if problems:
        print(json.dumps({"refused": True, "problems": problems, "written": 0}, indent=1, ensure_ascii=False))
        return 2
    path = (args.out_dir or RAW) / f"vicoba_calls_{args.date}.jsonl"
    path.parent.mkdir(parents=True, exist_ok=True)
    seen = {json.loads(line).get("record_id") for line in path.read_text(encoding="utf-8").splitlines() if line.strip()} if path.exists() else set()
    new = [rec for rec in out if rec["record_id"] not in seen]
    with path.open("a", encoding="utf-8") as f:
        for rec in new:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    print(json.dumps({"rows_used": len(out), "new_records": len(new), "rows_skipped": skipped, "file": str(path),
                      "next": "python scripts/master/link_vicoba_introducers.py, then --apply"}, indent=1))
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--date", default=date.today().isoformat(), help="the research date the merge will use")
    parser.add_argument("--welfare", action="store_true", help="a welfare call sheet: write welfare research contact records "
                        "(data/raw/welfare-research/research_K_calls_<date>.jsonl), add its glob to the run's run-config research_files if the date is new, then run the welfare pipeline with --rebuild-db")
    parser.add_argument("--vicoba", action="store_true", help="a VICOBA call sheet: write data/raw/contact-research/vicoba_calls_<date>.jsonl "
                        "(no numbers, no names) for link_vicoba_introducers.py")
    parser.add_argument("--out-dir", type=Path, help="with --vicoba: write the call records here instead (tests)")
    args = parser.parse_args()
    if args.vicoba:
        return vicoba(args)
    ws = load_workbook(args.workbook, read_only=True, data_only=True)["Calls"]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h or "") for h in next(rows)]
    col = {h.split(" (")[0]: i for i, h in enumerate(headers)}
    out, skipped = [], 0
    for r in rows:
        get = lambda k: str(r[col[k]] or "").strip() if k in col and col[k] < len(r) else ""  # noqa: E731
        if not get("organisation_id"):
            continue
        if not get("outcome").lower().startswith("reached") or not get("contact_name"):
            skipped += 1
            continue
        consent = get("agreed_to_receive_email").lower().startswith("y")
        number = (get("phones").split(",")[0] or "").strip()
        email = get("contact_work_email").lower()
        if email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[a-z]{2,}", email):
            email = ""
        excerpt = (f"Phone call on {get('call_date') or args.date} by {get('caller') or 'Silverleaf'}; answered by {get('answered_by_name') or 'the office'}"
                   f"{' (' + get('answered_by_role') + ')' if get('answered_by_role') else ''}; agreed to receive the email: {'yes' if consent else 'no'}.")
        out.append({"db": "master", "organisation_id": get("organisation_id"), "organisation_name": get("organisation"), "status": "found", "identity": "confirmed",
                    "searches_used": 0, "website": "", "emails": [], "phones": [], "postal_address": {}, "physical_address": {}, "socials": [],
                    "people": [{"name": get("contact_name"), "role": get("contact_role") or "staff welfare or benefits contact (by phone)", "email": email,
                                "phone": get("contact_work_phone"), "source_url": f"tel:{number}", "fetched": True, "excerpt": excerpt,
                                "pdpa_risk": "medium" if consent else "risky",
                                "pdpa_risk_reason": "given by the organisation by phone, with agreement to email" if consent else "given by phone without agreement to email"}],
                    "sources": [{"url": f"tel:{number}", "title": "Phone call", "fetched": True, "accessed_on": get("call_date") or args.date, "excerpt": excerpt,
                                 "facts": "staff welfare or benefits contact"}],
                    "notes": f"Method: phone call (call sheet). {get('notes')}".strip()})
    if args.welfare:
        out = [{"record_type": "contact", "slice": "K", "organisation_name": rec["organisation_name"], "contact_name": rec["people"][0]["name"],
                "role": rec["people"][0]["role"], "role_certainty": "confirmed", "emails": [rec["people"][0]["email"]] if rec["people"][0]["email"] else [],
                "email_type": "named" if rec["people"][0]["email"] else "", "phones": [rec["people"][0]["phone"]] if rec["people"][0]["phone"] else [],
                "phone_type": "office", "profile_url": "", "channel_attribution": rec["people"][0]["excerpt"],
                "sources": [{"url": rec["people"][0]["source_url"], "title": "Phone call", "source_date": "", "accessed_on": args.date, "evidence_basis": "phone call",
                             "fetched": True, "facts_supported": ["name", "role", "email"], "evidence_excerpt": rec["people"][0]["excerpt"]}],
                "verification_status": "verified", "pdpa_risk": rec["people"][0]["pdpa_risk"], "pdpa_risk_reason": rec["people"][0]["pdpa_risk_reason"],
                "notes": rec["notes"]} for rec in out]
        path = ROOT / "data" / "raw" / "welfare-research" / f"research_K_calls_{args.date}.jsonl"
    else:
        path = RAW / f"search_calls_{args.date}.jsonl"
    with path.open("a", encoding="utf-8") as f:
        for rec in out:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    print(json.dumps({"rows_used": len(out), "rows_skipped": skipped, "file": str(path)}, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
