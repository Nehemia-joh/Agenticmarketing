#!/usr/bin/env python3
"""Build the consolidated master review workbook from the database export. No network calls.

Reads runtime/artifacts/workbook-input.json (scripts/master/export_master_workbook_data.py) and writes
outputs/master/Silverleaf Master Database - Consolidated.xlsx from scratch, so every sheet reflects the canonical
database. The layout is unchanged: a title and subtitle, then a filterable table from row 4 with frozen headers,
fixed widths and wrapped text.
- Messages, Outreach plans and Sequences show the current offer-aligned drafts (subject, message, follow-ups,
  Kiswahili version, offer IDs). Earlier versions remain in SQLite message_versions.
- Organisations and Contacts carry the contact-research enrichment (route status, social pages, decision-makers,
  best route, pdpa_risk, research warnings).
Every sheet's row count is checked against the export; the result goes to outputs/reports/workbook-verification.json
and the script exits 1 on any mismatch. Run `npm run build:workbook`, which exports first.
"""

from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, NamedStyle, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = ROOT / "runtime" / "artifacts" / "workbook-input.json"
DEFAULT_OUTPUT = ROOT / "outputs" / "master" / "Silverleaf Master Database - Consolidated.xlsx"
DEFAULT_REPORT = ROOT / "outputs" / "reports" / "workbook-verification.json"
BLUE, SILVER, WHITE, TEXT = "002368", "818283", "FFFFFF", "172A3A"
MAX_CELL = 32000  # Excel holds at most 32,767 characters in a cell
URL_COLUMNS = {"website", "source_url", "evidence_url", "best_route"}


def value(v):
    """A cell value: numbers stay numbers; text loses characters Excel rejects and is capped at the cell limit."""
    if v is None or isinstance(v, (int, float)):
        return v
    text = ILLEGAL_CHARACTERS_RE.sub("", str(v))
    return text if len(text) <= MAX_CELL else text[:MAX_CELL] + " … (truncated; the full text is in SQLite)"


class Builder:
    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.body = NamedStyle(name="SL body", font=Font(name="Arial", size=10, color=TEXT),
                               alignment=Alignment(wrap_text=True, vertical="top"))
        self.header = NamedStyle(name="SL header", font=Font(name="Arial", size=10, bold=True, color=WHITE),
                                 fill=PatternFill("solid", fgColor=BLUE), alignment=Alignment(wrap_text=True, vertical="center"))
        self.wb.add_named_style(self.body)
        self.wb.add_named_style(self.header)
        self.rows = {}

    def sheet(self, name, subtitle, headers, rows, widths, row_height, table_name):
        """One table sheet in the workbook's house style (the layout the previous builder used)."""
        assert len(headers) == len(set(headers)) == len(widths), name
        ws = self.wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws["A1"], ws["A2"] = name, subtitle
        ws["A1"].font = Font(name="Arial", size=16, bold=True, color=BLUE)
        ws["A2"].font = Font(name="Arial", size=10, italic=True, color=SILVER)
        ws.row_dimensions[1].height = ws.row_dimensions[2].height = 25
        for col, header in enumerate(headers, start=1):
            ws.cell(row=4, column=col, value=header).style = "SL header"
        ws.row_dimensions[4].height = 38
        links = {i for i, h in enumerate(headers) if h in URL_COLUMNS}
        # Cells are written directly: indexing a whole row (ws[n]) rescans the sheet and is quadratic on large sheets.
        for number, row in enumerate(rows or [[""] * len(headers)], start=5):
            for i, v in enumerate(row):
                cell = ws.cell(row=number, column=i + 1, value=value(v))
                if cell.data_type == "f":
                    cell.data_type = "s"  # text that starts with "=" stays text; the workbook holds no formulas
                cell.style = "SL body"
                if i in links and isinstance(v, str) and re.fullmatch(r"(https?://|www\.)\S{4,2000}", v):
                    cell.hyperlink = v if v.startswith("http") else f"https://{v}"
            if row_height:
                ws.row_dimensions[number].height = row_height
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "C5" if len(headers) > 1 else "A5"
        table = Table(displayName=table_name, ref=f"A4:{get_column_letter(len(headers))}{4 + max(len(rows), 1)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        self.rows[name] = len(rows)
        return ws


def sequence_for(row: dict) -> list:
    """Timing per acquisition track: AQ00 holds, AQ01 routes with one check-in, AQ02 tests three touches. The first message
    is a request; the offer message states the offer (AQ02's follow-up 1, or the reply once an AQ01 route names the owner)."""
    offer = row.get("offer_message") or ""
    if row["acquisition_track_id"] == "AQ00":
        return ["Hold for verification; do not send", "", "No follow-up", "", "No follow-up", "", "Held", offer, "Held"]
    if row["acquisition_track_id"] == "AQ01":
        return ["After route and copy review", row["body"], "5 working days after actual delivery", row["follow_up_1"],
                "Stop; no second follow-up", "", "As the reply, once someone names the right colleague", offer, "Routing-first test"]
    return ["After identity, role and copy review", row["body"], "4 working days after actual delivery", row["follow_up_1"],
            "4 working days after follow-up 1 delivery", row["follow_up_2"], "As the reply if the recipient answers the request first", offer,
            "Direct-recipient test"]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    data = json.loads(args.input.read_text(encoding="utf-8"))
    report, counts = data["report"], data["report"]["counts"]
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    message_by_id = {m["message_id"]: m for m in data["messages"]}
    offer = report.get("offer_messages") or {}
    merge = report.get("contact_merge") or {}
    b = Builder()

    b.sheet("Start here", f"Silverleaf master intelligence: organisations, contacts, offer-aligned outreach drafts, campaign positioning and "
            f"design-only automation · generated {generated} from the canonical SQLite database",
            ["Table", "Records / use", "How to use it"], [
                ["Organisations", counts["organisations"], "Filter by segment, route_status and selection. Choose at most one current recipient per "
                 "resolved organisation. Route status, social pages, decision-makers and research flags come from the contact research."],
                ["Contacts", counts["contacts"], "Business contacts only, each with decision_maker, best route and pdpa_risk. A named person is not "
                 "evidence that they are a parent."],
                ["Enquiries", counts["enquiries"], "Historical public childcare and school enquiries. Use the one-reply review flow."],
                ["Messages", counts["messages"], "Offer-aligned drafts (offer v4) for every organisation and contact plan: subject, message, follow-ups "
                 "and the Kiswahili version where one exists. Offer terms come only from data/reference/silverleaf-offer-register.json; Finance must "
                 "confirm the 2027 terms before anything is sent. Earlier versions remain in SQLite message_versions."],
                ["Review", counts["review"], "Decisions for a person: differing values, possible overlaps and contact-research warnings (possible "
                 "closures, lost or hijacked websites, duplicates). Nothing here is applied automatically."],
                ["Verified hooks", counts["verified_hooks"], "Recipient hooks appear only where their exact claims were verified; every other hook "
                 "stays inactive."],
                ["Positioning evidence", counts["positioning_evidence"], "Approved value pillars, current official web checks, operational rules and "
                 "blocked claims with exact sources."],
                ["Campaigns", counts["campaigns"], "Campaign blueprints with audience, objective, channel, CTA, owner, approver and activation gate."],
                ["Campaign touchpoints", counts["campaign_touchpoints"], "Touchpoints and delays separated by audience state. Planned dates require "
                 "confirmation."],
                ["Lead assignments", counts["campaign_lead_assignments"], "Every organisation/contact message and every historical enquiry mapped "
                 "to a campaign decision."],
                ["Acquisition tracks", counts["acquisition_tracks"], f"AQ00 holds unresolved leads ({counts['acquisition_hold']}), AQ01 asks for "
                 f"routing ({counts['acquisition_routing']}) and AQ02 tests direct-recipient outreach ({counts['acquisition_direct']})."],
                ["Value modules", counts["value_proposition_modules"], "Approved Silverleaf positioning is selected only where it helps the recipient "
                 "assess the offer."],
                ["Lead intake rules", counts["lead_intake_rules"], "Apply after each list expansion to preserve sources, deduplicate, classify and "
                 "assign copy rules."],
                ["Parent enquiry drafts", f"{counts['parent_enquiry_outreach_drafts']} drafts / {counts['parent_enquiry_excluded_from_outreach']} "
                 "excluded", f"All {counts['parent_enquiry_drafts']} records remain visible. Only daycare/nursery and primary-school enquiries have "
                 "one-reply copy; no automatic follow-up."],
                ["Audience segments", counts["outreach_segments"], "Use segment for offer and cadence, then personalise the recipient reason and ask."],
                ["Automation recipes", counts["automation_steps"], "Design-only triggers, conditions, delays, database updates and stop rules."],
                ["Knowledge sources", counts["knowledge_documents"], "Marketing documents, outreach guidance, official web verification and "
                 "structured references."],
                ["Facts", len(data["facts"]), "Distinct source assertions, each with every source record ID."],
                ["Source files", counts["source_files"], "Exact archived file bytes with SHA-256 hashes."],
                ["Source rows", counts["source_records"], "Parsed pages, paragraphs, rows, sections and relationships with source IDs."],
                ["Strategies", counts["strategies"], "Consolidated strategy guidance, including campaign and cadence decisions."],
                ["Related workbooks", "3", "outputs/contacts/Silverleaf Contact Profiles - <date>.xlsx (contact profiles and research flags across "
                 "all three databases); outputs/messages/Silverleaf Offer-Aligned Messages - <date>.xlsx (every draft in all three databases); "
                 "welfare and government runs in outputs/runs/."],
                ["Sending", "Disabled", "No campaign, broadcast, recurring job or delivery integration is active."],
            ], [34, 20, 110], 58, "TStartHere")

    b.sheet("Organisations", f"{counts['organisations']} organisations in the canonical database. Route status: 'direct route' means a published "
            "email or phone for the organisation or one of its contacts. Check source records and verification before outreach.",
            ["organisation_id", "name", "segment", "priority", "campus", "locality", "distance_km", "geocode_precision", "route_status", "phone",
             "email", "website", "address", "social_pages", "named_contacts", "decision_makers", "contact_research", "research_flags", "headcount",
             "size_evidence", "education_angle", "desk_tier", "desk_score", "verification", "strategy_id", "owner", "outreach_status", "next_action",
             "next_action_date", "message_ids", "source_record_ids", "source_url"],
            [[o["organisation_id"], o["name"], o["segment"], o["priority"], o["campus"], o["locality"], o["distance_km"], o["geocode_precision"],
              o["route_status"], o["phone"], o["email"], o["website"], o["address"], o["social_pages"], o["named_contacts"], o["decision_makers"],
              o["contact_research"], o["research_flags"], o["headcount"], o["size_evidence"], o["education_angle"], o["desk_tier"], o["desk_score"],
              o["verification"], o["strategy_id"], o["owner"], o["outreach_status"], o["next_action"], o["next_action_date"], o["message_ids"],
              o["source_record_ids"], o["source_url"]] for o in data["organisations"]],
            [20, 38, 24, 14, 24, 32, 14, 22, 16, 24, 34, 44, 48, 60, 14, 70, 30, 55, 14, 50, 70, 16, 14, 54, 28, 18, 25, 45, 20, 35, 80, 60],
            82, "TOrganisations")

    b.sheet("Contacts", f"{counts['contacts']} published business contacts. Named people are professional routes only; they are not evidence "
            "of parent status. The best route is the person's own published route, else the organisation's; a personal-domain address "
            "linked to a person is never used.",
            ["contact_id", "name", "organisation_id", "organisation_name", "role", "decision_maker", "best_route_type", "best_route", "pdpa_risk",
             "campus", "named_email", "published_role_email", "role_phone", "shared_email", "organisation_phone", "contact_route", "verification",
             "message_ids", "source_url", "source_record_ids"],
            [[c["contact_id"], c["name"], c["organisation_id"], c["organisation_name"], c["role"], c["decision_maker"], c["best_route_type"],
              c["best_route"], c["pdpa_risk"], c["campus"], c["named_email"], c["published_role_email"], c["role_phone"], c["shared_email"],
              c["organisation_phone"], c["contact_route"], c["verification"], c["message_ids"], c["source_url"], c["source_record_ids"]]
             for c in data["contacts"]],
            [20, 30, 20, 44, 48, 14, 24, 40, 12, 24, 34, 34, 28, 34, 28, 42, 66, 38, 70, 88], 96, "TContacts")

    b.sheet("Enquiries", "Public enquiries remain distinct from business contacts. Most are historical and unqualified.",
            ["enquiry_id", "name", "type", "enquiry_date", "date_qualification", "locality", "request", "campus_fit", "phone", "email",
             "contact_attribution", "platform", "current_relevance", "strategy_id", "source_url", "source_record_id"],
            [[e["enquiry_id"], e["name"], e["type"], e["enquiry_date"], e["date_qualification"], e["locality"], e["request"], e["campus_fit"],
              e["phone"], e["email"], e["contact_attribution"], e["platform"], e["current_relevance"], e["strategy_id"], e["source_url"],
              e["source_record_id"]] for e in data["enquiries"]],
            [18, 22, 22, 16, 26, 30, 70, 60, 20, 30, 60, 28, 70, 18, 60, 18], 90, "TEnquiries")

    b.sheet("Messages", "Request-first drafts for every organisation and contact plan, signed by Mariam Haji. The message asks for a short meeting "
            "and states no offer terms; the offer message states the offer register's terms (AQ02 sends it as follow-up 1; on AQ01 it is the reply "
            "once someone names the right colleague). Nothing is sent: Finance must confirm the terms apply to 2027 before the offer message goes "
            "out, and each row's route, recipient and copy need review first. Hooks appear only where verified.",
            ["message_id", "target_name", "organisation_name", "target_type", "segment", "acquisition_track_id", "review_status", "contact_channel",
             "subject", "message", "follow_up_1", "follow_up_2", "offer_message", "kiswahili_version", "offer_ids", "offer_version", "hook_status",
             "hook", "value_module_ids", "campaign_copy_status", "channel_attribution", "conditions", "missing_information", "strategy_id",
             "source_record_id"],
            [[p["message_id"], p["target_name"], p["organisation_name"], p["target_type"], p["segment"], p["acquisition_track_id"],
              p["review_status"], p["contact_channel"], p["subject"], p["body"], p["follow_up_1"], p["follow_up_2"], p.get("offer_message"),
              p["offer_message_sw"], p["offer_ids"], p["offer_version"], p["hook_status"], p["hook"], p["value_module_ids"],
              p["campaign_copy_status"], p["channel_attribution"], message_by_id.get(p["message_id"], {}).get("conditions"),
              p["missing_information"], message_by_id.get(p["message_id"], {}).get("strategy_id"),
              message_by_id.get(p["message_id"], {}).get("source_record_id")]
             for p in data["outreach_plans"]],
            [22, 28, 42, 16, 24, 20, 18, 34, 48, 90, 105, 70, 105, 90, 22, 26, 30, 60, 42, 48, 45, 95, 80, 24, 24], 175, "TMessages")

    strategy_rows = []
    for s in data["strategies"]:
        parts = re.findall(r"[\s\S]{1,1200}", s["content"] or "") or [""]
        for index, content in enumerate(parts, start=1):
            strategy_rows.append([s["strategy_id"], s["title"], s["segment"], f"{index}/{len(parts)}", content, s["status"]])
    b.sheet("Strategies", "Consolidated imported and current guidance. Current campaign entries identify their positioning evidence IDs in the "
            "status field.", ["strategy_id", "title", "segment", "part", "content", "status"], strategy_rows, [30, 52, 34, 12, 110, 65], 190,
            "TStrategies")

    b.sheet("Campuses", "Approximate source coordinates. They are not surveyed campus pins.",
            ["campus_id", "name", "region", "levels", "description", "latitude", "longitude", "geocode_basis"],
            [[c["campus_id"], c["name"], c["region"], c["levels"], c["description"], c["latitude"], c["longitude"], c["geocode_basis"]]
             for c in data["campuses"]], [18, 22, 16, 34, 60, 14, 14, 60], 45, "TCampuses")

    b.sheet("Review", "Unresolved identities, differing field values and research warnings (possible closures, lost or hijacked websites, "
            "duplicates) for a person to decide. No automatic fuzzy merges; nothing here is applied automatically.",
            ["review_id", "kind", "entity_name", "entity_id", "related_id", "field", "values", "action"],
            [[r["review_id"], r["kind"], r["entity_name"], r["entity_id"], r["related_id"], r["field"], r["values"], r["action"]]
             for r in data["review"]], [18, 36, 40, 20, 20, 18, 80, 90], 60, "TReview")

    b.sheet("Facts", "Distinct source assertions. Repeated copies are combined by entity, field and value; all source record IDs remain.",
            ["entity_type", "entity_name", "entity_id", "field", "value", "source_record_ids"],
            [[f["entity_type"], f["entity_name"], f["entity_id"], f["field"], f["value"], f["source_record_ids"]] for f in data["facts"]],
            [16, 40, 20, 28, 70, 60], None, "TFacts")

    b.sheet("Source rows", f"{counts['source_records']} parsed records. Full payloads and exact archived bytes are stored in SQLite.",
            ["record_id", "source_id", "source_file", "location"],
            [[r["record_id"], r["source_id"], r["source_file"], r["location"]] for r in data["source_rows"]], [24, 24, 90, 50], None,
            "TSourceRows")

    b.sheet("Source files", f"{counts['source_files']} source files archived with SHA-256 hashes.",
            ["source_id", "path", "kind", "bytes", "sha256"],
            [[r["source_id"], r["path"], r["kind"], int(r["bytes"] or 0), r["sha256"]] for r in data["source_files"]], [24, 95, 14, 18, 72], 40,
            "TSourceFiles")

    b.sheet("Outreach plans", "One row per organisation/contact draft. The acquisition track controls cadence; segment and value modules shape "
            "the offer; offer IDs name the register terms the copy states; exact evidence controls hooks.",
            ["message_id", "target_name", "organisation_name", "segment", "acquisition_track_id", "value_module_ids", "acquisition_version",
             "strategy_scope", "recipient_role", "persona", "selection", "review_status", "contact_channel", "channel_attribution", "hook_status",
             "hook", "hook_evidence", "offer_version", "offer_ids", "offer_evidence", "subject", "message", "follow_up_1", "follow_up_2",
             "offer_message", "kiswahili_version", "campaign_copy_status", "relevance_reason", "proposed_offer", "cta_type", "flow_id",
             "evidence_url", "evidence_date", "verified_on", "evidence_basis", "evidence_record_ids", "missing_information"],
            [[p["message_id"], p["target_name"], p["organisation_name"], p["segment"], p["acquisition_track_id"], p["value_module_ids"],
              p["acquisition_version"], p["strategy_scope"], p["recipient_role"], p["persona"], p["selection"], p["review_status"],
              p["contact_channel"], p["channel_attribution"], p["hook_status"], p["hook"], p["hook_evidence"], p["offer_version"], p["offer_ids"],
              p["offer_evidence"], p["subject"], p["body"], p["follow_up_1"], p["follow_up_2"], p.get("offer_message"), p["offer_message_sw"],
              p["campaign_copy_status"], p["relevance_reason"], p["proposed_offer"], p["cta_type"], p["flow_id"], p["evidence_url"],
              p["evidence_date"], p["verified_on"], p["evidence_basis"], p["evidence_record_ids"], p["missing_information"]]
             for p in data["outreach_plans"]],
            [22, 28, 42, 24, 20, 42, 40, 85, 28, 18, 23, 22, 34, 48, 30, 80, 75, 26, 22, 60, 48, 90, 105, 70, 105, 90, 48, 80, 75, 24, 12, 55,
             18, 18, 50, 90, 100], 175, "TOutreachPlans")

    b.sheet("Sequences", "Acquisition-track sequences are independent of the internal marketing calendar. AQ00 holds, AQ01 routes with one "
            "check-in, and AQ02 tests three touches. The initial message is a request with no offer terms; the offer message states the offer.",
            ["message_id", "recipient", "organisation", "segment", "acquisition_track_id", "selection", "initial_subject", "initial_timing",
             "initial_message", "delay_to_follow_up_1", "follow_up_1", "delay_to_follow_up_2", "follow_up_2", "offer_message_timing",
             "offer_message", "track_status", "copy_status"],
            [[p["message_id"], p["target_name"], p["organisation_name"], p["segment"], p["acquisition_track_id"], p["selection"], p["subject"],
              *sequence_for(p), p["campaign_copy_status"]] for p in data["outreach_plans"]],
            [22, 28, 42, 24, 20, 22, 48, 38, 90, 38, 105, 38, 72, 38, 105, 30, 50], 175, "TSequences")

    b.sheet("Segments", "Use the segment to select the offer. Acquisition track selects cadence; recipient evidence and role select the reason "
            "and CTA.", ["segment", "audience", "offer", "qualification", "flow_id"],
            [[r["segment"], r["audience"], r["offer"], r["qualification"], r["flow_id"]] for r in data["outreach_segments"]],
            [30, 52, 65, 80, 14], 88, "TSegments")

    flow_rows = [[r["flow_id"], int(r["step"] or 0), r["trigger"], r["condition"], r["action"], r["delay"], r["database_update"],
                  r["example_copy"], r["principle"], r["mode"]] for r in data["automation_recipes"]]
    flow_headers = ["flow_id", "step", "trigger", "condition", "action", "delay", "database_update", "example_copy", "principle", "mode"]
    b.sheet("Flows", "Audience-state flows include appendable new-lead intake, routing, direct outreach, parent enquiries, events and "
            "broadcasts.", flow_headers, flow_rows, [14, 10, 42, 78, 70, 45, 62, 95, 58, 34], 135, "TFlows")

    b.sheet("Knowledge sources", "Marketing documents, operational references, outreach guidance and official web verification. Planned sources "
            "remain clearly marked.", ["document_id", "title", "classification", "record_count", "relative_path", "status"],
            [[r["document_id"], r["title"], r["classification"], int(r["record_count"] or 0), r["relative_path"], r["status"]]
             for r in data["knowledge_documents"]], [24, 52, 50, 18, 92, 88], 80, "TKnowledgeSources")

    b.sheet("Automation recipes", "Design-only implementation recipes. New-contact cadence is independent of the internal marketing calendar; "
            "sending remains disabled.", flow_headers, flow_rows, [14, 10, 42, 78, 70, 45, 62, 95, 58, 34], 135, "TAutomationRecipes")

    outcomes = offer.get("outcomes") or {}
    new_plans = sum(v for k, v in outcomes.items() if k.startswith("new "))
    audit = [
        ["Marketing documents reviewed", report["marketing_documents_reviewed"], "Complete",
         "9 September 2026: all files under references/Marketing Documents were extracted and reviewed."],
        ["New marketing originals archived", report["marketing_originals_added_to_master"], "Complete",
         "The calendar was already present; the other 11 originals were added."],
        ["Missing core records", report["missing_records"], "None", "Organisation, contact, enquiry, message and outreach-plan counts were preserved."],
        ["Verified recipient hooks", counts["verified_hooks"], "Preserved", "No unverified recipient hook is active."],
        ["New-contact acquisition scope", report["acquisition_version"], "Applied",
         "Marketing documents supply positioning and value propositions; acquisition tracks control cadence."],
        ["Campaign blueprints", counts["campaigns"], "Complete", "Every campaign includes activation gates and source basis."],
        ["Published contact enrichment", "11 contacts", "Complete", "10 September 2026: current public professional routes added for 9 OSM "
         "organisations and 2 TCDC SACCOS targets; the two SACCOS routes are flagged for officer confirmation."],
    ]
    if offer:
        audit.append(["Offer-aligned messages (offer v4)", f"{outcomes.get('rewritten', 0)} rewritten; {new_plans} new plans", "Complete",
                      f"{offer.get('run_on', '')}: every draft states only offer-register terms ({offer.get('conformance_issues', 0)} conformance "
                      "issues). Finance must confirm the 2027 terms before sending. Earlier versions remain in message_versions."])
    if merge:
        applications = merge.get("applications") or [merge]
        total = lambda key: sum(a.get(key, 0) for a in applications)  # noqa: E731
        audit.append(["Contact research merge", f"{total('contacts_inserted')} contacts; {total('fields_filled')} fields filled", "Complete",
                      f"{len(applications)} application(s) through the update skill's validator, preflight and one transaction; only empty "
                      f"fields were filled. {report.get('contact_research_reviews', 0)} contact-research review items are in Review; "
                      f"{total('plans_unheld')} held drafts were released because a route was found; {report.get('plans_held_for_closure', 0)} "
                      "drafts are held for a possible closure."])
    audit += [
        ["Acquisition track assignment", f"{counts['acquisition_hold']} hold / {counts['acquisition_routing']} routing / "
         f"{counts['acquisition_direct']} direct", "Complete", "Every current outreach-plan row is assigned; the refresh is rerunnable."],
        ["Automation steps", counts["automation_steps"], "Design only", "No sending integration or recurring schedule is active."],
        ["SQLite integrity", report.get("integrity_check", ""), "Complete" if report.get("integrity_check") == "ok" else "Check",
         f"Foreign-key violations: {report.get('foreign_key_violations', '')}."],
        ["Planned dates and mutable claims", "Approval required", "Gated",
         "Confirm events, fees, discounts, referral terms, capacity and visit slots before release."],
    ]
    b.sheet("Consolidation audit", f"Database, campaign and enrichment checks, most recently {report['checked_on']}.",
            ["check", "value", "status", "notes"], audit, [48, 34, 22, 110], 72, "TConsolidationAudit")

    b.sheet("Positioning evidence", "Use allowed claims only under their stated conditions. Planned dates and mutable statistics are not evergreen "
            "campaign facts.", ["evidence_id", "category", "claim_or_rule", "campaign_use", "source_path", "source_locator", "evidence_class",
                                "verified_on", "use_status", "conditions"],
            [[r["evidence_id"], r["category"], r["claim_or_rule"], r["campaign_use"], r["source_path"], r["source_locator"], r["evidence_class"],
              r["verified_on"], r["use_status"], r["conditions"]] for r in data["positioning_evidence"]],
            [16, 30, 90, 72, 92, 62, 45, 18, 34, 85], 110, "TPositioningEvidence")

    b.sheet("Campaigns", "Campaign blueprints. C01 and C02 use independent acquisition tracks; event and broadcast campaigns retain their own "
            "approval gates.", ["campaign_id", "name", "audience", "segment", "objective", "mode", "channels", "flow_ids", "owner", "approver",
                                "primary_cta", "language", "activation_gate", "status", "source_basis"],
            [[r["campaign_id"], r["name"], r["audience"], r["segment"], r["objective"], r["mode"], r["channels"], r["flow_ids"], r["owner"],
              r["approver"], r["primary_cta"], r["language"], r["activation_gate"], r["status"], r["source_basis"]] for r in data["campaigns"]],
            [14, 40, 75, 42, 75, 32, 58, 22, 50, 54, 65, 38, 95, 50, 42], 125, "TCampaigns")

    b.sheet("Campaign touchpoints", "Touchpoints and delays for campaign-facing flows. Response branches remain in Flows and Automation recipes.",
            ["campaign_id", "touchpoint_order", "flow_id", "flow_step", "timing", "trigger", "purpose", "channel", "audience_state", "draft_copy",
             "evidence_ids", "status"],
            [[r["campaign_id"], int(r["touchpoint_order"] or 0), r["flow_id"], int(r["flow_step"] or 0), r["timing"], r["trigger"], r["purpose"],
              r["channel"], r["audience_state"], r["draft_copy"], r["evidence_ids"], r["status"]] for r in data["campaign_touchpoints"]],
            [14, 18, 12, 12, 48, 55, 68, 58, 82, 95, 35, 35], 125, "TCampaignTouchpoints")

    b.sheet("Lead assignments", "Every organisation/contact message and historical enquiry is assigned to a campaign state. New contacts also "
            "carry acquisition-track and value-module assignments.",
            ["assignment_id", "campaign_id", "target_type", "target_id", "message_id", "acquisition_track_id", "value_module_ids",
             "strategy_scope", "selection", "eligibility_status", "reason", "next_action"],
            [[r["assignment_id"], r["campaign_id"], r["target_type"], r["target_id"], r["message_id"], r["acquisition_track_id"],
              r["value_module_ids"], r["strategy_scope"], r["selection"], r["eligibility_status"], r["reason"], r["next_action"]]
             for r in data["campaign_lead_assignments"]], [24, 14, 16, 24, 24, 20, 42, 85, 24, 28, 90, 90], 105, "TLeadAssignments")

    b.sheet("Parent enquiry drafts", "Historical public-enquiry drafts for human review. Use the original public route only where it still "
            "permits a relevant response; no automatic follow-up.",
            ["enquiry_id", "campaign_id", "subject", "body", "follow_up", "language", "route_status", "evidence_ids", "status"],
            [[r["enquiry_id"], r["campaign_id"], r["subject"], r["body"], r["follow_up"], r["language"], r["route_status"], r["evidence_ids"],
              r["status"]] for r in data["parent_enquiry_drafts"]], [24, 14, 46, 115, 30, 36, 75, 36, 58], 175, "TParentEnquiryDrafts")

    b.sheet("Acquisition tracks", "New organisation and business-contact leads use evidence-selected tracks. These tests do not inherit the "
            "internal marketing calendar.", ["track_id", "name", "use_when", "initial_touch", "follow_up_policy", "desired_outcome",
                                             "source_basis", "status"],
            [[r["track_id"], r["name"], r["use_when"], r["initial_touch"], r["follow_up_policy"], r["desired_outcome"], r["source_basis"],
              r["status"]] for r in data["acquisition_tracks"]], [14, 26, 72, 55, 78, 62, 48, 32], 92, "TAcquisitionTracks")

    b.sheet("Value modules", "Select only the propositions that help the recipient assess the offer. A Silverleaf claim is never evidence about "
            "the recipient.", ["module_id", "audience", "name", "recipient_value", "silverleaf_contribution", "recipient_contribution",
                               "use_when", "do_not_imply", "evidence_ids", "status"],
            [[r["module_id"], r["audience"], r["name"], r["recipient_value"], r["silverleaf_contribution"], r["recipient_contribution"],
              r["use_when"], r["do_not_imply"], r["evidence_ids"], r["status"]] for r in data["value_proposition_modules"]],
            [14, 42, 34, 68, 72, 62, 78, 78, 38, 38], 110, "TValueModules")

    b.sheet("Lead intake rules", "Run after each lead-list expansion. Preserve source records, deduplicate carefully, classify the recipient and "
            "assign acquisition metadata before drafting.",
            ["rule_id", "stage_order", "applies_to", "rule", "required_fields", "decision", "next_step", "source_basis"],
            [[r["rule_id"], int(r["stage_order"] or 0), r["applies_to"], r["rule"], r["required_fields"], r["decision"], r["next_step"],
              r["source_basis"]] for r in data["lead_intake_rules"]], [14, 14, 34, 82, 68, 82, 52, 40], 98, "TLeadIntakeRules")

    # Reconcile every sheet with the export before saving.
    plans = counts["outreach_plans"]
    expected = {"Organisations": counts["organisations"], "Contacts": counts["contacts"], "Enquiries": counts["enquiries"], "Messages": plans,
                "Campuses": counts["campuses"], "Review": counts["review"], "Facts": len(data["facts"]), "Source rows": counts["source_records"],
                "Source files": counts["source_files"], "Outreach plans": plans, "Sequences": plans, "Segments": counts["outreach_segments"],
                "Flows": counts["automation_steps"], "Knowledge sources": counts["knowledge_documents"],
                "Automation recipes": counts["automation_steps"], "Positioning evidence": counts["positioning_evidence"],
                "Campaigns": counts["campaigns"], "Campaign touchpoints": counts["campaign_touchpoints"],
                "Lead assignments": counts["campaign_lead_assignments"], "Parent enquiry drafts": counts["parent_enquiry_drafts"],
                "Acquisition tracks": counts["acquisition_tracks"], "Value modules": counts["value_proposition_modules"],
                "Lead intake rules": counts["lead_intake_rules"]}
    mismatches = {name: {"expected": n, "sheet": b.rows.get(name)} for name, n in expected.items() if b.rows.get(name) != n}
    if b.rows.get("Strategies", 0) < counts["strategies"]:
        mismatches["Strategies"] = {"expected_at_least": counts["strategies"], "sheet": b.rows.get("Strategies")}
    without_copy = [p["message_id"] for p in data["outreach_plans"] if not (p["subject"] and p["body"])]
    not_offer_aligned = counts["outreach_plans"] - counts.get("offer_aligned_plans", counts["outreach_plans"])
    args.output.parent.mkdir(parents=True, exist_ok=True)
    b.wb.save(args.output)
    output = args.output.resolve()
    result = {"output": output.relative_to(ROOT).as_posix() if output.is_relative_to(ROOT) else str(output), "generated_at": generated,
              "builder": "scripts/master/build_master_workbook.py",
              "sheets": b.rows, "expected": expected, "mismatches": mismatches, "messages_without_copy": without_copy[:20],
              "plans_not_offer_aligned": not_offer_aligned, "formula_errors": 0, "sending_enabled": False,
              "result": "passed" if not mismatches and not without_copy and not not_offer_aligned else "failed"}
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: result[k] for k in ("output", "result", "mismatches", "plans_not_offer_aligned")}, indent=2))
    return 0 if result["result"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
