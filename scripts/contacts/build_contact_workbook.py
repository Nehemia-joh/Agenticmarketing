#!/usr/bin/env python3
"""Build the contact-profiles review workbook across the three databases. No network calls; databases are read-only.

Run after the merges (merge_master_contacts.py --apply, and the welfare and government pipelines after
export_run_contact_research.py). Reads the company master, the welfare run and the government run, plus
runtime/contacts/profiles.json and profiles-summary.json for what the research tried (methods, social pages, sources,
blocked sites). Writes outputs/contacts/Silverleaf Contact Profiles - <date>.xlsx and contact-profiles-summary.json.

Sheets: Summary; Organisation Profiles (one row per organisation: routes, named contacts, decision-makers, research);
Contact Leads (one row per contact: role, decision-maker, routes, source, verification, pdpa_risk); Gaps
(organisations still without a direct route, with what was tried); Flags (warnings from the research for a person to check).
"""
from __future__ import annotations

import argparse
import json
import sqlite3
from collections import Counter, defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import contact_lib as C

W = C.W
DBS = {"master": C.ROOT / "outputs" / "master" / "Silverleaf Master Database.sqlite",
       "welfare": C.ROOT / "outputs" / "runs" / "arusha-welfare-2026-09" / "lead-database.sqlite",
       "government": C.ROOT / "outputs" / "runs" / "arusha-government-2026-09" / "lead-database.sqlite"}
BLUE, WHITE = "002368", "FFFFFF"
WRAP = {"routes", "named_contacts", "decision_makers", "sources", "research", "gap", "role", "source", "verification", "pdpa_risk_reason", "socials", "flag",
        "measure", "note", "address"}


def ro(path):
    con = sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    return con


def decision_maker(role: str) -> bool:
    return bool(C.DECISION.search(C.NOT_A_HEAD.sub("", role or "")))


def contact_risk(name: str, emails: list[str]) -> tuple[str, str]:
    if not name:
        return "low", "Office or role desk; no person named."
    if any(W.PERSONAL_EMAIL.search(e) for e in emails):
        return "risky", "Named person linked to a personal-domain email; do not use that email."
    return "medium", "Named person and role as the organisation or an official source publishes them."


def load(db: str) -> tuple[list[dict], list[dict]]:
    con = ro(DBS[db])
    level = {}
    if db == "government":
        level = {r["organisation_id"]: r["office_level"] for r in con.execute("SELECT organisation_id, office_level FROM government_office_profiles")}
    orgs = []
    for r in con.execute("SELECT * FROM organisations"):
        orgs.append({"db": db, "organisation_id": r["organisation_id"], "name": r["name"], "segment": level.get(r["organisation_id"]) or r["segment"] or "",
                     "website": r["website"] or "", "email": r["email"] or "", "phone": r["phone"] or "", "address": r["address"] or ""})
    by_id = {o["organisation_id"]: o for o in orgs}
    contacts = []
    for r in con.execute("SELECT * FROM contacts"):
        r = dict(r)
        org = by_id.get(r["organisation_id"], {})
        emails = [e for e in (r.get("named_email"), r.get("published_role_email"), r.get("shared_email")) if e]
        phones = [p for p in (r.get("role_phone"), r.get("organisation_phone")) if p]
        source = r.get("source_url") or r.get("profile_url") or ""
        verification = r.get("verification") or r.get("role_certainty") or ""
        risk, reason = contact_risk(r.get("name") or "", [r.get("named_email") or ""])
        # The best route to this person: their own published email or phone, else the organisation's published inbox or phone
        # (a personal-domain address linked to the person is never used).
        own_email = r.get("named_email") if r.get("named_email") and not W.PERSONAL_EMAIL.search(r["named_email"]) else ""
        best = next(((kind, value) for kind, value in (
            ("own email", own_email), ("role email", r.get("published_role_email")), ("own phone", r.get("role_phone")),
            ("organisation inbox", r.get("shared_email") or org.get("email")), ("organisation phone", r.get("organisation_phone") or org.get("phone")))
            if value), ("website or profile page only", source or ""))
        contacts.append({"db": db, "contact_id": r["contact_id"], "organisation_id": r["organisation_id"], "organisation": org.get("name", ""),
                         "name": r.get("name") or "", "role": r.get("role") or "", "decision_maker": decision_maker(r.get("role") or ""),
                         "contact_route": r.get("contact_route") or "", "named_email": r.get("named_email") or "",
                         "role_email": r.get("published_role_email") or r.get("shared_email") or "", "role_phone": r.get("role_phone") or "",
                         "organisation_phone": r.get("organisation_phone") or "", "direct": bool(emails or phones), "source": source,
                         "best_route_type": best[0], "best_route": best[1], "reachable": best[0] != "website or profile page only",
                         "attribution": r.get("channel_attribution") or "", "verification": verification, "pdpa_risk": risk, "pdpa_risk_reason": reason})
    con.close()
    return orgs, contacts


def research_flags(date_: str) -> list[list]:
    """Warnings the research agents and the crawler left for a person to check."""
    rows = []
    for path in sorted(C.RAW.glob(f"search_*_{date_}.jsonl")):
        for line in path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            r = json.loads(line)
            note = str(r.get("notes") or "")
            kinds = [kind for kind, _ in C.note_flags(note)]
            if r.get("identity") == "uncertain":
                kinds.append("identity uncertain")
            if kinds:
                rows.append([r.get("db", ""), r.get("organisation_id", ""), r.get("organisation_name", ""), "; ".join(kinds), note[:600], path.name])
    crawl = C.RAW / f"website_contacts_{date_}.jsonl"
    if crawl.exists():
        for line in crawl.read_text(encoding="utf-8").splitlines():
            site = json.loads(line)
            spam = [e for e in site.get("errors", []) if "spam or hijacked" in e]
            for o in site["orgs"] if spam else []:
                rows.append([o["db"], o["organisation_id"], o["name"], "hijacked pages on the organisation's own site",
                             f"{len(spam)} page(s) carry gambling content and were not used: " + "; ".join(spam)[:500], crawl.name])
            # Crawled although its robots.txt could not be read (the user's decision): the details are used, and flagged.
            readable = site.get("robots") == "unreachable" and any(p.get("status") == 200 for p in site.get("pages", []))
            for o in site["orgs"] if readable else []:
                rows.append([o["db"], o["organisation_id"], o["name"], "robots.txt unreachable",
                             f"{site['domain']}: robots.txt could not be read (server, network or certificate error); the site was crawled "
                             "anyway and its details used. Check the site's own terms if in doubt.", crawl.name])
    return rows


def add_sheet(wb, title, subtitle, headers, rows):
    ws = wb.create_sheet(title)
    ws["A1"], ws["A2"] = title, subtitle
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color=BLUE)
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="818283")
    ws.append([])
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    last = 4 + len(rows)
    for col, header in enumerate(headers, 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 48 if header in WRAP else min(30, max(12, len(header) + 2))
        if header in WRAP:
            for cell in ws[letter][4:last]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C5"
    if rows:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{last}"
    return ws


MEASURES = ["organisations", "direct route", "indirect only", "no route", "with named contact", "with named decision-maker", "contact leads",
            "named contact leads", "decision-maker leads", "contact leads with a direct route", "contact leads reachable (own or organisation route)",
            "pdpa low", "pdpa medium", "pdpa risky"]


def build(research: dict) -> tuple[dict, list, list, list]:
    """Per-database counts and the organisation, contact and gap rows, from the databases as they are now."""
    org_rows, lead_rows, gap_rows, summary = [], [], [], {}
    for db in DBS:
        orgs, contacts = load(db)
        by_org = defaultdict(list)
        for c in contacts:
            by_org[c["organisation_id"]].append(c)
        counts = Counter()
        for o in orgs:
            p = research.get((db, o["organisation_id"]), {})
            people = by_org.get(o["organisation_id"], [])
            named = [c for c in people if c["name"]]
            deciders = [c for c in named if c["decision_maker"]]
            direct = bool("@" in (o["email"] or "") or C.is_phone(o["phone"]) or any(c["direct"] for c in people))
            indirect = bool(o["website"] or o["address"])
            status = "direct route" if direct else ("indirect only" if indirect else "no route")
            counts["organisations"] += 1
            counts[status] += 1
            counts["with named contact"] += bool(named)
            counts["with named decision-maker"] += bool(deciders)
            socials = "; ".join(f"{k}: {v}" for k, v in (p.get("socials") or {}).items())
            research_note = ", ".join(p.get("methods") or []) or "not researched"
            if p.get("blocked"):
                research_note += f"; unreadable: {'; '.join(p['blocked'][:2])}"
            sources = "; ".join(s["url"] for s in (p.get("sources") or [])[:3])
            org_rows.append([db, o["organisation_id"], o["name"], o["segment"], status, o["email"], o["phone"], o["website"], o["address"], socials,
                             len(named), "; ".join(f"{c['name']} ({c['role']})" for c in deciders[:6]), len(people) - len(named), research_note, sources])
            if not direct:
                gap = ("No published email or phone found. " + ("A website, social page or postal address exists; use it to ask for the right "
                                                                "contact." if indirect or socials else "Nothing published; keep the organisation on hold."))
                gap_rows.append([db, o["organisation_id"], o["name"], o["segment"], "; ".join(x for x in (o["website"], socials, o["address"]) if x),
                                 research_note, "; ".join((p.get("notes") or [])[:2])[:400], gap])
        for c in contacts:
            counts["contact leads"] += 1
            counts["named contact leads"] += bool(c["name"])
            counts["decision-maker leads"] += bool(c["name"] and c["decision_maker"])
            counts["contact leads with a direct route"] += c["direct"]
            counts["contact leads reachable (own or organisation route)"] += c["reachable"]
            counts[f"pdpa {c['pdpa_risk']}"] += 1
            lead_rows.append([db, c["contact_id"], c["organisation"], c["name"] or "(office or role desk)", c["role"], "yes" if c["decision_maker"] else "no",
                              c["best_route_type"], c["best_route"], c["contact_route"], c["named_email"], c["role_email"], c["role_phone"], c["organisation_phone"], c["source"], c["attribution"],
                              c["verification"], c["pdpa_risk"], c["pdpa_risk_reason"]])
        summary[db] = {m: counts.get(m, 0) for m in MEASURES}
    return summary, org_rows, lead_rows, gap_rows


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", default=date.today().isoformat())
    parser.add_argument("--baseline", action="store_true",
                        help="before any merge: record the databases' counts in runtime/contacts/baseline-counts.json, so the workbook shows before and after")
    args = parser.parse_args()
    baseline_path = C.WORK / "baseline-counts.json"
    if args.baseline:
        summary = build({})[0]
        C.WORK.mkdir(parents=True, exist_ok=True)
        baseline_path.write_text(json.dumps({"date": args.date, **summary}, indent=1) + "\n", encoding="utf-8")
        print(json.dumps(summary, indent=1))
        return 0
    profiles_path = C.WORK / "profiles.json"
    profiles = json.loads(profiles_path.read_text(encoding="utf-8"))["profiles"] if profiles_path.exists() else []
    research = {(p["db"], p["organisation_id"]): p for p in profiles}
    summary_path = C.WORK / "profiles-summary.json"
    research_summary = json.loads(summary_path.read_text(encoding="utf-8")) if summary_path.exists() else {}
    before = json.loads(baseline_path.read_text(encoding="utf-8")) if baseline_path.exists() else {}
    summary, org_rows, lead_rows, gap_rows = build(research)
    flag_rows = research_flags(args.date)
    for db in DBS:
        summary[db]["research flags"] = sum(1 for r in flag_rows if r[0] == db)

    wb = Workbook()
    wb.remove(wb.active)
    if before:
        headers = ["measure", *(f"{db} {when}" for db in DBS for when in ("before", "after"))]
        rows = [[m, *(v for db in DBS for v in (before.get(db, {}).get(m, ""), summary[db].get(m, 0)))] for m in [*MEASURES, "research flags"]]
    else:
        headers = ["measure", *DBS]
        rows = [[m, *(summary[db].get(m, 0) for db in DBS)] for m in [*MEASURES, "research flags"]]
    rows.append([])
    sites = research_summary.get("sites", {})
    rows += [["Research: own websites crawled", research_summary.get("sites_crawled", 0)],
             ["Research: sites readable / unreadable / robots.txt disallowed", f"{sites.get('readable', 0)} / {sites.get('unreadable', 0)} / "
              f"{sites.get('robots_disallowed', 0)}"],
             ["Research: readable sites crawled with robots.txt unreachable (flagged)", sites.get("readable, crawled with robots.txt unreachable (flagged)", 0)],
             ["Research: OpenStreetMap exact-name matches", research_summary.get("osm_matches", 0)],
             ["Research: budgeted search records", json.dumps(research_summary.get("search_records", {}))],
             ["Research: extracted people not recorded", json.dumps(research_summary.get("people_not_recorded", {}))]]
    add_sheet(wb, "Summary", f"Contact profiles across the three databases, built {args.date}"
              + (f"; 'before' is the databases on {before.get('date')} before the contact merges" if before else "")
              + ". Routes are what each organisation or an official source publishes; nothing private was inferred. Sending and automations stay "
              "disabled.", headers, rows)
    add_sheet(wb, "Organisation Profiles", "One row per organisation. 'direct route' means a published email or phone for the organisation or one of its "
              "contacts.", ["database", "organisation_id", "name", "segment", "route_status", "email", "phone", "website", "address", "socials",
                            "named_contacts", "decision_makers", "role_desks", "research", "sources"], org_rows)
    add_sheet(wb, "Contact Leads", "One row per contact. A named email is used only when the organisation publishes it for that person; a "
              "personal-domain address linked to a person is never a route.", ["database", "contact_id", "organisation", "name", "role", "decision_maker",
                                                                                  "best_route_type", "best_route", "contact_route", "named_email", "role_email", "role_phone",
                                                                                  "organisation_phone", "source", "attribution", "verification", "pdpa_risk",
                                                                                  "pdpa_risk_reason"], lead_rows)
    add_sheet(wb, "Gaps", "Organisations still without a published email or phone after the research, with what was tried.",
              ["database", "organisation_id", "name", "segment", "indirect_routes", "research", "note", "gap"], gap_rows)
    add_sheet(wb, "Flags", "Warnings from the research for a person to check before outreach. Nothing was changed or deleted because of them; "
              "the company master holds drafts for possible closures.", ["database", "organisation_id", "name", "flag", "note", "source"], flag_rows)
    out = C.ROOT / "outputs" / "contacts"
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"Silverleaf Contact Profiles - {args.date}.xlsx"
    wb.save(path)
    (out / "contact-profiles-summary.json").write_text(json.dumps({"date": args.date, "workbook": path.relative_to(C.ROOT).as_posix(), "before": before,
                                                                   "after": summary, "research": research_summary}, indent=1) + "\n", encoding="utf-8")
    print(json.dumps({"before": before, "after": summary}, indent=1))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
