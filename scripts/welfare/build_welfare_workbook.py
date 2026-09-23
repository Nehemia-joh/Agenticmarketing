#!/usr/bin/env python3
"""Generate the welfare review workbook from the run SQLite database only.

Sheets follow skills/silverleaf-create-lead-list/references/workbook-layout.md plus welfare sheets. The workbook is
written to outputs/runs/<run-id>/ and never inside outputs/master/ (the company-leads master stays separate).
"""
from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from collections import Counter
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import welfare_lib as W

sys.path.insert(0, str(W.ROOT / "scripts" / "messaging"))
import offer_lib as OL  # noqa: E402  (offer-aligned drafts: sheet rows)

BLUE, WHITE = "002368", "FFFFFF"
RISK_FILL = {"low": "E2F0D9", "medium": "FFF2CC", "risky": "F8CBAD"}
WRAP_COLS = {"description", "notes", "evidence_excerpt", "request", "size_evidence", "education_angle", "track_reason", "known_funders_or_partners",
             "foreign_charity_registrations", "social_media", "services", "red_flags", "pdpa_risk_reason", "vision", "projects", "detail",
             "action", "text", "channel_attribution", "schooling_arrangement_published", "funding_model_published", "all_source_urls", "missing_information",
             "subject", "body", "follow_up_1", "english_meaning", "conditions", "recipient"}
ORG_HEADERS = ["organisation_id", "record_key", "name", "welfare_fit", "segment", "subtype", "care_model", "proposed_welfare_track", "track_reason",
               "desk_tier", "priority_score", "factors_known", "verification_status", "pdpa_risk", "pdpa_risk_reason",
               "campus", "distance_km", "transport_band", "nearest_primary_campus", "distance_to_primary_km", "geocode_precision",
               "latitude", "longitude", "locality", "address", "website", "email", "all_emails", "phone", "all_phones", "social_media",
               "children_served_published", "size_evidence", "children_served_as_of", "age_range_published", "gender_served", "services",
               "schooling_arrangement_published", "runs_own_school", "funding_model_published", "known_funders_or_partners",
               "foreign_charity_registrations", "operator_or_umbrella", "religious_affiliation_published", "founded", "registration_published",
               "nis_reg_no", "nis_reg_date", "nis_profile_url", "licence_status", "volunteer_programme", "volunteer_fee_published",
               "safeguarding_policy_url", "staff_count_published", "latest_activity_date", "red_flags", "catchment_note", "osm_id",
               "master_organisation_id", "notes", "source_url", "evidence_basis", "evidence_excerpt", "all_source_urls"]
CONTACT_HEADERS = ["contact_id", "name", "role", "role_certainty", "organisation", "organisation_id", "pdpa_risk", "pdpa_risk_reason",
                   "named_email", "published_role_email", "shared_email", "all_emails", "email_type", "organisation_phone", "role_phone",
                   "all_phones", "phone_type", "profile_url", "contact_route", "channel_attribution", "verification_status", "source_url",
                   "evidence_excerpt", "notes"]
ENQUIRY_HEADERS = ["enquiry_id", "name", "author_type_stated", "type", "enquiry_date", "date_qualification", "request", "locality", "campus_fit",
                   "platform", "source_url", "phone", "email", "contact_attribution", "current_relevance", "pdpa_risk", "pdpa_risk_reason", "notes"]
NIS_HEADERS = ["nis_id", "name", "welfare_relevance", "welfare_fit", "relevance_signal", "matched_organisation", "matched_organisation_id", "reg_no", "reg_date",
               "level", "region", "district", "years_experience", "latest_project_end", "nearest_campus_registry_coords",
               "distance_km_registry_coords", "profile_url", "profile_fetched", "vision"]
REVIEW_HEADERS = ["review_id", "kind", "entity_type", "entity_id", "field", "values_json", "action", "status"]


def payloads(con, entity_type):
    out = {}
    for eid, payload in con.execute("SELECT es.entity_id, sr.payload_json FROM entity_sources es JOIN source_records sr "
                                    "ON sr.source_record_id = es.source_record_id WHERE es.entity_type = ?", (entity_type,)):
        out.setdefault(eid, json.loads(payload))
    return out


def urls_from(sources_json: str) -> list[str]:
    try:
        return [s.get("url") for s in json.loads(sources_json or "[]") if isinstance(s, dict) and s.get("url")]
    except json.JSONDecodeError:
        return []


def add_sheet(wb, title, subtitle, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws["A1"], ws["A2"] = title, subtitle
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color=BLUE)
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="818283")
    ws.append([])
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    last_row = 4 + len(rows)
    for col, header in enumerate(headers, 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = (widths or {}).get(header, 48 if header in WRAP_COLS else min(40, max(12, len(header) + 2)))
        if header in WRAP_COLS:
            for cell in ws[letter][4:last_row]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if header.endswith("url") or header in ("website", "source_url", "profile_url", "nis_profile_url"):
            for cell in ws[letter][4:last_row]:
                if isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = Font(name="Arial", size=10, color="1F4E9A", underline="single")
        if header == "pdpa_risk":
            for cell in ws[letter][4:last_row]:
                if cell.value in RISK_FILL:
                    cell.fill = PatternFill("solid", fgColor=RISK_FILL[cell.value])
    ws.freeze_panes = "C5"
    if rows:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{last_row}"
    return ws


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-id", required=True)
    args = parser.parse_args()
    cfg = W.load_config(args.run_id)
    db, workbook_path = cfg["paths"]["db"], cfg["paths"]["workbook"]
    master_dir = (W.ROOT / "outputs" / "master").resolve()
    if master_dir in workbook_path.resolve().parents:
        raise SystemExit("Refusing to write the welfare workbook inside outputs/master: it must stay separate from the company-leads master.")
    con = sqlite3.connect(f"file:{db.as_posix()}?mode=ro", uri=True)
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    org_payload, con_payload = payloads(con, "organisation"), payloads(con, "contact")

    org_cols = [d[0] for d in con.execute("SELECT * FROM organisations LIMIT 0").description]
    org_rows, org_names = [], {}
    for row in con.execute("SELECT * FROM organisations ORDER BY desk_score DESC, distance_km"):
        o = dict(zip(org_cols, row))
        p = org_payload.get(o["organisation_id"], {})
        merged = {**p, **{k: v for k, v in o.items() if v not in (None, "")}, "name": o["name"],
                  "all_source_urls": "\n".join(urls_from(p.get("all_sources_json")))}
        org_names[o["organisation_id"]] = o["name"]
        org_rows.append([merged.get(h, "") for h in ORG_HEADERS])

    contact_cols = [d[0] for d in con.execute("SELECT c.*, o.name AS organisation FROM contacts c JOIN organisations o USING(organisation_id) LIMIT 0").description]
    contact_rows = []
    for row in con.execute("SELECT c.*, o.name AS organisation FROM contacts c JOIN organisations o USING(organisation_id) ORDER BY o.name, c.name"):
        c = dict(zip(contact_cols, row))
        merged = {**con_payload.get(c["contact_id"], {}), **{k: v for k, v in c.items() if v not in (None, "")}}
        merged["name"], merged["organisation"] = c["name"], c["organisation"]
        contact_rows.append([merged.get(h, "") for h in CONTACT_HEADERS])

    enquiry_sql = "SELECT e.*, sr.payload_json FROM enquiries e JOIN source_records sr ON sr.source_record_id = e.source_record_id"
    enquiry_cols = [d[0] for d in con.execute(enquiry_sql + " LIMIT 0").description]
    enquiry_rows = []
    for row in con.execute(enquiry_sql + " ORDER BY enquiry_date DESC"):
        e = dict(zip(enquiry_cols, row))
        p = json.loads(e.pop("payload_json"))
        merged = {**p, **{k: v for k, v in e.items() if v not in (None, "")}}
        enquiry_rows.append([merged.get(h, "") for h in ENQUIRY_HEADERS])

    rel_headers = ["relationship_id", "from_organisation", "relationship_type", "to_organisation", "from_organisation_id", "to_organisation_id",
                   "verification_status", "pdpa_risk", "notes", "source_urls"]
    rel_rows = [list(r[:9]) + ["\n".join(urls_from(r[9]))] for r in con.execute(
        "SELECT relationship_id, from_organisation, relationship_type, to_organisation, from_organisation_id, to_organisation_id, "
        "verification_status, pdpa_risk, notes, sources_json FROM organisation_relationships ORDER BY to_organisation")]
    nis_rows = [list(r) for r in con.execute(
        f"SELECT {', '.join(NIS_HEADERS)} FROM registry_ngos ORDER BY CASE welfare_relevance WHEN 'strong' THEN 0 WHEN 'medium' THEN 1 "
        "WHEN 'medium_projects_only' THEN 2 WHEN 'weak' THEN 3 ELSE 4 END, distance_km_registry_coords")]
    review_rows = [list(r) for r in con.execute(f"SELECT {', '.join(REVIEW_HEADERS)} FROM review ORDER BY kind, entity_id")]

    evidence_rows, sources = [], {}
    for entity_type, payload_map in (("organisation", org_payload), ("contact", con_payload)):
        for eid, p in payload_map.items():
            for s in json.loads(p.get("all_sources_json") or "[]"):
                if not isinstance(s, dict) or not s.get("url"):
                    continue
                label = org_names.get(eid) or f"{p.get('contact_name') or p.get('role')} ({p.get('organisation_name')})"
                evidence_rows.append([entity_type, eid, label, s.get("url"), s.get("title", ""), s.get("source_date", ""), s.get("accessed_on", ""),
                                      s.get("evidence_basis", ""), "yes" if s.get("fetched") else "no", ", ".join(s.get("facts_supported") or []),
                                      s.get("evidence_excerpt", "")])
                src = sources.setdefault(s["url"], {"title": s.get("title", ""), "basis": set(), "fetched": False, "cited": 0, "first": label})
                src["basis"].add(s.get("evidence_basis", ""))
                src["fetched"] |= bool(s.get("fetched"))
                src["cited"] += 1
    for (url,) in con.execute("SELECT DISTINCT source_url FROM enquiries"):
        sources.setdefault(url, {"title": "public enquiry", "basis": {"public_post"}, "fetched": True, "cited": 1, "first": "enquiry"})
    source_rows = [[u, u.split("/")[2] if "//" in u else "", v["title"], v["cited"], ", ".join(sorted(b for b in v["basis"] if b)),
                    "yes" if v["fetched"] else "no", v["first"]] for u, v in sorted(sources.items(), key=lambda kv: -kv[1]["cited"])]
    coverage_rows = [list(r) for r in con.execute("SELECT slice, line_no, text FROM research_coverage ORDER BY slice, line_no")]

    counts = {t: con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in
              ("organisations", "contacts", "enquiries", "organisation_relationships", "registry_ngos", "review", "source_records")}
    col = {h: i for i, h in enumerate(ORG_HEADERS)}
    seg, track = Counter(r[col["segment"]] for r in org_rows), Counter(r[col["proposed_welfare_track"]] for r in org_rows)
    band, tier = Counter(r[col["transport_band"]] for r in org_rows), Counter(r[col["desk_tier"]] for r in org_rows)
    crisk = Counter(r[CONTACT_HEADERS.index("pdpa_risk")] for r in contact_rows)
    slices = sorted({r[0] for r in coverage_rows})
    wb = Workbook()
    ws = wb.active
    ws.title = "Read Me"
    lines = [(f"Silverleaf welfare leads — research run {args.run_id}", "title"),
             (f"Generated {generated} from {db.name} (SQLite is the source of truth; this workbook is a generated review view). "
              "Separate from the company-leads master in outputs/master/.", "sub"), ("", ""), ("Scope", "h"),
             (cfg.get("scope", "Welfare institutions near Silverleaf's campuses."), ""),
             ("Drafts only, request first: the Outreach Plans sheet holds a meeting request for each organisation (sponsorship requests for "
              "funders only), signed by Mariam Haji, and the offer register v4 terms in the follow-up. Nothing was sent or scheduled, and every "
              "automation remains disabled.", ""), ("", ""),
             ("Coverage and accuracy", "h"),
             (f"Sources: {len(slices)} research slices (web), the NGOs Information System register ({counts['registry_ngos']} NGOs pinned within "
              f"{cfg.get('register_radius_km', 30)} km) and OpenStreetMap. Research period: {cfg.get('research_period', cfg['research_date'])}.", "")]
    lines += [(note, "") for note in cfg.get("readme_notes", [])]
    lines += [("", ""), ("Counts", "h"),
              (f"Organisations {counts['organisations']} · Contacts {counts['contacts']} · Parent/guardian enquiries {counts['enquiries']} · "
               f"Relationships {counts['organisation_relationships']} · NGO register rows {counts['registry_ngos']} · Review items {counts['review']} · "
               f"Source records {counts['source_records']}", ""),
              ("Organisations by segment: " + "; ".join(f"{k or 'unclassified'} {v}" for k, v in seg.most_common()), ""),
              ("Organisations by proposed track: " + "; ".join(f"{k} {v}" for k, v in track.most_common()), ""),
              ("Organisations by priority tier: " + "; ".join(f"{k} {v}" for k, v in sorted(tier.items())), ""),
              ("Organisations by transport band (nearest campus): " + "; ".join(f"{k} {v}" for k, v in band.most_common()), ""),
              ("Contacts by data-protection risk: " + "; ".join(f"{k} {v}" for k, v in crisk.most_common()), ""), ("", ""),
              ("Data-protection risk labels (pdpa_risk)", "h"),
              ("low — organisation-level facts and generic organisational routes published by the organisation.", ""),
              ("medium — a named individual in a professional role as published by the organisation, an official register or their own professional profile; named work email; organisation-published mobile.", ""),
              ("risky — private individuals (parents, guardians), personal-domain emails or personal mobiles of named people, personal social-media data, regulator-listed individuals, anything touching household circumstances or children. Do not contact risky rows until Silverleaf's data-protection owner or counsel has approved a lawful basis.", ""),
              ("Tanzania's Personal Data Protection Act 2022 treats data related to children as sensitive and (per DLA Piper's summary) requires explicit consent for direct marketing. Not legal advice.", ""),
              ("", ""), ("Deliberately not collected", "h"),
              ("Any child's name, photo, age-with-identity, story, health information or sponsorship profile; anything about the parents or relatives of children in care; private individuals' addresses; closed-group or logged-in content; and any cross-source profiling of individuals.", ""),
              ("", ""), ("Status and track definitions", "h"),
              ("verification_status: verified (key facts confirmed on a fetched own/official page) · needs_review (conflicting, partial or registry-only) · unverified (snippet or low-quality source) · historical (no sign of operation in about 3 years).", ""),
              ("proposed_welfare_track (plans/b2b-welfare-leads-plan.md §5.4): WA00 Hold — unresolved identity, route, location or red flags · WA01 Routing — usable route; decision-maker and licence status still to confirm · WA02 Consultative — requires licence confirmation by the authority, so no row is WA02 yet. Excluded — out of scope.", ""),
              ("priority_score (0–100) and desk_tier P1–P4 use the plan's rubric; factors_known shows how many of the six factors had evidence. A low score often means 'unknown', not 'poor fit'.", ""),
              ("Distances are straight-line from approximate campus centroids. NGO-register map pins are self-reported and can be wrong; see geocode_precision.", ""),
              ("", ""), ("Sheets", "h"),
              ("Organisations · Contacts · Parent Enquiries · Relationships · NGO Register · Review Queue · Evidence · Sources · Search Coverage · then Outreach Plans (offer-aligned drafts) and the create-skill layout sheets (Campaigns, Touchpoints, Assignments, Strategies, Automation), which hold design-only material.", "")]
    for i, (text, kind) in enumerate(lines, 1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(name="Arial", size=16 if kind == "title" else 11 if kind == "h" else 10, bold=kind in ("title", "h"),
                         italic=kind == "sub", color=BLUE if kind in ("title", "h") else "000000")
    ws.column_dimensions["A"].width = 150

    add_sheet(wb, "Organisations", "One row per welfare institution or funder. Filter proposed_welfare_track, desk_tier, transport_band and pdpa_risk.",
              ORG_HEADERS, org_rows, {"name": 38, "segment": 26, "subtype": 26, "track_reason": 40})
    add_sheet(wb, "Contacts", "Published business contacts only, each attributed to the source that published it. Check pdpa_risk before any use.", CONTACT_HEADERS, contact_rows)
    add_sheet(wb, "Parent Enquiries", "Public parent/guardian enquiries as published. Every row is pdpa_risk = risky; do not contact without an approved lawful basis.",
              ENQUIRY_HEADERS, enquiry_rows)
    add_sheet(wb, "Relationships", "Funding, operating and partnership links between organisations (warm paths).", rel_headers, rel_rows)
    add_sheet(wb, "NGO Register", "Every NGO in the NGOs Information System within the register radius (registry pins), with welfare relevance and matches.", NIS_HEADERS, nis_rows)
    add_sheet(wb, "Review Queue", "Conflicts, possible duplicates, unmatched contacts, outliers and other decisions for a person.", REVIEW_HEADERS, review_rows)
    add_sheet(wb, "Evidence", "One row per entity and source, with the facts it supports.",
              ["entity_type", "entity_id", "entity", "url", "title", "source_date", "accessed_on", "evidence_basis", "fetched", "facts_supported", "evidence_excerpt"], evidence_rows)
    add_sheet(wb, "Sources", "Unique source URLs.", ["url", "domain", "title", "times_cited", "evidence_basis", "fetched", "first_entity"], source_rows)
    add_sheet(wb, "Search Coverage", "What each research slice searched, what was blocked, and known gaps.", ["slice", "line_no", "text"], coverage_rows, {"text": 140})
    note = "Design only; nothing is scheduled or enabled. See plans/b2b-welfare-leads-plan.md."
    outreach_rows = OL.outreach_sheet_rows(con)
    add_sheet(wb, "Outreach Plans", "Request-first drafts: body asks for a short meeting and states no offer terms; follow_up_1 states the offer "
              "register v4 terms. One per in-scope organisation, plus replies to in-fit parent enquiries. Finance must confirm the 2027 terms before "
              "a follow-up is sent; needs_review rows are held for the reason in conditions. Nothing is sent.",
              OL.OUTREACH_HEADERS, outreach_rows, {"target": 36, "subject": 40, "body": 90, "follow_up_1": 70, "conditions": 70, "recipient": 30})
    add_sheet(wb, "Campaigns", note, ["campaign_id", "name", "audience", "activation_gate", "status"],
              [["C10", "Welfare institutional placements", "Verified welfare institutions within the catchment",
                "Approved terms memo (Finance); safeguarding visit protocol; confirmed capacity for the next intake; PDPC registration confirmed", "Design only; not activated"]])
    add_sheet(wb, "Touchpoints", note, ["flow_id", "step", "trigger", "action", "pipeline_stage"],
              [["F14", 1, "WA02 selected after review", "Call or email the published route", "contacted"],
               ["F14", 2, "No reply after 4 working days", "One follow-up", "contacted"],
               ["F14", 3, "Conversation held", "Record aggregate age-band counts; admissions confirms level and capacity", "needs_assessed"],
               ["F14", 4, "Visit agreed", "Two-person visit under the safeguarding protocol; no photography", "visit_held"],
               ["F14", 5, "Visit held", "Proposal on approved terms within 5 working days", "proposal"],
               ["F14", 6, "Proposal accepted", "Admissions assesses each child in Ed-admin, outside the lead database", "placement"],
               ["F14", 7, "Eight weeks before the next intake window", "Check next year's places", "renewal"]])
    add_sheet(wb, "Assignments", note, ["status"], [["No campaign assignments in this research-only run."]])
    add_sheet(wb, "Strategies", note, ["module_id", "name", "silverleaf_contribution", "do_not_imply", "gate"],
              [["VM08", "Placement planning", "20–30 minute conversation using published levels and intake windows", "Guaranteed places or capacity", "Admissions confirms capacity"],
               ["VM09", "Institutional billing", "Consolidated invoice on the published four-instalment schedule", "Discount, credit or deferral", "Finance approval"],
               ["VM10", "Sponsor-ready reporting", "Termly progress summary in an approved format to the authorised institution", "Outcome guarantees; donor sharing without consent", "Academic and safeguarding approval"],
               ["VM11", "Safeguarding and wellbeing", "Share the child-protection policy; explain the wellness programme", "Certification or specialist therapy", "Policy confirmed current"],
               ["VM12", "Transport", "Published transport bands", "Route availability", "Operations confirms the route"],
               ["VM13", "Documented partner rate", "NGO partner discount OF03 (3-18% per child when a home places all its primary-age children), from the offer register", "Any term not in the offer register", "Finance confirms the 2027 terms"]])
    add_sheet(wb, "Automation", note, ["flow_id", "status"], [["F14", "DISABLED — design only"]])
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(workbook_path)
    sheet_counts = {"Organisations": len(org_rows), "Contacts": len(contact_rows), "Parent Enquiries": len(enquiry_rows), "Relationships": len(rel_rows),
                    "NGO Register": len(nis_rows), "Review Queue": len(review_rows), "Evidence": len(evidence_rows), "Sources": len(source_rows),
                    "Outreach Plans": len(outreach_rows)}
    db_counts = {"Organisations": counts["organisations"], "Contacts": counts["contacts"], "Parent Enquiries": counts["enquiries"],
                 "Relationships": counts["organisation_relationships"], "NGO Register": counts["registry_ngos"], "Review Queue": counts["review"],
                 "Outreach Plans": con.execute("SELECT COUNT(*) FROM outreach_plans").fetchone()[0]}
    mismatch = {k: (sheet_counts[k], v) for k, v in db_counts.items() if sheet_counts[k] != v}
    con.close()
    print(json.dumps({"workbook": str(workbook_path.relative_to(W.ROOT)), "sheet_rows": sheet_counts, "mismatches": mismatch}, indent=1))
    return 0 if not mismatch else 1


if __name__ == "__main__":
    raise SystemExit(main())
