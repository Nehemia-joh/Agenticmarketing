#!/usr/bin/env python3
"""Capture the official council and regional websites through their public JSON API (GWF CORE), politely.

The sites render in the browser, so plain page fetches return only a title. Their API returns the same content:
/api/profiles (leaders), /api/statistics, /api/footer (official contact), /api/menus, /api/pages/slug/<slug>,
/api/files, /api/news and /api/announcements. One request at a time per host, at least 1.5 s apart (gov_lib), with
retries and an on-disk cache under runtime/government/<run-id>/http-cache/ (never committed).

Captures are sanitised before they are written: CMS editor names and emails (createdBy), photos, biographies and the
body text of news items are dropped (news keeps its title, date, category and convening tags), and the SHA-256 of
every original response is kept for traceability. Attachments on the
councillor, ward and contact pages are fetched into the cache and their text extracted: XLSX rows, PDF text, or a
note that a scanned PDF needs transcription.

Output: data/raw/government-research/council_sites/<host>_<date>.json, one file per host.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import urllib.parse

import gov_lib as G

W = G.W
PAGE_PATTERN = re.compile(r"counc?il+ors?|concilor|diwani|madiwani|\bwards?\b|\bkata\b|division|tarafa|village|vijiji|kijiji|\bmtaa\b|mitaa|"
                          r"street|community|maendeleo|social|ustawi|contact|mawasiliano|history|historia|about|location|administrative|"
                          r"structure|muundo|viongozi|leaders|profile|region-districts|district|council", re.I)
ATTACH_PAGE = re.compile(r"counc?il+ors?|concilor|diwani|madiwani|\bwards?\b|\bkata\b|tarafa|division|contact|mawasiliano", re.I)
CONVENING_TERMS = {"public_meeting": r"mkutano wa hadhara|mikutano ya hadhara|public meeting|baraza|mabaraza",
                   "village_or_mtaa_assembly": r"mkutano mkuu wa (kijiji|mtaa)|mikutano mikuu|village assembl",
                   "ward_meeting": r"mkutano wa kata|kamati ya maendeleo ya kata|\bwdc\b|ward development",
                   "parents_or_education": r"wazazi|walezi|parents|elimu|shule|education",
                   "exhibition_or_public_event": r"maonesho|maonyesho|nane ?nane|exhibition|maadhimisho|kilele cha|wiki ya|siku ya",
                   "officials_training": r"watendaji wa (kata|vijiji|mitaa)|maafisa tarafa|wenyeviti wa (vijiji|mitaa)",
                   "election_period": r"uchaguzi|election|kampeni|campaign"}


def fetch_json(url: str, cache, errors: list):
    try:
        body = W.polite_request(url, cache_dir=cache, retries=2, backoff=5.0)
    except Exception as exc:  # noqa: BLE001 - record and continue with the next endpoint
        errors.append({"url": url, "error": repr(exc)[:200]})
        return None, ""
    try:
        return json.loads(body), hashlib.sha256(body).hexdigest()
    except json.JSONDecodeError:
        errors.append({"url": url, "error": f"not JSON ({body[:60]!r})"})
        return None, hashlib.sha256(body).hexdigest()


def flatten_menus(items, trail=()) -> list[dict]:
    out = []
    for item in items or []:
        title = str(item.get("title") or "").strip()
        slug = item.get("slug") or (item.get("page") or {}).get("slug")
        path = trail + (title,)
        if slug:
            out.append({"path": " > ".join(path), "title": title, "title_sw": item.get("swahiliTitle") or "", "slug": slug})
        out += flatten_menus(item.get("children"), path)
    return out


def attachment_text(url: str, file_name: str, cache, errors: list) -> dict:
    record = {"file_name": file_name, "url": url}
    try:
        body = W.polite_request(url, cache_dir=cache, retries=2, backoff=5.0)
    except Exception as exc:  # noqa: BLE001
        errors.append({"url": url, "error": repr(exc)[:200]})
        record["status"] = "fetch_failed"
        return record
    record.update(sha256=hashlib.sha256(body).hexdigest(), bytes=len(body))
    lower = file_name.lower()
    try:
        if lower.endswith((".xlsx", ".xlsm")) or body[:2] == b"PK" and not lower.endswith(".docx"):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [" ".join(str(c).split()) for c in row if c not in (None, "")]
                    if cells:
                        rows.append(f"{ws.title} | " + " | ".join(cells))
            record.update(status="text_extracted", kind="xlsx", text="\n".join(rows))
        elif lower.endswith(".pdf") or body[:4] == b"%PDF":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(body))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
            status = "text_extracted" if len(text.strip()) > 40 else "scanned_needs_transcription"
            record.update(status=status, kind="pdf", pages=len(reader.pages), text=text.strip())
        else:
            record.update(status="not_parsed", kind=lower.rsplit(".", 1)[-1])
    except Exception as exc:  # noqa: BLE001
        record.update(status="parse_failed", error=repr(exc)[:200])
    return record


def tag_convening(text: str) -> list[str]:
    low = text.lower()
    return [tag for tag, pattern in CONVENING_TERMS.items() if re.search(pattern, low)]


def collect_host(site: dict, cfg: dict) -> dict:
    host, cache, errors = site["host"], cfg["paths"]["cache"], []
    base = f"https://{host}"
    hashes = {}
    out = {"host": host, "organisation": site["name"], "retrieved": cfg["research_date"], "base_url": base}
    for key, path in (("profiles", "/api/profiles"), ("statistics", "/api/statistics"), ("footer", "/api/footer"), ("menus", "/api/menus")):
        data, digest = fetch_json(base + path, cache, errors)
        hashes[base + path] = digest
        out[key] = data
    if out["profiles"] is None and out["menus"] is None:
        out.update(status="unreachable", errors=errors, response_sha256=hashes)
        return out
    out["profiles"] = [{k: p.get(k) for k in ("id", "name", "title", "titleSwahili", "order", "published", "createdAt", "updatedAt")}
                       for p in (out["profiles"] or []) if isinstance(p, dict)]
    stats = out["statistics"] or {}
    out["statistics"] = [{"label": s.get("label"), "value": s.get("value"), "suffix": s.get("suffix")} for s in stats.get("data", [])] if isinstance(stats, dict) else []
    footer = out["footer"] if isinstance(out["footer"], dict) else {}
    out["footer"] = G.sanitise({k: v for k, v in footer.items() if k not in ("translations",)})
    menu = flatten_menus(out["menus"] if isinstance(out["menus"], list) else [])
    out["menus"] = menu
    pages, seen = [], set()
    for item in menu:
        slug = item["slug"]
        if slug in seen or not PAGE_PATTERN.search(item["path"] + " " + slug):
            continue
        seen.add(slug)
        url = f"{base}/api/pages/slug/{urllib.parse.quote(slug)}"
        data, digest = fetch_json(url, cache, errors)
        hashes[url] = digest
        if not isinstance(data, dict):
            continue
        content = data.get("content") or {}
        page = {"slug": slug, "menu_path": item["path"], "title": data.get("title") or item["title"], "api_url": url,
                "public_url": f"{base}/{slug}", "createdAt": data.get("createdAt"), "updatedAt": data.get("updatedAt"),
                "text_en": G.html_to_text((content.get("en") or {}).get("html", "")) if isinstance(content, dict) else "",
                "text_sw": G.html_to_text((content.get("sw") or {}).get("html", "")) if isinstance(content, dict) else "",
                "attachments": []}
        for att in data.get("attachments") or []:
            entry = {"file_name": att.get("fileName"), "url": att.get("url"), "file_type": att.get("fileType"), "file_size": att.get("fileSize")}
            if att.get("url") and ATTACH_PAGE.search(item["path"] + " " + slug):
                entry.update(attachment_text(att["url"], att.get("fileName") or "", cache, errors))
            page["attachments"].append(entry)
        pages.append(page)
    out["pages"] = pages
    files, page_no = [], 1
    while True:
        url = f"{base}/api/files?page={page_no}&limit=100"
        data, digest = fetch_json(url, cache, errors)
        hashes[url] = digest
        if not isinstance(data, dict):
            break
        files += [{"name": f.get("name"), "date": f.get("date"), "url": f.get("url"), "file_type": f.get("fileType"), "size": f.get("size")}
                  for f in data.get("data", [])]
        if page_no >= int(data.get("totalPages") or 1):
            break
        page_no += 1
    out["files"] = files
    for key in ("news", "announcements"):
        items, page_no = [], 1
        while True:
            url = f"{base}/api/{key}?page={page_no}&limit=50"
            data, digest = fetch_json(url, cache, errors)
            hashes[url] = digest
            if not isinstance(data, dict):
                break
            for n in data.get("data", []):
                text = G.html_to_text(n.get("content") or n.get("excerpt") or "")
                # The body text is only used to tag the item; it is not kept, so people named in news stories are not stored.
                items.append({"id": n.get("id"), "title": " ".join(str(n.get("title") or "").split()), "date": str(n.get("date") or "")[:10],
                              "category": n.get("category"), "listing_url": url, "text_chars": len(text),
                              "convening_tags": tag_convening(f"{n.get('title', '')} {text}"),
                              "attachments": [a.get("fileName") for a in n.get("attachments") or []]})
            if page_no >= int(data.get("totalPages") or 1) or page_no >= 20:
                break
            page_no += 1
        out[key] = items
    out.update(status="captured", errors=errors, response_sha256=hashes,
               sanitised="createdBy/updatedBy (CMS editors), images, biographies and footer translations removed")
    return out


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--host", action="append", help="limit to these hosts (repeatable)")
    parser.add_argument("--refresh", action="store_true", help="re-capture hosts that already have a file for this date")
    args = parser.parse_args()
    cfg = G.load_config(args.run_id)
    out_dir = cfg["paths"]["raw"] / "council_sites"
    out_dir.mkdir(parents=True, exist_ok=True)
    sites = [s for s in cfg["councils"] + cfg["regions"] if s.get("collect_site", True)]
    if args.host:
        sites = [s for s in sites if s["host"] in args.host]
    for site in sites:
        path = out_dir / f"{site['host']}_{cfg['research_date']}.json"
        if path.exists() and not args.refresh:
            print(f"reuse {path.relative_to(G.ROOT)}")
            continue
        result = collect_host(site, cfg)
        path.write_text(json.dumps(result, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
        print(f"{site['host']}: {result['status']}; profiles {len(result.get('profiles') or [])}, pages {len(result.get('pages') or [])}, "
              f"files {len(result.get('files') or [])}, news {len(result.get('news') or [])}, errors {len(result.get('errors') or [])}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
