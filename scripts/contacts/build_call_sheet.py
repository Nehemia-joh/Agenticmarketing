#!/usr/bin/env python3
"""Build a call sheet: the top-tier employers that still lack a named staff-benefits contact, with a two-minute script to ask who it is.

A short phone call is the most reliable way to learn who looks after staff welfare or benefits, and the person can agree to receive the
email, which is consent-based. Reads the master (read-only): companies in reach tiers P1 and P2 (rank_employer_reach.py) whose chosen
recipient is the company team, or a named person who is not HR or the general manager, and that publish a phone number.
Writes outputs/contacts/Silverleaf Call Sheet - <date>.xlsx:
- "Script": the call script in English and Kiswahili (native-speaker review needed) and the rules for the call;
- "Calls": one row per company, most reach first, with the numbers to call, who to ask for, and blank columns for the answer.
Fill in the blank columns and import them with import_call_results.py. Nothing is dialled or sent.

--vicoba builds outputs/contacts/Silverleaf VICOBA Call Sheet - <date>.xlsx (docs/plans/vicoba-integration-plan.md decision 3;
docs/methodology/contact-research.md, "VICOBA savings groups"): one row per introducer (an NGO field office, bank branch or faith
office linked through data/reference/vicoba-introducers.json, or a CDO office once its council's introduction is recorded) whose
office publishes a phone, listing the groups it covers, and one row per group officer whose role phone a public body or the group
publishes for the role. Only complete groups appear, women's groups first. The sheet holds office numbers and numbers published for
the role only, each with its basis and source; the builder fails if any other number would reach it. There is no answer column for
a name or a number: a number given on the call goes only in the partnerships call log outside the repository. Import the answers
with import_call_results.py --vicoba.
"""
from __future__ import annotations

import argparse
import json
import re
import sqlite3
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts" / "master"))
import resolve_master_recipients as R  # noqa: E402

MASTER = ROOT / "outputs" / "master" / "Silverleaf Master Database.sqlite"
OUT = ROOT / "outputs" / "contacts"
PHONE = re.compile(r"\+?\d[\d ()-]{7,}\d")
HEADERS = ["priority", "reach_tier", "reach_score", "organisation", "organisation_id", "kind", "phones", "company_email", "current_recipient",
           "ask_for", "answered_by_name", "answered_by_role", "contact_name", "contact_role", "contact_work_email", "contact_work_phone",
           "agreed_to_receive_email (yes/no)", "outcome (reached / call back / wrong number / not interested / no answer)", "call_date", "caller", "notes"]
SCRIPT_EN = """Good morning / afternoon. My name is [caller], from Silverleaf Academy, the school in Arusha, Usa River and Boma Ng'ombe.
We are setting up partnerships with employers to provide an education benefit for the children of their staff.
Could you tell me who looks after staff welfare or benefits at [company]?
(If named) Thank you. Could I have their work email, so that Mariam Haji, our Marketing and Partnership Coordinator, can send them a short note?
Is it all right for us to email them about this?
Thank you very much for your help."""
SCRIPT_SW = """Habari za asubuhi / mchana. Naitwa [jina], kutoka Silverleaf Academy, shule iliyoko Arusha, Usa River na Boma Ng'ombe.
Tunaanzisha ushirikiano na waajiri ili kutoa fursa ya elimu kwa watoto wa wafanyakazi wao.
Je, unaweza kunieleza ni nani anayeshughulikia ustawi au marupurupu ya wafanyakazi katika [kampuni]?
(Akitajwa) Asante. Naweza kupata barua pepe yake ya kazini, ili Mariam Haji, Mratibu wetu wa Masoko na Ushirikiano, amtumie ujumbe mfupi?
Je, ni sawa tukimtumia barua pepe kuhusu jambo hili?
Asante sana kwa msaada wako."""
WELFARE_EN = """Good morning / afternoon. My name is [caller], from Silverleaf Academy, the school in Arusha, Usa River and Boma Ng'ombe.
We would like to set up a partnership with homes and programmes like [organisation] on the education of the children in their care.
Could you tell me who leads [organisation], or who decides on the children's schooling?
(If named) Thank you. Could I have their work email, so that Mariam Haji, our Marketing and Partnership Coordinator, can send them a short note?
Is it all right for us to email them about this?
Thank you very much for your help."""
WELFARE_SW = """Habari za asubuhi / mchana. Naitwa [jina], kutoka Silverleaf Academy, shule iliyoko Arusha, Usa River na Boma Ng'ombe.
Tungependa kuanzisha ushirikiano na vituo na programu kama [kituo] kuhusu elimu ya watoto walio chini ya uangalizi wao.
Je, unaweza kunieleza ni nani kiongozi wa [kituo], au ni nani anayeamua kuhusu masomo ya watoto?
(Akitajwa) Asante. Naweza kupata barua pepe yake ya kazini, ili Mariam Haji, Mratibu wetu wa Masoko na Ushirikiano, amtumie ujumbe mfupi?
Je, ni sawa tukimtumia barua pepe kuhusu jambo hili?
Asante sana kwa msaada wako."""
RULES = ["Call office numbers in working hours; one call, and at most one call-back. Stop at once if they are not interested.",
         "Ask only for the work name, role and work email or phone of the person who handles staff welfare or benefits. Never ask about anyone's children or family.",
         "Write down only what the person on the phone gives, and whether they agreed to the email. Do not guess spellings: ask them to spell names and emails.",
         "Do not promise discounts or terms on the call; the offer comes in the follow-up email after the first reply.",
         "The Kiswahili script is a draft for native-speaker review."]


def phones(*values) -> str:
    found = []
    for v in values:
        for m in PHONE.findall(str(v or "")):
            m = re.sub(r"[ ()-]", "", m)
            if len(re.sub(r"\D", "", m)) >= 9 and m not in found:
                found.append(m)
    return ", ".join(found[:3])


def save(rows, script_en: str, script_sw: str, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Script"
    ws["A1"] = f"Silverleaf call sheet, {date.today().isoformat()}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A3"], ws["A4"] = "English", script_en
    ws["A6"], ws["A7"] = "Kiswahili (draft for native-speaker review)", script_sw
    ws["A9"] = "Rules for the call"
    for i, rule in enumerate(RULES, 10):
        ws[f"A{i}"] = f"- {rule}"
    for cell in ("A3", "A6", "A9"):
        ws[cell].font = Font(bold=True)
    for cell in ("A4", "A7"):
        ws[cell].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = ws.row_dimensions[7].height = 120
    ws.column_dimensions["A"].width = 140
    calls = wb.create_sheet("Calls")
    calls.append(HEADERS)
    for row in rows:
        calls.append(row)
    widths = [8, 10, 10, 40, 16, 22, 34, 32, 40, 40, 22, 22, 26, 26, 32, 20, 18, 30, 12, 16, 40]
    for i, w in enumerate(widths, 1):
        calls.column_dimensions[calls.cell(1, i).column_letter].width = w
    fill = PatternFill("solid", fgColor="FFF2CC")
    for col in range(11, len(HEADERS) + 1):
        calls.cell(1, col).fill = fill
    calls.freeze_panes = "E2"
    if rows:
        tab = Table(displayName="TCalls", ref=f"A1:{calls.cell(1, len(HEADERS)).column_letter}{len(rows) + 1}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        calls.add_table(tab)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def welfare_sheet(run_id: str) -> int:
    """High- and medium-priority homes and programmes whose draft goes to the organisation route, with a phone to call."""
    db = ROOT / "outputs" / "runs" / run_id / "lead-database.sqlite"
    con = sqlite3.connect(f"file:{db.as_posix()}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    phones_by_org = {}
    for c in con.execute("SELECT organisation_id, organisation_phone, role_phone FROM contacts"):
        phones_by_org.setdefault(c["organisation_id"], []).extend([c["organisation_phone"], c["role_phone"]])
    rows = []
    for r in con.execute("""SELECT o.*, p.recipient, p.review_status FROM organisations o JOIN outreach_plans p USING(organisation_id)
                            WHERE p.target_type='organisation' AND o.priority IN ('High','Medium') AND o.segment NOT IN ('Welfare funder','Out of scope')
                            ORDER BY CASE o.priority WHEN 'High' THEN 0 ELSE 1 END, o.distance_km, o.name"""):
        if "(organisation route)" not in (r["recipient"] or ""):
            continue  # already addressed to a named leader
        numbers = phones(r["phone"], *phones_by_org.get(r["organisation_id"], []))
        if not numbers:
            continue
        rows.append([len(rows) + 1, r["priority"], r["distance_km"] if r["distance_km"] is not None else "", r["name"], r["organisation_id"], r["segment"],
                     numbers, r["email"] or "", "organisation team (no named person)", "the director or manager of the home or programme"] + [""] * 11)
    path = ROOT / "outputs" / "runs" / run_id / f"Silverleaf Welfare Call Sheet - {date.today().isoformat()}.xlsx"
    save(rows, WELFARE_EN, WELFARE_SW, path)
    print(f"{len(rows)} organisations to call: {path}")
    return 0


# --- VICOBA savings groups (decided 13 October 2026) ------------------------------------------------------------------------------
VICOBA_HEADERS = ["order", "row_kind", "women_group_first", "organisation", "organisation_id", "groups_covered", "group_ids", "ward", "council",
                  "number", "number_basis", "number_source", "ask_for", "channel", "band",
                  "reached (office / Mwenyekiti / Katibu / Mweka Hazina / other role)",
                  "outcome (reached / call back / wrong number / not interested / no answer)", "committee_meeting_agreed (yes/no)",
                  "letter_or_sms_agreed (yes/no)", "introduces_groups (yes/no; introducer rows)", "meeting_day", "meeting_time",
                  "member_list_offered (yes/no; always refused)", "call_date", "caller", "notes (no numbers, no names)"]
VICOBA_ANSWERS_FROM = 16
VICOBA_OFFICER_SW = """Habari, naitwa [jina] kutoka Silverleaf Academy, Arusha. Nimepata namba hii kupitia [chanzo kilichoichapisha].
Je, ninaongea na [Mwenyekiti/Katibu] wa [Kikundi]?
Tungependa kuomba dakika 10 hadi 15 katika kikao chenu cha wiki, kuzungumza kwa Kiswahili kuhusu maandalizi ya mtoto kuanza shule. Je, kamati ingependa hilo?
Mnakutana siku gani na saa ngapi?
Je, ni sawa nikutumie barua fupi au ujumbe mfupi (SMS) wa kuthibitisha?"""
VICOBA_OFFICER_EN = """(English meaning) Hello, my name is [caller], from Silverleaf Academy, Arusha. I found this number through [the source that published it].
Am I speaking with the [chair/secretary] of [group]?
We would like to ask for 10 to 15 minutes at one of your weekly meetings, to talk in Kiswahili about getting children ready for school. Would the committee like that?
On which day and at what time do you meet?
Is it all right if I send you a short letter or an SMS to confirm?"""
VICOBA_INTRODUCER_EN = """Good morning / afternoon. My name is [caller], from Silverleaf Academy, the school in Arusha, Usa River and Boma Ng'ombe.
We would like to offer the savings groups your [programme / branch / office] works with a short talk in Kiswahili, at one of their regular meetings, on getting children ready for school and planning for their education.
Could you introduce us to one or two group committees, or tell us who in your office works with the groups?
We do not ask for member lists or any member's details; each group decides for itself.
Thank you very much for your help."""
VICOBA_INTRODUCER_SW = """Habari za asubuhi / mchana. Naitwa [jina], kutoka Silverleaf Academy, shule iliyoko Arusha, Usa River na Boma Ng'ombe.
Tungependa kuvipa vikundi vya kuweka akiba na kukopa mnavyofanya navyo kazi elimu fupi ya dakika 10 hadi 15 kwa Kiswahili, katika kikao chao cha kawaida, kuhusu maandalizi ya mtoto kuanza shule na kupanga elimu yake.
Je, mnaweza kututambulisha kwa kamati ya kikundi kimoja au viwili, au kutueleza ni nani ofisini kwenu anayefanya kazi na vikundi?
Hatuombi orodha ya wanachama wala taarifa zao; kila kikundi kitaamua chenyewe.
Asante sana kwa msaada wako."""
VICOBA_RULES = ["Call only the numbers on this sheet: office numbers and numbers published for the role. Working hours; one call and at most one call-back. Stop at once if they are not interested.",
                "Women's groups come first in every wave (their rows are marked).",
                "Never ask for, or write down, a member's name, phone, savings, loans or children. If a member list is offered, refuse it politely and write 'yes' under member_list_offered.",
                "A number someone gives you on the call is never written on this sheet: write it only in the partnerships call log outside the repository. The import refuses a sheet with a phone number or an email in any answer.",
                "Write no names: the role (Mwenyekiti, Katibu, Mweka Hazina) is enough.",
                "No terms on the call: the offer follows in writing after the committee replies. Never offer a loan, a subsidy or fee financing.",
                "Never visit a group meeting unannounced: a VICOBA meeting is private and cash changes hands.",
                "No ward official is called before the council's introduction (GA01) is recorded in the government run.",
                "The Kiswahili scripts are drafts for native-speaker review."]


def vicoba_rows(database: Path, government_db: str | None) -> list[list]:
    """The VICOBA call rows: introducer office lines and group officers' role phones published for the role, nothing else."""
    sys.path.insert(0, str(ROOT / "scripts" / "master"))
    import add_vicoba_groups as V  # noqa: E402
    import link_vicoba_introducers as L  # noqa: E402
    con = sqlite3.connect(f"file:{database.as_posix()}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    groups = V.load_vicoba_groups(con)
    order = V.wave_order(groups)
    props = {oid: L.band_properties(g) for oid, g in groups.items()}
    try:
        sys.path.insert(0, str(ROOT / "scripts" / "messaging"))
        from first_touch import vicoba_case  # noqa: E402  (T4: the band rule shared with the exporter)
    except ImportError:
        vicoba_case = None
    basis = {r[0] for r in con.execute("SELECT entity_id FROM facts WHERE entity_type='contact' AND field='phone_basis' AND value='published_for_role'")}
    risky = {r[0] for r in con.execute("SELECT entity_id FROM facts WHERE entity_type='contact' AND field='pdpa_risk' AND value='risky'")}
    officers = {}
    for c in con.execute("SELECT * FROM contacts"):
        if c["organisation_id"] in groups and phones(c["role_phone"]) and c["contact_id"] in basis and c["contact_id"] not in risky:
            officers.setdefault(c["organisation_id"], []).append(dict(c))
    orgs = {r["organisation_id"]: dict(r) for r in con.execute("SELECT organisation_id, name, segment, phone, source_url FROM organisations")}
    office_lines = {}
    for c in con.execute("SELECT organisation_id, organisation_phone, source_url FROM contacts"):
        if phones(c["organisation_phone"]):
            office_lines.setdefault(c["organisation_id"], []).append((c["organisation_phone"], c["source_url"] or ""))
    gov = L.GovRuns(L.parse_gov(government_db))
    position = {oid: i for i, oid in enumerate(order)}

    def band(oid):
        if vicoba_case is None:
            return ""
        try:
            return str(vicoba_case(groups[oid]))
        except Exception:  # the band is T4's; a failure leaves the column blank
            return ""

    rows, allowed = [], set()
    introducers: dict[str, list] = {}
    for oid in order:
        g, p = groups[oid], props[oid]
        if g["record_status"] != "complete":
            continue  # held groups (incomplete or stale) are not called
        if p["officer_route"] == "published_for_role":
            for c in officers.get(oid, []):
                number = phones(c["role_phone"])
                allowed |= {re.sub(r"\D", "", n) for n in number.split(", ")}
                rows.append([position[oid], "group_officer", "yes" if g["women_group"] else "", g["name"], oid, g["name"], oid,
                             L.last(g, "ward"), L.last(g, "council"), number, "published_for_role: role phone published for the role",
                             c["source_url"] or "", c["role"], "call", band(oid)])
        # A welfare-run introducer is called within its own run's conversation, never from this sheet; an introducer that is not
        # ready (council introduction pending, to confirm, outside the region, another conversation first) is not called.
        if p["introducer_status"] == "linked" and p["linked_introducer_kind"] not in ("employer", "group_published") \
                and p["linked_introducer_db"] != "welfare" and p["introducer_ready"] == "yes" and not p["introducer_outcome"].startswith("declined"):
            introducers.setdefault(p["linked_introducer_ref"], []).append(oid)
    for ref, covered in introducers.items():
        covered.sort(key=position.get)
        first = covered[0]
        kind = props[first]["linked_introducer_kind"]
        if ref.startswith("gov:"):
            office, _ = gov.office(ref)
            numbers = [(office["phone"], f"government run {office['run']}: office phone")] if office and phones(office["phone"]) else []
            name = office["name"] if office else ref
        else:
            org = orgs.get(ref) or {}
            numbers = [(org.get("phone"), org.get("source_url") or "")] if phones(org.get("phone")) else []
            numbers += office_lines.get(ref, [])
            name = org.get("name", ref)
        if not numbers:
            continue
        number = phones(*[n for n, _ in numbers])
        allowed |= {re.sub(r"\D", "", n) for n in number.split(", ")}
        ask = {"cdo_ward": "the Ward Community Development Officer (Afisa Maendeleo ya Jamii wa Kata)",
               "cdo_council": "the Head of Community Development (Mkuu wa Idara ya Maendeleo ya Jamii)",
               "bank_branch": "the branch manager, or whoever looks after group accounts",
               "faith_office": "the office that hosts the savings groups"}.get(kind, "the programme or field officer who works with savings groups")
        rows.append([position[first], "introducer", "yes" if any(groups[o]["women_group"] for o in covered) else "", name, ref,
                     "; ".join(groups[o]["name"] for o in covered), "; ".join(covered),
                     "; ".join(dict.fromkeys(L.last(groups[o], "ward") for o in covered)), "; ".join(dict.fromkeys(L.last(groups[o], "council") for o in covered)),
                     number, "office: published by the organisation", "; ".join(dict.fromkeys(s for _, s in numbers if s)), ask,
                     props[first]["channel"], band(first)])
    rows.sort(key=lambda r: (r[0], r[1] != "group_officer"))
    # The guard: every number on the sheet is an office line or a role phone published for the role, as read above.
    withheld = {re.sub(r"\D", "", n) for oid in groups for n in phones(orgs.get(oid, {}).get("phone")).split(", ") if n}
    for r in rows:
        for n in r[9].split(", "):
            d = re.sub(r"\D", "", n)
            if d not in allowed or d in withheld or r[10].split(":")[0] not in ("office", "published_for_role"):
                raise SystemExit(f"VICOBA call sheet refused: a number on row '{r[3]}' is neither an office line nor published for the role.")
    for i, r in enumerate(rows, 1):
        r[0] = i
    con.close()
    return [r + [""] * (len(VICOBA_HEADERS) - len(r)) for r in rows]


def save_vicoba(rows, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Script"
    ws["A1"] = f"Silverleaf VICOBA call sheet, {date.today().isoformat()}"
    ws["A1"].font = Font(bold=True, size=13)
    blocks = [("Group officer (role phone published for the role): Kiswahili (draft for native-speaker review)", VICOBA_OFFICER_SW),
              ("Group officer: English meaning", VICOBA_OFFICER_EN),
              ("Introducer (NGO field office, bank branch, faith office, CDO office): English", VICOBA_INTRODUCER_EN),
              ("Introducer: Kiswahili (draft for native-speaker review)", VICOBA_INTRODUCER_SW)]
    row = 3
    for title, text in blocks:
        ws.cell(row, 1, title).font = Font(bold=True)
        ws.cell(row + 1, 1, text).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row + 1].height = 110
        row += 3
    ws.cell(row, 1, "Rules for the call").font = Font(bold=True)
    for i, rule in enumerate(VICOBA_RULES, row + 1):
        ws.cell(i, 1, f"- {rule}")
    ws.column_dimensions["A"].width = 140
    calls = wb.create_sheet("Calls")
    calls.append(VICOBA_HEADERS)
    for r in rows:
        calls.append(r)
    for i, w in enumerate([7, 14, 10, 36, 18, 40, 30, 18, 24, 24, 34, 34, 34, 9, 14] + [20] * 11, 1):
        calls.column_dimensions[calls.cell(1, i).column_letter].width = w
    fill = PatternFill("solid", fgColor="FFF2CC")
    for col in range(VICOBA_ANSWERS_FROM, len(VICOBA_HEADERS) + 1):
        calls.cell(1, col).fill = fill
    calls.freeze_panes = "E2"
    if rows:
        tab = Table(displayName="TVicobaCalls", ref=f"A1:{calls.cell(1, len(VICOBA_HEADERS)).column_letter}{len(rows) + 1}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        calls.add_table(tab)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def vicoba_sheet(database: Path, government_db: str | None, out: Path | None) -> int:
    rows = vicoba_rows(database, government_db)
    path = out or OUT / f"Silverleaf VICOBA Call Sheet - {date.today().isoformat()}.xlsx"
    save_vicoba(rows, path)
    print(json.dumps({"rows": len(rows), "group_officers": sum(1 for r in rows if r[1] == "group_officer"),
                      "introducers": sum(1 for r in rows if r[1] == "introducer"), "file": str(path)}, indent=1))
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--database", type=Path, default=MASTER)
    parser.add_argument("--tiers", default="P1,P2")
    parser.add_argument("--run-id", help="a welfare run (for example arusha-welfare-2026-09): its high- and medium-priority homes and programmes")
    parser.add_argument("--vicoba", action="store_true", help="the VICOBA call sheet: introducer office lines and role phones published for the role")
    parser.add_argument("--government-db", help="with --vicoba: a government run database copy, <path> or <run-id>=<path> (read only)")
    parser.add_argument("--out", type=Path, help="with --vicoba: write the workbook here instead of outputs/contacts/")
    args = parser.parse_args()
    if args.vicoba:
        return vicoba_sheet(args.database, args.government_db, args.out)
    if args.run_id:
        return welfare_sheet(args.run_id)
    tiers = [t.strip() for t in args.tiers.split(",")]
    con = sqlite3.connect(f"file:{args.database.as_posix()}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    contacts = {r["contact_id"]: dict(r) for r in con.execute("SELECT * FROM contacts")}
    by_org = {}
    for c in contacts.values():
        by_org.setdefault(c["organisation_id"], []).append(c)
    rows = []
    for r in con.execute("""SELECT e.tier, e.score, o.*, p.segment AS plan_segment, p.target_type, p.target_id, p.target_name, p.recipient_role
                            FROM employer_reach e JOIN organisations o USING(organisation_id)
                            JOIN outreach_plans p ON p.organisation_id = e.organisation_id AND p.selection = 'Candidate for review'
                            WHERE e.tier IN (%s) ORDER BY e.tier, e.score DESC, o.name""" % ",".join("?" * len(tiers)), tiers):
        role = r["recipient_role"] or ""
        rank = R.role_rank(role) if r["target_type"] == "contact" else None
        if r["target_type"] == "contact" and rank in (0, 2) and "@" in str(contacts.get(r["target_id"], {}).get("named_email") or ""):
            continue  # already a named HR or general manager with their own email
        numbers = phones(r["phone"], *[c["organisation_phone"] for c in by_org.get(r["organisation_id"], [])],
                         *[c["role_phone"] for c in by_org.get(r["organisation_id"], [])])
        if not numbers:
            continue
        ask = "the human resources or administration manager" if r["plan_segment"] in ("Healthcare employers", "Education employers", "Corporate employers") \
            or re.search(r"hotel|lodge|camp|resort", f"{r['name']} {r['segment'] or ''}", re.I) else "the owner, managing director or general manager"
        current = f"{r['target_name']} ({role})" if r["target_type"] == "contact" else "company team (no named person)"
        rows.append([len(rows) + 1, r["tier"], r["score"], r["name"], r["organisation_id"], r["plan_segment"], numbers,
                     re.split(r"[;, ]+", str(r["email"] or "").strip())[0] if r["email"] else "", current, ask] + [""] * 11)
    path = OUT / f"Silverleaf Call Sheet - {date.today().isoformat()}.xlsx"
    save(rows, SCRIPT_EN, SCRIPT_SW, path)
    print(f"{len(rows)} companies to call: {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
